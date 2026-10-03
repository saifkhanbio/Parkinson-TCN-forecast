"""Build a source bibliography and complete input-file citation crosswalk.

Only writes this citation package. Original data and frozen results are read-only.
The bibliography uses verified metadata; unknown dates remain unknown.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import xml.etree.ElementTree as ET
from collections import Counter
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches, Pt

BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[1]
DATE = '2026-10-01'
REFS = []


def sha(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def csv_write(path, rows, fields=None):
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields or list(rows[0]))
        w.writeheader()
        w.writerows(rows)


def ref(key, title, author, publisher, url, role, note, year='', kind='dataset', **extra):
    r = dict(id=key, title=title, authors=[{'literal': author}], publisher=publisher,
             url=url, role=role, note=note, year=str(year), kind=kind, doi='',
             accessed=DATE, verification='Official source or preserved original inspected', **extra)
    REFS.append(r)
    return r


def article(key, role, note):
    from acquire_references import ARTICLES
    target = ARTICLES[key].lower()
    for a in ET.parse(BASE / 'metadata/pubmed_articles.xml').findall('./PubmedArticle'):
        ids = {x.get('IdType'): x.text for x in a.findall('./PubmedData/ArticleIdList/ArticleId')}
        if ids.get('doi', '').lower() != target:
            continue
        p = a.find('./MedlineCitation/Article')
        authors = []
        for au in p.findall('./AuthorList/Author'):
            collective = au.findtext('CollectiveName')
            if collective:
                authors.append({'literal': collective})
            else:
                authors.append({'family': au.findtext('LastName', ''), 'given': au.findtext('ForeName', ''),
                                'initials': au.findtext('Initials', '')})
        # The collaborator groups are the journal's lead author names.
        groups = [x for x in authors if 'literal' in x]
        if key.startswith('gbd_') and groups:
            authors = groups
        r = dict(id=key, title=''.join(p.find('ArticleTitle').itertext()).rstrip('.'),
                 authors=authors, publisher='', journal=p.findtext('./Journal/ISOAbbreviation'),
                 volume=p.findtext('./Journal/JournalIssue/Volume'),
                 issue=p.findtext('./Journal/JournalIssue/Issue', ''), pages=p.findtext('./Pagination/MedlinePgn'),
                 year=p.findtext('./Journal/JournalIssue/PubDate/Year'), doi=ids['doi'],
                 pmid=ids['pubmed'], pmcid=ids.get('pmc', ''), url='https://doi.org/' + ids['doi'],
                 role=role, note=note, kind='article', accessed=DATE,
                 verification='PubMed ArticleIdList and PMC open-access article metadata agree')
        pmc = json.loads((BASE / ('metadata/' + key + '_pmc_skill.json')).read_text())['records'][0]
        assert pmc['doi'].lower() == target and str(pmc['pmid']) == ids['pubmed']
        assert not pmc['is_retracted']
        r['license'] = pmc['license_code']
        REFS.append(r)
        return r
    raise ValueError('PubMed DOI not found: ' + key)


def datacite(key, role, note):
    d = json.loads((BASE / ('metadata/' + key + '_datacite.json')).read_text())['data']['attributes']
    r = ref(key, d['titles'][0]['title'], '', d['publisher'], 'https://doi.org/' + d['doi'],
            role, note, year=d['publicationYear'])
    r['doi'] = d['doi']
    r['authors'] = [({'literal': x['name'].strip()} if x.get('nameType') != 'Personal' else
                     {'family': x['familyName'], 'given': x['givenName'],
                      'initials': ''.join(t[0] for t in x['givenName'].split())}) for x in d['creators']]
    r['verification'] = 'DataCite registered creators, title, publisher, publication year and DOI'
    return r


def make_catalog():
    ihme = 'Institute for Health Metrics and Evaluation'
    un = 'United Nations, Department of Economic and Social Affairs, Population Division'
    ga = 'General Authority for Statistics'
    moh = 'Ministry of Health, Kingdom of Saudi Arabia'
    chi = 'Council of Health Insurance'
    shc = 'Saudi Health Council'
    who = 'World Health Organization'
    used, method, audit, context, pending = ('analysis_source', 'source_methods', 'supplementary_audit',
                                           'documentation_or_background', 'catalog_only_not_acquired')
    ref('D1', 'Global Burden of Disease Study 2023 (GBD 2023) Results',
        'Global Burden of Disease Collaborative Network', ihme,
        'https://vizhub.healthdata.org/gbd-results/', used,
        'Primary age-specific outcomes, native numbers and uncertainty bounds; separate age-standardized-rate benchmark. '
        'Dataset DOI is not inferred from other GHDx products. Original export dates remain unverified. '
        'Native citation.txt is preserved verbatim as citation evidence; final publication-year field requires confirmation.')
    article('gbd_nonfatal', method, 'Cite for prevalence, incidence, YLDs, DALYs and non-fatal estimation methods.')
    article('gbd_fatal', method, 'Cite for cause-specific deaths and YLL estimation methods.')
    article('gbd_demography', method, 'Cite when explaining GBD population/mortality methods or denominator coherence; '
            'GBD-implied populations in this study are calculated from matching counts and rates, not a separate demographic export.')
    datacite('hierarchy', used, 'Cause, location and REI workbook dated 23 October 2025. '
             'World Bank classifications here are carried by the IHME hierarchy; no separate World Bank outcome dataset was used.')
    ref('D2', 'World Population Prospects 2024: Online Edition', un, 'United Nations',
        'https://population.un.org/wpp/', used,
        'Medium-variant age-sex population, official notes and correction archive; probabilistic total and age-sex workbooks '
        'used in supplementary demographic uncertainty audits. Values originally in thousands. Marginal quantiles are not joint draws.', year=2024)
    ref('wpp_methods', 'World Population Prospects 2024: Methodology of the United Nations population estimates and projections',
        un, 'United Nations', 'https://population.un.org/wpp/assets/Files/WPP2024_Methodology-Report_Final.pdf', method,
        'Suggested citation in the retained report: UN DESA/POP/2024/DC/NO. 10, July 2024 [Advance unedited version]. '
        'CC BY 3.0 IGO.', year=2024, kind='report')
    ref('D3', 'Population: DF_PSS_DEM_POP, version 1.0', 'GCC Statistical Centre', 'GCC-Stat',
        'https://dp.marsa.gccstat.org/dataset/population', used,
        'Age-sex-nationality demographic comparisons and denominator checks. Original public CSV retrieved 26 September 2026; '
        'reference periods vary. Original publication/update date not established. Saudi values overlap GASTAT; not independent validation.')
    ref('D4', 'Population Estimates Publication 2024 and accompanying Population estimates EN workbook', ga, ga,
        'https://www.stats.gov.sa/documents/d/guest/population-estimates-en', used,
        'Official Saudi 2023/2024 age-sex-nationality estimates used for denominator checks. Raw copies retrieved 26 September 2026. '
        '2024 is the product edition, not an inferred publication date. Indexed hosting dates require confirmation against the exact retained file.')
    ref('gastat_methods', 'Methodology and Quality Report for Population Projections and Estimates Statistics', ga, ga,
        'https://www.stats.gov.sa/en/w/methodology-and-quality-report-for-population-projections-and-estimates-statistics', method,
        'Definitions, population coverage and estimation methods; PDF downloaded 1 October 2026. Publication date not stated in retained copy.', kind='report')
    for year in [2023, 2024]:
        ref('D5' if year == 2023 else 'D6', 'Statistical Yearbook ' + str(year), moh, moh,
            f'https://www.moh.gov.sa/Ministry/Statistics/book/Documents/Statistical-Yearbook-{year}.xlsx', context,
            ('Neurology workforce Table 2-5; all-condition rehabilitation sheet 43. Also Table 1-3: 2022 population by age, sex and nationality, credited to GASTAT.'
             if year == 2023 else 'Neurology workforce Table 2-19; all-condition rehabilitation Table 4-55.') +
            ' Edition year is in the title; original publication date not established. Retrieved 26 September 2026. '
            'Health-system descriptive context is not Parkinson-specific service utilization.', kind='report')
    ref('gastat_series', 'Population estimates by gender, nationality, and age group 2010–2024', ga, ga,
        'https://www.stats.gov.sa/', audit,
        'User-supplied CSV audited 1 October 2026: 840 values agree with the retained GCC-Stat extract. '
        'Exact dataset landing/download URL and original retrieval time are missing; institutional URL is a fallback only.')
    ref('moh_population', 'Population by Nationality, Gender, and Age groups, 2022G', moh, moh,
        'https://hdp.moh.gov.sa/', audit,
        'User-supplied CSV. All 68 age-sex-nationality values were matched to Statistical Yearbook 2023, Table 1-3 '
        '(source credited there to GASTAT). Cite D5 alongside this portal export. Exact export URL and download date unverified.')
    ref('who_mortality', 'WHO Mortality Database', who, who,
        'https://www.who.int/data/data-collection-tools/who-mortality-database', audit,
        'Public-file update 23 February 2026; downloaded 1 October 2026. ICD-10 parts 3 and 6, population, availability, '
        'country-code and notes files. File-specific versions differ. Saudi G20 extracts are unscored; 2024 is reserved. '
        'Registered underlying-cause deaths are not a clinical incidence or prevalence cohort. GBD input overlap remains unverified.', year=2026)
    ref('who_methods', 'WHO Mortality Database: Documentation', who,
        'World Health Organization, Department of Data, Digital Health, Analytics and AI',
        'https://www.who.int/data/data-collection-tools/who-mortality-database', method,
        'Document dated 23 February 2026, obtained from the official documentation ZIP. Cite for field formats, age coding and interpretation.',
        year=2026, kind='report')
    ref('chi_data', 'Open Data: health-insurance beneficiary statistics, selected quarterly releases 2021–2026', chi, chi,
        'https://www.chi.gov.sa/en/open-data/Pages/Indicators-and-statistics.aspx', audit,
        'Retained periods: 2021 Q4, 2022 Q1, 2023 Q4, 2024 Q1, 2024 Q2, 2025 Q4 and 2026 Q2. '
        'Exact versions are recorded by filename/hash in the crosswalk. Age boundaries vary; beneficiary snapshots are not person-time. '
        'Individual resource URLs and publication dates have not been recovered.')
    ref('chi_class', 'CHI beneficiary class to the Age group 2024', chi, chi,
        'https://www.chi.gov.sa/en/open-data/Pages/Indicators-and-statistics.aspx', audit,
        'CSV containing age-by-beneficiary-class counts; no sex field. 2024 is the reference year. '
        'Exact resource URL, update date and snapshot definition require custodian metadata.')
    ref('chi_nphies', 'Top 10 by gender — NPHIES', chi, chi,
        'https://www.chi.gov.sa/en/open-data/Pages/Indicators-and-statistics.aspx', audit,
        'User-supplied workbook, descriptive title retained. Visible top-10 summaries cover November 2022–November 2023. '
        'Embedded aggregate service records have a different time scope and no patient ID, age, sex or diagnosis fields. '
        'This is not a Parkinson patient dataset; exact resource/version URL is missing.')
    ref('chi_terms', 'Requirements for using the open data of the Council of Health Insurance', chi, chi,
        'https://www.chi.gov.sa/en/open-data/Pages/requirements.aspx', context,
        'Official attribution requirements consulted 1 October 2026. Cite CHI as the data source. '
        'A downloaded page is not patient-data access approval.', kind='webpage')
    datacite('nhrsp_ncd', context, 'Project CHRS0762022. Only a codebook was supplied; no participant observations used. '
             'Registered publication year 2021; survey fieldwork April 2019–February 2020. No Parkinson-specific item identified in the inspected codebook.')
    datacite('nhrsp_hiss', context, 'Project CHRS0622021. Registered creators: Nasser BinDhim, Nora Al thumairi, Mada Basyouni; '
             'publisher Sharik Health Association; publication year 2020. Only the codebook and data-collection tool were supplied, not participant records.')
    article('safiri', context, 'Legacy D7. Article and supplements retained as literature/background. '
            'Older-release numerical comparisons are retired and are not restored by this citation package.')
    article('menasa', context, 'Legacy D8. Protocol and questionnaire documentation only; no registry participant data used.')
    article('saudi_genetics', context, 'Selected Saudi clinical/genetic study and public supplements reviewed for context; '
            'not a population denominator, national outcome series or independent model-validation cohort.')
    for key, title, author, url in [
        ('moh_access', 'Health Data Platform: Health Data Sharing Procedures', moh, 'https://hdp.moh.gov.sa/en/data-sharing'),
        ('chi_access', 'Data Sharing Request', chi, 'https://www.chi.gov.sa/en/open-data/Pages/access-data.aspx'),
        ('nhrsp_access', 'National Health Research and Studies Portal: access workflow and user guide', shc, 'https://shc.gov.sa/ar/EServices/Pages/nhrsp.aspx'),
        ('gastat_access', 'Microdata request and user guide', ga, 'https://www.stats.gov.sa/en/request-for-scientific-use-files'),
    ]:
        ref(key, title, author, author, url, context,
            'Access documentation only. A form, user guide or data-sharing agreement template is not approval or acquisition of patient records.', kind='webpage')
    iqvia = ref('iqvia_abstract', 'Effect of COVID-19 Pandemic on Utilization of Parkinson’s Medications in Saudi Arabia: A Repeated Cross-Sectional Study',
        'Orayj K', 'Movement Disorders',
        'https://www.mdsabstracts.org/abstract/effect-of-covid-19-pandemic-on-utilization-of-parkinsons-medications-in-saudi-arabia-a-repeated-cross-sectional-study/',
        context, 'Conference abstract, Mov Disord 2025;40(suppl 1). Saudi IQVIA sales 2019–2023 discussed in the abstract; '
        'underlying commercial data not acquired. DOI/abstract number not confirmed.', year=2025, kind='conference_abstract')
    iqvia['authors'] = [{'family':'Orayj','given':'K','initials':'K'}]
    ref('kfmc_report', 'Research report 2016: Parkinson Disease and Movement Disorder Registry in KFMC project listing', moh, moh,
        'https://www.moh.gov.sa/Ministry/MediaCenter/Publications/Pages/MOH-2016-VIEW.pdf', context,
        'Historical project listing only. Current registry operation and data access not established.', year=2016, kind='report')
    ref('chi_guidance', 'Parkinson Disease: indication update', chi, chi,
        'https://www.chi.gov.sa/Style%20Library/IDF_Branding/Indication/155%20-%20Parkinson%20Disease-Indication%20Update.pdf', context,
        'Clinical guidance reviewed for Saudi evidence context; not a patient dataset. Exact issue/revision date not established.', kind='report')
    for key in ['gbd_demographic_catalog', 'gbd_nonfatal_catalog', 'sdi_catalog']:
        datacite(key, pending, 'Catalog and citation metadata only. The DOI applies to this named product, '
                 'not automatically to a Results Tool query. No additional data from this catalog product were acquired for the analyses.')


def authors_text(r):
    names = [a.get('literal') or (a['family'] + ' ' + a.get('initials', '')) for a in r['authors']]
    return ', '.join(names if len(names) <= 6 else names[:3] + ['et al'])


def cite(r):
    if r['kind'] == 'article':
        return f"{authors_text(r)}. {r['title']}. {r['journal']} {r['year']};{r['volume']}:{r['pages']}. https://doi.org/{r['doi']}"
    typ = {'dataset': 'dataset', 'report': 'report', 'webpage': 'Internet', 'conference_abstract': 'conference abstract'}[r['kind']]
    year = r['year'] or 'publication date not stated'
    return f"{authors_text(r)}. {r['title']} [{typ}]. {r['publisher']}; {year}. {r['url']} (1 October 2026, date of source/citation review)."


def copy_documents():
    jobs = [
        ('supporting_data/2026-09-30_source_audit/raw/un_wpp2024_methodology.pdf', 'wpp_methods.pdf', 'wpp_methods'),
        ('supporting_data/2026-09-26/raw/GASTAT_Population_Estimates_2024_EN.pdf', 'gastat_population_2024.pdf', 'D4'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/gastat_population_methods.pdf', 'gastat_methods.pdf', 'gastat_methods'),
        ('supporting_data/2026-10-01_saudi_evidence/processed/WHO_Mortality_Database_Documentation.pdf', 'who_methods.pdf', 'who_methods'),
        ('More data/IHME_GBD_2023_HIERARCHIES_INFO_SHEET_Y2025M10D23.pdf', 'gbd_hierarchy_information_sheet.pdf', 'hierarchy'),
        ('More data/IHME-GBD_2023_DATA-e22a15d9-1/citation.txt', 'native_gbd_export_citation.txt', 'D1'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/moh_research_2016.pdf', 'kfmc_project_report_2016.pdf', 'kfmc_report'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/chi_parkinson_guidance.pdf', 'chi_parkinson_guidance.pdf', 'chi_guidance'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/iqvia_saudi_pd_2025_abstract.html', 'iqvia_abstract.html', 'iqvia_abstract'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/nhrsp_userguide.pdf', 'nhrsp_userguide.pdf', 'nhrsp_access'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/gastat_microdata_guide.pdf', 'gastat_microdata_guide.pdf', 'gastat_access'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/moh_data_request_form.pdf', 'moh_data_request_form.pdf', 'moh_access'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/moh_data_sharing.html', 'moh_data_sharing.html', 'moh_access'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/chi_access_data.html', 'chi_access_data.html', 'chi_access'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/shc_nhrsp_access.html', 'nhrsp_access.html', 'nhrsp_access'),
        ('supporting_data/2026-10-01_saudi_evidence/raw/gastat_microdata_request.html', 'gastat_access.html', 'gastat_access'),
    ]
    rows = []
    for source, filename, key in jobs:
        src, dst = ROOT / source, BASE / 'documents' / filename
        original_hash = sha(src)
        shutil.copyfile(src, dst)
        assert sha(dst) == original_hash
        rows.append(dict(file='documents/' + filename, reference_id=key, copied_from=source,
                         sha256=original_hash, bytes=src.stat().st_size,
                         operation='Exact copy of preserved reference document; not a new source retrieval'))
    csv_write(BASE / 'copied_document_manifest.csv', rows)


def bib_ris():
    bib, ris = [], []
    def esc(t):
        return str(t).replace('\\', '\\textbackslash{}').replace('&', '\\&').replace('%', '\\%').replace('_', '\\_')
    for r in REFS:
        authors = ' and '.join(('{' + esc(a['literal']) + '}' if 'literal' in a else
                                esc(a['family'] + ', ' + a['given'])) for a in r['authors'])
        fields = {'author': authors, 'title': '{' + esc(r['title']) + '}', 'year': r['year'],
                  'publisher': esc(r['publisher']), 'doi': r['doi'], 'url': r['url'], 'urldate': r['accessed'],
                  'journal': r.get('journal', ''), 'volume': r.get('volume', ''),
                  'number': r.get('issue', ''), 'pages': (r.get('pages') or '').replace('-', '--'),
                  'note': esc(r['role'] + '. ' + r['note'])}
        typ = 'article' if r['kind'] == 'article' else 'misc'
        bib.append('@' + typ + '{' + r['id'] + ',\n' + ',\n'.join('  ' + k + ' = {' + str(v) + '}' for k,v in fields.items() if v) + '\n}')
        rt = {'article': 'JOUR', 'dataset': 'DATA', 'report': 'RPRT', 'webpage': 'ELEC', 'conference_abstract': 'CONF'}[r['kind']]
        lines = ['TY  - ' + rt, 'ID  - ' + r['id']]
        for a in r['authors']:
            lines.append('AU  - ' + (a['literal'] + ',' if 'literal' in a else a['family'] + ', ' + a['given']))
        for k, v in [('TI', r['title']), ('PY', r['year']), ('PB', r['publisher']), ('DO', r['doi']),
                     ('UR', r['url']), ('Y2', DATE), ('JO', r.get('journal')), ('VL', r.get('volume')),
                     ('IS', r.get('issue')), ('SP', r.get('pages')), ('N1', r['role'] + '. ' + r['note'])]:
            if v: lines.append(k + '  - ' + v)
        ris.append('\n'.join(lines + ['ER  -', '']))
    (BASE / 'dataset_references.bib').write_text('\n\n'.join(bib) + '\n', encoding='utf-8')
    (BASE / 'dataset_references.ris').write_text('\n'.join(ris), encoding='utf-8')


def doc(title):
    d = Document()
    s = d.styles['Normal']; s.font.name = 'Times New Roman'; s.font.size = Pt(11)
    s.paragraph_format.space_after = Pt(6)
    for sec in d.sections:
        sec.top_margin = Inches(.8); sec.bottom_margin = Inches(.8)
    d.add_heading(title, 0)
    return d


def build_guide():
    nums = {r['id']: i + 1 for i,r in enumerate(REFS)}
    def c(*keys): return '[' + ', '.join(str(nums[k]) for k in keys) + ']'
    passages = [
        ('Primary outcomes and age-standardized benchmark',
         'Parkinson’s disease prevalence, incidence, deaths, years lived with disability, years of life lost and disability-adjusted life years were obtained from GBD 2023. '
         'The primary analysis used sex-specific age-group estimates for ages 45 years and older; age-standardized rates were evaluated in a separate benchmark. '
         'The extracted measures, metrics, age categories, locations and observation years are documented in Supplementary Material S1. '
         + c('D1', 'gbd_nonfatal', 'gbd_fatal')),
        ('Geographic metadata',
         'Cause and location identifiers were checked against the official GBD 2023 hierarchies. ' + c('hierarchy')),
        ('Population estimates and projections',
         'Demographic features and population scenarios used the 2024 revision of United Nations World Population Prospects. '
         'Medium-variant age-sex estimates and projections, accompanying notes and the official correction archive were retained. '
         'Supplementary uncertainty analyses used the published probabilistic age-sex population quantiles, which are marginal summaries rather than joint population trajectories. '
         + c('D2', 'wpp_methods')),
        ('National population and denominator checks',
         'Saudi population checks used GASTAT population estimates and the GCC-Stat population dataflow. '
         'These sources may reproduce the same underlying national statistics and were not treated as independent replications. '
         'GBD-implied population denominators were calculated from matched native counts and rates; no separately downloaded GBD demographic dataset was assumed. '
         + c('D3', 'D4', 'gastat_methods', 'D1', 'gbd_demography')),
        ('Health-system context — use only where reported',
         'Descriptive neurology workforce and rehabilitation indicators came from the Saudi Ministry of Health Statistical Yearbooks. '
         'Rehabilitation activity covered all conditions and was not interpreted as Parkinson-specific utilization. ' + c('D5', 'D6')),
        ('New Saudi population audit — supplementary only',
         'Additional population exports were inspected for demographic consistency. The MOH 2022 age-sex-nationality values matched the corresponding table in the 2023 Statistical Yearbook, '
         'which attributes the source to GASTAT. The GASTAT 2010–2024 export duplicated values already represented in the GCC-Stat extract. '
         'These checks did not introduce new Parkinson outcomes into model fitting. ' + c('gastat_series', 'moh_population', 'D5', 'D3')),
        ('WHO mortality acquisition audit — supplementary only',
         'Saudi records coded G20 were extracted from the WHO Mortality Database public files and inspected with the accompanying age-format and availability documentation. '
         'These registered underlying-cause deaths have not been used to score forecasts. Their coverage and possible overlap with GBD input data require evaluation before any claim of independent validation. '
         + c('who_mortality', 'who_methods')),
        ('Insurance and claims audit — supplementary only',
         'CHI beneficiary summaries and the NPHIES workbook were examined for potential supporting evidence. The retained releases provide insurance-coverage or aggregate service information '
         'with varying age classifications; they do not supply the age-sex Parkinson outcome series required for primary forecast validation. ' + c('chi_data', 'chi_class', 'chi_nphies')),
        ('Survey and registry documentation — supplementary/background only',
         'HISS and ABSHER NCD codebooks and the MENASA registry questionnaire were reviewed to assess potential future data availability. '
         'No participant-level observations from these projects were acquired for training or validation. ' + c('nhrsp_hiss', 'nhrsp_ncd', 'menasa')),
        ('Saudi clinical evidence — background only',
         'A published Saudi genetic study and its public supplementary tables were reviewed for clinical context; its selected clinical sample was not used as a national prevalence denominator or forecast-validation cohort. '
         + c('saudi_genetics')),
    ]
    d = doc('Dataset citations and reporting guide')
    d.add_paragraph('Citation audit completed 1 October 2026. This supplement covers all retained dataset types and their documentation. '
                    'Acquisition or documentation review does not imply inclusion in model fitting. Original analytic results remain unchanged.')
    d.add_heading('How to use this supplement', 1)
    d.add_paragraph('Use the reporting passages below only for work actually described in the manuscript. The bracketed numbers refer to this supplement’s bibliography; '
                    'renumber them in order of first appearance when merging with the full manuscript bibliography. Cite original providers in derived tables and figure captions. '
                    'The file-to-citation crosswalk uses supplementary archive member names, not machine-specific storage paths.')
    for title, text in passages:
        d.add_heading(title, 2); d.add_paragraph(text)
    d.add_heading('Suggested source acknowledgments', 1)
    d.add_paragraph('The authors acknowledge IHME and the GBD Collaborative Network; the United Nations Department of Economic and Social Affairs, Population Division; '
                    'GASTAT; GCC-Stat; the Saudi Ministry of Health; WHO; and CHI for the source material used or reviewed. '
                    'Specify each provider’s actual contribution and omit providers whose material is not reported. '
                    'Analyses and interpretations are those of the authors and do not imply endorsement by the data providers.')
    d.add_paragraph('For CHI-derived displays: “Source: Council of Health Insurance (CHI), Saudi Arabia; authors’ calculations.” '
                    'Retain the exact quarterly reference period. ' + c('chi_terms'))
    d.add_paragraph('A citation does not confer a new redistribution licence. Six article PDFs are marked CC BY in the downloaded PMC metadata; '
                    'the retained UN methods report specifies CC BY 3.0 IGO. Other source terms remain provider-specific. '
                    'This package supplies citations and reference documents; it does not publish restricted data or submit access requests.')
    d.add_heading('References', 1)
    txt = ['Dataset citations and reporting guide', 'Reviewed 1 October 2026', '']
    for title, text in passages: txt.extend([title, text, ''])
    for i,r in enumerate(REFS, 1):
        text = f"{i}. {cite(r)}"
        d.add_paragraph(text)
        d.add_paragraph('Reference key: ' + r['id'] + '. Reporting role: ' + r['role'] + '. ' + r['note'])
        txt.extend([text, r['id'] + ': ' + r['note'], ''])
    d.save(BASE / 'dataset_citations_and_reporting.docx')
    (BASE / 'dataset_citations_and_reporting.txt').write_text('\n'.join(txt), encoding='utf-8')
    csv_write(BASE / 'citation_placement.csv', [dict(section=t, suggested_text=p) for t,p in passages])


def match_moh_population():
    """Match every CSV datum to an explicit cell in the original MOH table."""
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    source = ROOT / 'supporting_data/2026-09-26/raw/saudi_moh_yearbook_2023.xlsx'
    with ZipFile(source) as z:
        strings = [''.join(e.itertext()) for e in ET.fromstring(z.read('xl/sharedStrings.xml')).findall('s:si', ns)]
        table = {}
        for c in ET.fromstring(z.read('xl/worksheets/sheet4.xml')).findall('.//s:c', ns):
            value = c.findtext('s:v', '', ns)
            table[c.get('r')] = strings[int(value)] if value and c.get('t') == 's' else value
    ages = {table['A' + str(i)]: i for i in range(6, 23)}
    assert 'Table 1-3' in table.values()
    cols = {('Saudi', 'MALE'): 'B', ('Saudi', 'FEMALE'): 'C',
            ('NonSaudi', 'MALE'): 'E', ('NonSaudi', 'FEMALE'): 'F'}
    rows = []
    with (ROOT / 'data/raw/moh/Population by Nationality, Gender, and Age groups, 2022G.csv').open(encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            age = row['Age Group'].strip().lstrip("'")
            cell = cols[(row['Variable1'].strip(), row['Variable2'].strip())] + str(ages[age])
            value, original = int(row['Variable_Value']), int(table[cell])
            assert value == original
            rows.append(dict(age=age, nationality=row['Variable1'].strip(), sex=row['Variable2'].strip(),
                             csv_value=value, yearbook_value=original, source_table='1-3',
                             sheet_xml='xl/worksheets/sheet4.xml', source_cell=cell, equal=True))
    assert len(rows) == 68
    csv_write(BASE / 'metadata/moh_2022_yearbook_cell_match.csv', rows)
    return len(rows)


def crosswalk():
    refs = {r['id']: r for r in REFS}
    assigned, provenance = {}, {}
    def assign(path, ids, kind='source', detail=''):
        p = Path(path)
        rel = p.relative_to(ROOT).as_posix() if p.is_absolute() else p.as_posix()
        if (ROOT / rel).is_file():
            assert all(k in refs for k in ids), (rel, ids)
            assigned[rel] = (ids, kind, detail)
    # Import known provenance rather than silently resetting retrieval dates.
    for manifest in (ROOT / 'supporting_data').glob('*/manifest*.json'):
        raw = json.loads(manifest.read_text(encoding='utf-8-sig'))
        if not isinstance(raw, list): continue
        for row in raw:
            if not isinstance(row, dict) or not row.get('sha256'): continue
            path = row.get('path', '')
            if path:
                path = path.replace('\\', '/')
                p = Path(path)
                if not p.is_absolute(): p = manifest.parent / p
                if p.exists() and ROOT in p.parents:
                    provenance[p.relative_to(ROOT).as_posix()] = row
    legacy = {'D1':'D1','D2':'D2','D3':'D3','D4':'D4','D5':'D5','D6':'D6','D7':'safiri','D8':'menasa','D9':'sdi_catalog'}
    with (ROOT / 'study_design/data_file_citations.csv').open(encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            ids = [legacy[x.strip()] for x in row['reference_labels'].split(';')]
            assign(row['file_path'], ids, row['file_kind'], row['citation_status'])
    for p in (ROOT / 'age_standard').iterdir():
        if p.suffix.lower() in ('.csv','.xlsx'): assign(p, ['D1'], 'analytic_input', 'Original workbook extraction history incomplete')
    for p in (ROOT / 'More data').rglob('*'):
        if p.is_file() and ':Zone.Identifier' not in p.name:
            assign(p, ['hierarchy'] if 'HIERARCH' in p.name else ['D1'], 'user_supplied_source',
                   'Native file; original extraction date unverified; review date is not download date')
    for folder, key in [('2026-09-30_hierarchy_verification','hierarchy'),('2026-09-30_asr_verification','D1')]:
        for sub in ['raw','processed']:
            for p in (ROOT / 'supporting_data' / folder / sub).glob('*'):
                if p.is_file(): assign(p,[key], 'original_copy' if sub=='raw' else 'derived', 'Verified source copy or derived metadata')
    audit_keys = {'gbd2023_nonfatal.html':'gbd_nonfatal','gbd2023_demographics.html':'gbd_demography',
                  'gbd2023_hierarchy_record.html':'hierarchy','gbd2023_demographic_record.html':'gbd_demographic_catalog',
                  'gbd2023_nonfatal_record.html':'gbd_nonfatal_catalog','gbd2023_catalog.html':'D1',
                  'un_wpp2024_methodology.pdf':'wpp_methods'}
    for sub in ['raw','processed']:
        for p in (ROOT / 'supporting_data/2026-09-30_source_audit' / sub).glob('*'):
            assign(p,[audit_keys.get(p.name,'D2')], 'source_document' if p.suffix in ('.pdf','.html','.json') else
                   ('source' if sub=='raw' else 'derived'), 'Source audit; probabilistic population products remain separate from joint draws')
    old_new = {'WHO_MDB':'who_mortality','MOH_ACCESS':'moh_access','CHI_ACCESS':'chi_access','NHRSP_ACCESS':'nhrsp_access',
               'GASTAT_ACCESS':'gastat_access','GASTAT_METHODS':'gastat_methods','ALMUBARAK_2015':'saudi_genetics',
               'MENASA_2025':'menasa','IQVIA_ABSTRACT':'iqvia_abstract','KFMC_PROJECT':'kfmc_report','CHI_GUIDANCE':'chi_guidance'}
    evidence = ROOT / 'supporting_data/2026-10-01_saudi_evidence'
    with (evidence / 'file_to_citation.csv').open(encoding='utf-8-sig') as f:
        for row in csv.DictReader(f):
            p = evidence / row['file']; key = old_new[row['reference_id']]
            if 'documentation' in p.name and key == 'who_mortality': key = 'who_methods'
            assign(p,[key], 'source', 'Acquired for Saudi evidence audit; no forecast scoring')
            prov = provenance.setdefault(p.relative_to(ROOT).as_posix(),{})
            prov.setdefault('url',row['url']);prov.setdefault('sha256',row['sha256'])
    for p in (evidence / 'processed').glob('*'):
        if p.name=='validation.json': continue
        key = 'who_methods' if p.suffix=='.pdf' else ('who_mortality' if p.name.startswith('who_') else 'saudi_genetics')
        assign(p,[key], 'derived_or_extracted', 'Unscored evidence audit; 2024 WHO deaths reserved')
    for p in (ROOT / 'data/raw').rglob('*'):
        if not p.is_file() or ':Zone.Identifier' in p.name: continue
        name, folder = p.name.lower(),p.parent.name
        if folder=='chi':
            key = 'chi_terms' if 'requirements' in name else ('chi_nphies' if 'nphies' in name else ('chi_class' if 'beneficiary class' in name else 'chi_data'))
        elif folder=='moh': key = 'moh_access' if p.suffix=='.html' else 'moh_population'
        elif folder=='gastat': key = 'gastat_methods' if 'methodology' in name else ('gastat_access' if 'microdata' in name else 'gastat_series')
        elif folder=='nhrsp': key = 'nhrsp_hiss' if 'hiss' in name else ('nhrsp_ncd' if 'ncds' in name else 'nhrsp_access')
        else: raise ValueError('Unmapped provider: '+folder)
        assign(p,[key,'D5'] if key=='moh_population' else [key], 'user_supplied_source',
               'Inspected 1 October 2026; original file retrieval time/URL not supplied')
    for p in (ROOT / 'data/processed/source_inventory_recheck_v1').glob('*.csv'):
        assign(p,['D1'],'derived','GBD source inventory check')
    assign('data/processed/design_v1/country_crosswalk.csv',['D1','D2'],'derived','Study-authored linkage of source identifiers; checked against hierarchy separately')
    # Scope is every retained original input, plus source-derived files and prior registered outputs.
    raw_paths = list((ROOT/'More data').rglob('*')) + list((ROOT/'age_standard').glob('*')) + list((ROOT/'data/raw').rglob('*'))
    for folder in (ROOT/'supporting_data').glob('*/raw'): raw_paths.extend(folder.rglob('*'))
    raw_paths = sorted({p for p in raw_paths if p.is_file() and ':Zone.Identifier' not in p.name})
    missing = [p.relative_to(ROOT).as_posix() for p in raw_paths if p.relative_to(ROOT).as_posix() not in assigned]
    assert not missing, missing
    rows = []
    for path,(keys,kind,detail) in sorted(assigned.items()):
        p = ROOT / path; digest = sha(p); prov = provenance.get(path,{})
        if prov.get('sha256'): assert digest==prov['sha256'], 'Raw source changed: '+path
        if path.startswith('data/raw/'):
            acquisition = ''
        else: acquisition = prov.get('checked_utc',prov.get('retrieved_utc',''))
        direct_url = prov.get('url','')
        rows.append(dict(archive_member=path, reference_ids=';'.join(keys), file_kind=kind,
                         reporting_roles=';'.join(sorted({refs[k]['role'] for k in keys})),
                         exact_download_url=direct_url, citation_urls=';'.join(refs[k]['url'] for k in keys),
                         acquisition_time_if_documented=acquisition, review_date=DATE, bytes=p.stat().st_size,
                         sha256=digest, note=detail))
    csv_write(BASE / 'file_to_citation_crosswalk.csv', rows)
    # Compare new downloads with the preceding inspection's independent hash inventory.
    previous = ROOT/'reports/raw_download_review_2026-10-01/download_inventory.csv'
    old = list(csv.DictReader(previous.open(encoding='utf-8-sig')))
    assert len(old)==19
    for r in old: assert sha(ROOT/r['file'])==r['sha256']
    csv_write(BASE/'raw_source_snapshot.csv',[dict(archive_member=p.relative_to(ROOT).as_posix(),bytes=p.stat().st_size,sha256=sha(p)) for p in raw_paths])
    return {'mapped_files':len(rows),'original_input_files':len(raw_paths),'unmapped_original_files':len(missing),
            'previous_19_download_hashes_unchanged':True,'prior_source_manifest_hashes_checked':len(provenance)}


def missing_report():
    items = [
        ('M01','GBD Results export metadata','The supplied native citation is retained verbatim. Confirm the final publication-year field and original extraction date/query identifiers with IHME or the original exporter.', 'D1','Partial metadata; dataset source and methods are identified'),
        ('M02','Original global ASR workbook provenance','Recover original export query, construction history and extraction date. Regional checks do not independently establish the source history of every global location.', 'D1','Partial provenance'),
        ('M03','GASTAT population product dating','Confirm the publication/update date of the exact retained 2024 workbook/report. Indexed website dates are not silently substituted for the file version.', 'D4','Publication date incomplete'),
        ('M04','GCC-Stat original update date','A dataflow version and acquisition date are known; a durable original publication/update date is not.', 'D3','Publication date incomplete'),
        ('M05','New GASTAT CSV resource','Exact dataset/download URL and retrieval timestamp for the 2010–2024 CSV were not supplied. Values match existing GCC-Stat data but this does not reconstruct the original download URL.', 'gastat_series','Resource metadata incomplete'),
        ('M06','MOH population CSV export','All 68 data cells now match MOH Statistical Yearbook 2023 Table 1-3, credited to GASTAT. The portal export URL and acquisition timestamp remain unknown.', 'moh_population;D5','Parent source resolved; export provenance incomplete'),
        ('M07','CHI quarterly, class and NPHIES resources','Recover exact dataset/resource URLs, publication/update dates and snapshot/version definitions. Citation uses the verified provider collection until individual resource metadata are available.', 'chi_data;chi_class;chi_nphies','Resource metadata incomplete'),
        ('M08','NHRSP documentation versions','Project titles, registered authors/publishers, publication years and DOIs are now verified. Local codebook/questionnaire version dates and original download timestamps remain unknown.', 'nhrsp_ncd;nhrsp_hiss','Dataset citation resolved; document version incomplete'),
        ('M09','Undated institutional documents','Original publication dates remain unspecified for MOH yearbook copies, CHI indication guidance and access guides. Edition/reference years are stated in titles without being presented as publication years.', 'D5;D6;chi_guidance;moh_access;chi_access;nhrsp_access;gastat_access','Non-blocking bibliographic dates'),
        ('M10','IQVIA conference abstract locator','The author, title, 2025 supplement and official abstract page are identified; DOI/abstract number not established. No underlying sales dataset obtained.', 'iqvia_abstract','Non-blocking abstract locator'),
        ('M11','Dataset-specific reuse/deposit terms','Provider attribution is preserved, but do not infer an open redistribution licence from public availability. Exact resource licences for new CHI/NHRSP documentation should accompany any future public data deposit.', 'chi_data;chi_class;chi_nphies;nhrsp_ncd;nhrsp_hiss','Deposit metadata; no external deposit made'),
        ('M12','Supplementary archive identifier','No final journal supplement URL or archive DOI exists yet; add the deposited identifier when available.', 'all','Future submission metadata'),
    ]
    rows=[dict(id=i,item=t,required_detail=x,reference_ids=r,status=s) for i,t,x,r,s in items]
    csv_write(BASE/'missing_reference_metadata.csv',rows)
    d=doc('Missing reference and provenance metadata')
    d.add_paragraph('All retained dataset types have a source citation in the accompanying register. The items below concern incomplete dates, resource locators, versions or future deposit identifiers. '
                    'They are not scientific contradictions. No release-based contradiction analysis is introduced.')
    for row in rows:
        d.add_heading(row['id']+' — '+row['item'],2)
        d.add_paragraph(row['required_detail']);d.add_paragraph('References: '+row['reference_ids']+'. Status: '+row['status']+'.')
    d.save(BASE/'missing_reference_metadata.docx')
    return len(rows)


def validate_and_package(info):
    # Verify every successful acquisition against its preserved HTTP response body.
    downloaded=0
    for f in BASE.glob('download_manifest_*.json'):
        for r in json.loads(f.read_text()):
            if r.get('status')=='downloaded':
                assert sha(BASE/r['file'])==r['sha256'];downloaded+=1
    for r in REFS:
        if r['kind']=='article':
            assert (BASE/('documents/'+r['id']+'.pdf')).read_bytes().startswith(b'%PDF')
            xml=ET.parse(BASE/('documents/'+r['id']+'.xml'))
            doi=xml.findtext('./front/article-meta/article-id[@pub-id-type="doi"]')
            assert doi.lower()==r['doi'].lower()
    assert (BASE/'dataset_references.ris').read_text().count('TY  - ')==len(REFS)
    assert (BASE/'dataset_references.ris').read_text().count('ER  -')==len(REFS)
    assert len(re.findall(r'^@', (BASE/'dataset_references.bib').read_text(),re.M))==len(REFS)
    d=Document(BASE/'dataset_citations_and_reporting.docx')
    text='\n'.join(p.text for p in d.paragraphs)
    assert '/home/' not in text and '/mnt/' not in text
    for r in REFS: assert r['title'] in text
    for f in ['dataset_citations_and_reporting.docx','missing_reference_metadata.docx']:
        with ZipFile(BASE/f) as z: assert z.testzip() is None
    info.update(reference_count=len(REFS), verified_article_pdfs=6, article_xml_doi_matches=6,
                datacite_records=6, successful_new_document_or_metadata_downloads=downloaded,
                ris_and_bibtex_record_counts_match=True, docx_readback_passed=True,
                page_rendering_performed=False, analyses_or_forecasts_changed=False,
                patient_data_acquired=False, raw_inputs_preserved=True, generated_on=DATE)
    (BASE/'validation.json').write_text(json.dumps(info,indent=2),encoding='utf-8')
    (BASE/'README.txt').write_text(
        'DATASET CITATIONS — 1 OCTOBER 2026\n\n'
        'Start with dataset_citations_and_reporting.docx. It contains reporting passages and all 35 references.\n'
        'Import dataset_references.ris into Zotero/EndNote or use dataset_references.bib.\n'
        'source_register.csv distinguishes analytic inputs, source methods, supplementary audits, background documentation and catalog-only records.\n'
        'file_to_citation_crosswalk.csv maps every retained original input and relevant derived file to those references. Paths are supplementary archive member names.\n'
        'missing_reference_metadata.docx/CSV lists unresolved export/version metadata, not missing disease outcomes or scientific contradictions.\n'
        'documents/ holds six newly downloaded open-access paper PDFs and XML, official pages and exact copies of retained reference documents.\n'
        'metadata/ holds PubMed, PMC, Crossref and DataCite evidence, plus the 68-cell MOH population source match.\n'
        'Download manifests and copied_document_manifest.csv distinguish new HTTP retrievals from reuse of existing documents.\n'
        'Validation checks document hashes, DOI identities, citation counts and source coverage; Word page rendering was not performed.\n'
        'No model was refitted and no primary result was altered. No acquired codebook is represented as participant observations.\n'
        'Original data files are indexed, not duplicated into this reference package. Source-specific redistribution terms still apply.\n\n'
        'REBUILD\n'
        'python3 study_design/dataset_citations_2026-10-01/build_citation_package.py\n'
        'Requires Python 3.9+ and python-docx; acquisition additionally requires requests and the configured NCBI skill adapters.\n',encoding='utf-8')
    files=[p for p in BASE.rglob('*') if p.is_file() and '__pycache__' not in p.parts and p.suffix!='.zip' and p.name!='package_checksums.csv']
    csv_write(BASE/'package_checksums.csv',[dict(file=p.relative_to(BASE).as_posix(),bytes=p.stat().st_size,sha256=sha(p)) for p in sorted(files)])
    archive=BASE/'dataset_citation_package.zip'
    with ZipFile(archive,'w',compression=ZIP_DEFLATED,compresslevel=6) as z:
        for p in sorted(files+[BASE/'package_checksums.csv']): z.write(p,'dataset_citations_2026-10-01/'+p.relative_to(BASE).as_posix())
    with ZipFile(archive) as z: assert z.testzip() is None
    print(json.dumps(info))


def main():
    make_catalog()
    assert len({r['id'] for r in REFS}) == len(REFS)
    copy_documents()
    bib_ris()
    build_guide()
    (BASE / 'references.json').write_text(json.dumps(REFS, indent=2, ensure_ascii=False), encoding='utf-8')
    rows = [{k:r.get(k,'') for k in ['id','title','kind','role','year','doi','url','verification','note']} |
            {'citation':cite(r)} for r in REFS]
    csv_write(BASE / 'source_register.csv', rows)
    info=crosswalk()
    info['moh_2022_cells_matched_to_official_yearbook']=match_moh_population()
    info['missing_metadata_items']=missing_report()
    validate_and_package(info)


if __name__ == '__main__':
    main()
