"""Create a standalone publication figure without touching manuscript files."""

from __future__ import annotations

import csv
import hashlib
import json
import os
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", "/tmp/gbd_standalone_forecast_mpl")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from matplotlib.lines import Line2D
from matplotlib.ticker import MultipleLocator
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "forecast.csv"
REFERENCES = ROOT / "study_design/forecast_comparison_2026-10-03.json"
OUT = ROOT / "reports/standalone_forecast_comparison_v1"
STEM = "saudi_forecasts_and_published_context"
INK = "#253646"
MUTED = "#52616D"
GRID = "#E5EAEE"
SCENARIOS = [
    ("gbd_2023_aligned_un_growth", "GBD 2023 baseline + UN growth", "#0072B2", "-", "o"),
    ("un_medium_unaligned", "UN medium population", "#00866B", (0, (1.2, 1.6)), "s"),
    ("national_2024_un_growth", "National 2024 baseline + UN growth", "#C44E00", (0, (5, 2.4)), "^"),
]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def manuscript_hashes():
    return {str(p.relative_to(ROOT)): sha(p)
            for p in sorted((ROOT / "manuscript").rglob("*")) if p.is_file()}


def setup_style():
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 11,
        "text.color": INK, "axes.labelcolor": INK,
        "xtick.color": MUTED, "ytick.color": MUTED,
        "axes.spines.top": False, "axes.spines.right": False,
        "axes.edgecolor": "#A4AFB7", "axes.linewidth": .65,
        "axes.titlesize": 12, "axes.titleweight": "bold",
        "axes.labelsize": 10.5, "xtick.labelsize": 10,
        "ytick.labelsize": 10, "pdf.fonttype": 42,
        "svg.fonttype": "none", "savefig.facecolor": "white",
        "figure.facecolor": "white", "axes.facecolor": "white",
    })


def separate_labels(values, minimum_gap, upper):
    positions = sorted(enumerate(values), key=lambda item: item[1])
    y = np.array([item[1] for item in positions], dtype=float)
    for i in range(1, len(y)):
        y[i] = max(y[i], y[i - 1] + minimum_gap)
    if y[-1] > upper:
        y -= y[-1] - upper
    result = np.empty(len(y))
    for (original_index, _), value in zip(positions, y):
        result[original_index] = value
    return result


def forecast_panels(fig, counts, top, bottom):
    grid = fig.add_gridspec(2, 3, left=.068, right=.965, top=top, bottom=bottom,
                           wspace=.27, hspace=.40)
    for row, outcome in enumerate(["prevalence", "incidence"]):
        ymax = 41 if row == 0 else 5.85
        for col, sex in enumerate(["Male", "Female", "Both"]):
            ax = fig.add_subplot(grid[row, col])
            letter = "ABCDEF"[row * 3 + col]
            sex_label = "Both sexes" if sex == "Both" else sex
            ax.set_title(f"{letter}   {sex_label}", loc="left", pad=11)
            ax.set_axisbelow(True)
            ax.grid(axis="y", color=GRID, linewidth=.7)
            ax.spines["left"].set_visible(False)
            ax.tick_params(axis="both", length=0, pad=6)
            ax.set_xlim(2023.83, 2029.34)
            ax.set_ylim(0, ymax)
            ax.set_xticks(range(2024, 2029))
            ax.yaxis.set_major_locator(MultipleLocator(10 if row == 0 else 1))
            if col == 0:
                ax.set_ylabel("Prevalent cases (thousands)" if row == 0
                              else "Annual incident cases (thousands)", labelpad=9)
            endpoints = []
            for scenario, _, color, linestyle, marker in SCENARIOS:
                data = counts[(counts.outcome == outcome) & (counts.sex == sex)
                              & (counts.population_scenario == scenario)].sort_values("forecast_year")
                assert list(data.forecast_year) == [2024, 2025, 2026, 2027, 2028]
                x = data.forecast_year.to_numpy()
                y = data.forecast_value.to_numpy() / 1000
                ax.plot(x, y, color=color, linestyle=linestyle, linewidth=1.8,
                        marker=marker, markersize=4.8, markerfacecolor="white",
                        markeredgewidth=1.1, clip_on=False)
                endpoints.append((y[-1], color, float(data.forecast_value.iloc[-1])))
            label_ys = separate_labels([e[0] for e in endpoints], ymax * .078, ymax * .92)
            for (point_y, color, count), label_y in zip(endpoints, label_ys):
                ax.plot([2028.06, 2028.34, 2028.41], [point_y, label_y, label_y],
                        color=color, linewidth=.7, alpha=.85)
                ax.text(2028.47, label_y, f"{count:,.0f}", color=color,
                        ha="left", va="center", fontsize=9.3, fontweight="medium")
    handles = [Line2D([], [], color=color, ls=ls, marker=marker, mfc="white",
                      markersize=5, lw=1.8, label=label)
               for _, label, color, ls, marker in SCENARIOS]
    fig.legend(handles=handles, loc="upper left", bbox_to_anchor=(.066, top + .61 / fig.get_figheight()),
               ncol=3, frameon=False, columnspacing=1.85, handlelength=2.65,
               handletextpad=.6, borderaxespad=0, fontsize=10.2)


def context_panels(fig, published, bottom=.128, top=.315, letters=("G", "H")):
    # Only published counts share this axis. Saudi values are in panels A-F.
    ax = fig.add_axes([.297, bottom, .345, top - bottom])
    ax.set_axisbelow(True)
    ax.grid(axis="x", color=GRID, linewidth=.7)
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="both", length=0, pad=6)
    ax.set_xlim(0, 33)
    ax.set_ylim(-.48, 3.66)
    ax.set_xticks([0, 10, 20, 30])
    ax.set_xlabel("Published prevalent cases (millions)", labelpad=9)
    specifications = [
        ("su_2025_bmj", "Global", 2050, "Su et al. (2025)\nGlobal · all ages · 2050"),
        ("su_2025_bmj", "Global", 2030, "Su et al. (2025)\nGlobal · all ages · 2030"),
        ("su_2025_bmj", "North Africa and Middle East", 2050,
         "Su et al. (2025)\nN. Africa & Middle East · all ages · 2050"),
        ("marras_2018_npj", "United States of America", 2030,
         "Marras et al. (2018)\nUnited States · age ≥45 · 2030"),
    ]
    labels = []
    for y, (study, geo, year, label) in zip([3, 2, 1, 0], specifications):
        match = published[(published.study_id == study) & (published.geography == geo)
                          & (published.forecast_year == year)]
        assert len(match) == 1
        r = match.iloc[0]
        val = r.forecast_value / 1e6
        has_bounds = pd.notna(r.lower_bound)
        if has_bounds:
            ax.errorbar(val, y, xerr=np.array([[val - r.lower_bound / 1e6],
                                             [r.upper_bound / 1e6 - val]]),
                        fmt="o", color=INK, markersize=5.4, capsize=3.2,
                        linewidth=1.2, markeredgewidth=1)
        else:
            ax.plot(val, y, marker="D", color=INK, markersize=5.4,
                    markerfacecolor="white", markeredgewidth=1.2)
        number = f"{val:.1f}" if year == 2030 and geo == "Global" else f"{val:.3f}"
        ax.text(val, y + .24, number, ha="center", va="bottom", fontsize=10.5)
        labels.append(label)
    ax.set_yticks([3, 2, 1, 0], labels, fontsize=9.4)
    for tick in ax.get_yticklabels():
        tick.set_linespacing(1.5)
    fig.text(.068, top + .04, f"{letters[0]}   Published prevalence projections", fontsize=12, weight="bold")

    ax = fig.add_axes([.785, bottom + .015, .174, top - bottom - .055])
    ax.set_axisbelow(True)
    ax.grid(axis="x", color=GRID, linewidth=.7)
    ax.spines["left"].set_visible(False)
    ax.tick_params(axis="both", length=0, pad=6)
    ax.set_xlim(0, 165)
    ax.set_ylim(-.52, 1.56)
    ax.set_xticks([0, 50, 100, 150])
    for y, outcome in [(1, "prevalence"), (0, "incidence")]:
        match = published[(published.study_id == "wu_2025_frontiers") & (published.outcome == outcome)]
        assert len(match) == 1
        value = float(match.forecast_value.iloc[0])
        ax.plot(value, y, "D", ms=5.5, mfc="white", mec=INK, mew=1.2)
        ax.text(value, y + .23, f"{value:.2f}", va="bottom", ha="center", fontsize=10.5)
    ax.set_yticks([1, 0], ["Prevalence", "Incidence"], fontsize=9.4)
    ax.set_xlabel("Age-standardised rate\n(per 100,000 population)", labelpad=9, fontsize=10)
    fig.text(.71, top + .04, f"{letters[1]}   Published rate projections", fontsize=12, weight="bold")
    fig.text(.71, top + .013, "Wu et al. (2025) · Global · 2026", fontsize=10, color=MUTED)


def export_figure(fig, stem):
    paths = []
    for ext in ("pdf", "svg", "png", "tiff"):
        p = OUT / f"{stem}.{ext}"
        kwargs = {"dpi": 600 if ext == "tiff" else 300}
        if ext == "tiff":
            kwargs["pil_kwargs"] = {"compression": "tiff_lzw"}
        fig.savefig(p, **kwargs)
        paths.append(p)
    # Lightweight inspection preview; not the publication raster.
    preview = OUT / f"{stem}_preview.png"
    fig.savefig(preview, dpi=125)
    paths.append(preview)
    return paths


def write_caption(references):
    paragraphs = [
        ("Standalone figure. Saudi Parkinson’s disease forecasts through 2028 and published projection context.", True),
        ("Panels A–F show annual adapted temporal convolutional network (TCN) point forecasts for Saudi residents aged 45 years and older, including citizens and non-citizens. Male, female and combined-sex results are shown for prevalence (A–C) and annual incidence (D–F). Disease-rate forecasts use data through 2023 and are unchanged across the three population scenarios: GBD 2023 population baselines advanced with United Nations World Population Prospects 2024 growth; unaligned UN medium populations; and national 2024 population baselines advanced with UN growth. The national open-age 80+ total is allocated using the GBD-aligned within-80+ age distribution. Axes begin at zero and share limits within each outcome row. Endpoint labels give 2028 counts, rounded to whole cases solely for display. All stored precision is retained in forecast.csv. Close overlap of the UN and national prevalence curves reflects similar totals under those scenarios.", False),
        ("Panels G–H provide published numerical context. G shows prevalent counts for the world in 2030 and 2050 and North Africa and the Middle East in 2050 from Su et al., and for United States residents aged 45 years and older in 2030 from Marras et al. Su et al. used Bayesian model averaging; Marras et al. applied stable age- and sex-specific prevalence to projected populations. H shows Wu et al.’s global 2026 prevalence and incidence age-standardised rates, forecast with ARIMA. All published values shown combine both sexes. Published counts and standardised rates use separate axes. Values retain the precision of the source tables or text; million/thousand units are converted arithmetically. The longer published projection years are retained rather than interpolated to 2028.", False),
        ("Horizontal whiskers in G are Su et al.’s published 95% uncertainty intervals. Open diamonds identify published point values for which numeric forecast bounds were not extracted for this export; they do not indicate zero uncertainty. No uncertainty ribbons are plotted for the Saudi scenarios: differences between population assumptions are not prediction intervals. The available published values do not provide a verified Saudi 45+ comparator matched on year, sex and estimand, so the figure supports contextual comparison only. It does not establish agreement, superiority, external validation or calibrated future intervals. The retained study’s historical interval-reliability limitations remain unchanged.", False),
        ("This figure is an independent reporting artifact. It has not been added to, numbered within, or referenced from the manuscript. The older national-2022 population scenario and alternative within-80+ allocations remain available in forecast.csv; the figure displays the original reference, UN medium and latest national anchor for clarity.", False),
        ("References", True),
    ]
    for source in references["sources"].values():
        paragraphs.append((source["citation"] + " https://doi.org/" + source["doi"], False))
    paragraphs.append(("Current-study sources: GBD 2023; UN World Population Prospects 2024; Saudi General Authority for Statistics population estimates. Disease-rate forecasts and demographic scenarios were extracted from the preserved current-study export, forecast.csv. Source review date: 3 October 2026.", False))
    text = "\n\n".join(p for p, _ in paragraphs) + "\n"
    (OUT / "figure_caption.txt").write_text(text, encoding="utf-8")
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    for text, bold in paragraphs:
        p = doc.add_paragraph()
        p.add_run(text).bold = bold
    doc.save(OUT / "figure_caption.docx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    protected_before = manuscript_hashes()
    source_before = sha(SOURCE)
    d = pd.read_csv(SOURCE, float_precision="round_trip")
    count_mask = (
        d.record_type.eq("current_study_count_forecast")
        & d.population_scenario.isin([s[0] for s in SCENARIOS])
        & d.within_80plus_allocation.isin(["native_age_detail", "gbd_aligned_within_80plus"])
    )
    counts = d.loc[count_mask].copy()
    published = d.loc[d.record_type.eq("published_context_forecast")].copy()
    assert len(counts) == 90 and len(published) == 6
    assert counts.age_scope.eq("45+").all()
    assert counts.geography.eq("Saudi Arabia").all()
    assert published.direct_difference.isna().all()
    # Verify all plotted forecasts include complete annual, sex and outcome cells.
    for _, group in counts.groupby(["population_scenario", "outcome", "sex"]):
        assert sorted(group.forecast_year) == [2024, 2025, 2026, 2027, 2028]
    setup_style()
    fig = plt.figure(figsize=(12.6, 12.2))
    fig.text(.068, .974, "Parkinson’s disease forecasts for Saudi Arabia", fontsize=20, weight="bold", va="top")
    fig.text(.068, .943, "Ages ≥45 years  |  Annual projections, 2024–2028  |  Adapted TCN", fontsize=11.5, color=MUTED)
    forecast_panels(fig, counts, top=.857, bottom=.443)
    fig.text(.068, .405, "Published projections in their original populations", fontsize=14, weight="bold")
    fig.text(.068, .383, "Different populations, ages and forecast years; contextual comparison only.", fontsize=10.7, color=MUTED)
    context_panels(fig, published, bottom=.119, top=.312)
    fig.text(.068, .060, "A–F: population scenarios, not uncertainty intervals. G: whiskers = published 95% uncertainty intervals.", fontsize=9.3, color=MUTED)
    fig.text(.068, .043, "G–H: open diamonds = point values; numeric forecast bounds not extracted. Published results combine both sexes.", fontsize=9.3, color=MUTED)
    fig.text(.068, .023, "Sources: current study; Su et al., BMJ 2025; Wu et al., Front Aging Neurosci 2025; Marras et al., NPJ Parkinsons Dis 2018.", fontsize=8.7, color=MUTED)
    outputs = export_figure(fig, STEM)
    plt.close(fig)

    # Larger Saudi-only companion for independent use or close inspection.
    fig = plt.figure(figsize=(12.6, 8.0))
    fig.text(.068, .962, "Saudi Parkinson’s disease forecasts, 2024–2028", fontsize=20, weight="bold", va="top")
    fig.text(.068, .913, "Ages ≥45 years  |  Same disease-rate forecasts under three population scenarios", fontsize=11.2, color=MUTED)
    forecast_panels(fig, counts, top=.791, bottom=.118)
    fig.text(.068, .053, "Point projections, conditional on population assumptions. Endpoint labels show 2028 counts; no prediction intervals are shown.", fontsize=9.5, color=MUTED)
    fig.text(.068, .026, "National scenario uses GBD-aligned age weights within the 80+ group. UN and national prevalence trajectories nearly overlap.", fontsize=9.5, color=MUTED)
    outputs.extend(export_figure(fig, "saudi_forecasts_detail"))
    plt.close(fig)
    write_caption(json.loads(REFERENCES.read_text(encoding="utf-8")))

    # Preserve the literal input numeric strings in the figure-data subset.
    with SOURCE.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames
        raw = list(reader)
    included = [r for r in raw if r["record_type"] == "published_context_forecast" or (
        r["record_type"] == "current_study_count_forecast"
        and r["population_scenario"] in [s[0] for s in SCENARIOS]
        and r["within_80plus_allocation"] in ["native_age_detail", "gbd_aligned_within_80plus"])]
    assert len(included) == 96
    with (OUT / "plotted_data.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(included)
    protected_after = manuscript_hashes()
    assert protected_before == protected_after
    assert sha(SOURCE) == source_before
    raster_checks = {}
    for p in outputs:
        if p.suffix in [".png", ".tiff"]:
            with Image.open(p) as im:
                raster_checks[p.name] = {"pixels": list(im.size), "dpi": [float(x) for x in im.info.get("dpi", [])]}
                if p.suffix == ".tiff":
                    assert min(im.info["dpi"]) >= 599
    report = {
        "purpose": "Standalone publication graphics; not included in manuscript",
        "source_sha256": source_before, "plotted_rows": len(included),
        "current_study_rows": len(counts), "published_context_rows": len(published),
        "manuscript_files_unchanged": len(protected_before),
        "manuscript_hashes": protected_before,
        "source_file_unchanged": True, "raw_precision_preserved_in_plotted_data": True,
        "no_cross_study_numeric_differences": True, "no_scenario_uncertainty_bands": True,
        "vector_exports": [p.name for p in outputs if p.suffix in [".pdf", ".svg"]],
        "raster_checks": raster_checks,
        "outputs": {p.name: {"sha256": sha(p), "bytes": p.stat().st_size} for p in outputs},
    }
    (OUT / "validation.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k not in ["manuscript_hashes", "outputs"]}, indent=2))


if __name__ == "__main__":
    main()
