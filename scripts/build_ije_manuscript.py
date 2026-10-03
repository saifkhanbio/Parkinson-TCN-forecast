"""Build an additive IJE manuscript package from frozen study ledgers."""
import os
os.environ.setdefault("MPLCONFIGDIR", "/tmp/gbd_ije_matplotlib")
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import zipfile

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import TwoSlopeNorm
from matplotlib.lines import Line2D
from PIL import Image, ImageOps, ImageDraw
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuscript/ije_draft_v1"
FIG = OUT / "figures"
TITLE = "Sex-specific transfer learning and forecast reliability for Parkinson’s disease in Saudi Arabia"
COL = {"TCN": "#222222", "Local": "#0072B2", "Non-neural": "#D55E00",
       "Mixture": "#009E73", "Male": "#0072B2", "Female": "#D55E00"}
AGES = ["45-49", "50-54", "55-59", "60-64", "65-69", "70-74", "75-79", "80-84", "85-89", "90-94", "95+"]
COUNTRIES = ["Saudi Arabia", "Bahrain", "Kuwait", "Oman", "Qatar", "United Arab Emirates"]
SOURCES = {}
REFERENCES = json.loads((OUT / "references.json").read_text())
BODY = (OUT / "body.md").read_text()
ABSTRACT = (OUT / "abstract.md").read_text()
FRONT_MATTER = json.loads((OUT / "front_matter.json").read_text())
CITE_RE = re.compile(r"\[\s*(@[a-z_]+(?:;@[a-z_]+)*)\]")
KEYS = []
for match in CITE_RE.finditer(BODY):
    for key in match.group(1).replace("@", "").split(";"):
        if key not in KEYS:
            KEYS.append(key)
REFNUM = {key: i + 1 for i, key in enumerate(KEYS)}
assert set(KEYS) == set(REFERENCES)


def word_count(text, include_headings=False):
    text = CITE_RE.sub("", text)
    text = re.sub(r"^#+\s*" if include_headings else r"^#+[^\n]+", "", text, flags=re.M)
    return len(text.split())


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1048576), b""):
            h.update(block)
    return h.hexdigest()


def read(name):
    SOURCES[name] = sha(ROOT / name)
    return pd.read_csv(ROOT / name, float_precision="round_trip")


def n(value, decimals=0):
    text = f"{float(value):,.{decimals}f}"
    if abs(float(value)) < 10000:
        return text.replace(",", "")
    return text.replace(",", " ")


def rich(paragraph, text):
    position = 0
    for match in CITE_RE.finditer(text):
        paragraph.add_run(text[position:match.start()])
        keys = match.group(1).replace("@", "").split(";")
        run = paragraph.add_run(",".join(str(REFNUM[key]) for key in keys))
        run.font.superscript = True
        position = match.end()
    paragraph.add_run(text[position:])


def document(landscape=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2
    normal.paragraph_format.space_after = Pt(0)
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        doc.styles[name].font.name = "Times New Roman"
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[name].paragraph_format.keep_with_next = True
    sec.header.paragraphs[0].text = "Parkinson’s disease forecasting | Draft"
    sec.header.paragraphs[0].style = "Caption"
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = TITLE
    doc.core_properties.author = "Research team — authorship details pending"
    return doc


def add_markdown(doc, text):
    for block in text.strip().split("\n\n"):
        if block.startswith("# "):
            doc.add_heading(block[2:], level=1)
        elif block.startswith("## "):
            doc.add_heading(block[3:], level=2)
        elif block.startswith("### "):
            doc.add_heading(block[4:], level=3)
        else:
            rich(doc.add_paragraph(), block.replace("\n", " "))


def add_table(doc, title, headers, rows, footnote="", sizes=None):
    doc.add_heading(title, level=1)
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    if sizes:
        available = (doc.sections[-1].page_width - doc.sections[-1].left_margin
                     - doc.sections[-1].right_margin) / Cm(1)
        sizes = [size * min(1, available / sum(sizes)) for size in sizes]
        for column, size in zip(table.columns, sizes):
            column.width = Cm(size)
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "bottom", "left", "right", "insideH", "insideV"]:
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single" if edge in ["top", "bottom"] else "nil")
        element.set(qn("w:sz"), "8")
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for cell, title_cell in zip(table.rows[0].cells, headers):
        cell.text = title_cell
        for run in cell.paragraphs[0].runs:
            run.bold = True
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            rich(cell.paragraphs[0], str(value))
    for row in table.rows:
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for i, cell in enumerate(row.cells):
            if sizes:
                cell.width = Cm(sizes[i])
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(5)
                for run in p.runs:
                    run.font.size = Pt(10)
    if footnote:
        p = doc.add_paragraph()
        rich(p, footnote)
        p.paragraph_format.line_spacing = 1.15
        for run in p.runs:
            run.font.size = Pt(10)
    return table


def panel(ax, letter, title):
    ax.set_title(f"{letter}  {title}", loc="left", fontsize=15, pad=13, weight="bold")
    ax.spines[["top", "right"]].set_visible(False)


def save_figure(fig, number):
    prefix = FIG / f"figure_{number}"
    fig.savefig(prefix.with_suffix(".png"), dpi=300, facecolor="white")
    fig.savefig(prefix.with_suffix(".tiff"), dpi=600, facecolor="white", pil_kwargs={"compression": "tiff_lzw"})
    fig.savefig(prefix.with_suffix(".eps"), format="eps", facecolor="white")
    plt.close(fig)


def figures(age, endpoint, donors, learning, rates, burdens, population, shares):
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 14, "axes.labelsize": 15,
        "xtick.labelsize": 12, "ytick.labelsize": 13, "legend.fontsize": 12,
        "axes.linewidth": .8, "lines.linewidth": 2.2, "savefig.dpi": 300, "ps.fonttype": 42,
        "pdf.fonttype": 42})
    fig, axes = plt.subplots(2, 2, figsize=(12, 10), gridspec_kw={"height_ratios": [1, 1.12]})
    fig.subplots_adjust(left=.16, right=.97, top=.89, bottom=.20, hspace=.72, wspace=.42)
    for j, sex in enumerate(["Male", "Female"]):
        ax = axes[0, j]
        d = age[age.sex.eq(sex)]
        base = d[d.comparator.eq("local_champion")].set_index("age").loc[AGES]
        ax.axvspan(6.5, 10.5, color="#eeeeee", zorder=0)
        ax.plot(range(11), base.tcn_absolute_log_error, "o-", color=COL["TCN"], label="Adapted TCN")
        for role, label, marker in [("local_champion", "Local", "s"), ("nonneural_champion", "Non-neural", "^")]:
            part = d[d.comparator.eq(role)].set_index("age").loc[AGES]
            ax.plot(range(11), part.comparator_absolute_log_error, marker + "--", color=COL[label], label=label + " comparator")
        ax.set_xticks(range(11), [x.replace("-", "–") for x in AGES], rotation=55, ha="right")
        ax.set_ylabel("Absolute log error")
        ax.set_xlabel("Age group (years)")
        ax.set_xlim(-.3, 10.4)
        ax.set_ylim(0, .095)
        panel(ax, "AB"[j], f"Saudi prevalence: {sex.lower()}")
    handles, labels = axes[0, 0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", ncol=3, frameon=False, bbox_to_anchor=(.56, .99))
    for j, outcome in enumerate(["prevalence", "incidence"]):
        ax = axes[1, j]
        order = [(s, c) for s in ["Male", "Female"] for c in ["local_champion", "nonneural_champion"]]
        array = np.array([[endpoint.loc[(endpoint.target == country) & (endpoint.outcome == outcome)
                    & (endpoint.sex == sex) & (endpoint.comparator == comp), "relative_improvement_percent"].item()
                    for sex, comp in order] for country in COUNTRIES])
        im = ax.imshow(array, cmap="RdBu", norm=TwoSlopeNorm(vmin=-70, vcenter=0, vmax=70), aspect="auto")
        for row in range(6):
            for col in range(4):
                value = array[row, col]
                ax.text(col, row, f"{value:+.1f}", ha="center", va="center", color="white" if abs(value) > 45 else "#111111", fontsize=13)
        ax.set_xticks(range(4), ["Male\nlocal", "Male\nnon-neural", "Female\nlocal", "Female\nnon-neural"])
        ax.set_yticks(range(6), ["Saudi Arabia", "Bahrain", "Kuwait", "Oman", "Qatar", "UAE"])
        panel(ax, "CD"[j], f"GCC {outcome}")
    cax = fig.add_axes([.28, .082, .55, .018])
    fig.colorbar(im, cax=cax, orientation="horizontal", label="TCN relative error improvement (%) | positive = lower error")
    save_figure(fig, 1)

    fig, axes = plt.subplots(3, 2, figsize=(12, 12))
    fig.subplots_adjust(left=.17, right=.96, bottom=.07, top=.95, hspace=.57, wspace=.35)
    names = ["GCC only", "Historically similar", "Random 101", "Random 211", "Random 307", "Random 401", "Random 503"]
    d = donors[donors.arm.ne("all_six")]
    for j, sex in enumerate(["Male", "Female"]):
        ax = axes[0, j]
        ax.axvline(0, color="black", linewidth=1)
        ax.barh(np.arange(7), d[sex], color=COL[sex], height=.65)
        for i, value in enumerate(d[sex]):
            ax.text(value + (1.1 if value >= 0 else -1.1), i, f"{value:+.1f}", ha="left" if value >= 0 else "right", va="center", fontsize=11)
        ax.set_yticks(range(7), names)
        ax.invert_yaxis()
        ax.set_xlim(-23, 59)
        ax.set_xlabel("Error change vs all donors (%)\nPositive = higher error")
        panel(ax, "AB"[j], f"Donor restriction: {sex.lower()}")
    for row, outcome in enumerate(["prevalence", "incidence"], 1):
        for col, sex in enumerate(["Male", "Female"]):
            ax = axes[row, col]
            part = learning[(learning.outcome == outcome) & (learning.sex == sex) & (learning.horizon == 5) & (learning.age_group == "45+")]
            for family, label, style, color in [("tcn_adapted", "Adapted", "o-", COL[sex]),
                    ("tcn_unadapted", "Unadapted", "s--", "#555555")]:
                sub = part[part.family == family].sort_values("history_years")
                assert len(sub) == 3
                ax.plot(sub.history_years, sub.mean_absolute_log_error, style, label=label, color=color)
            ax.set_xticks([15, 20, 29])
            ax.set_xlim(14, 30)
            ax.set_ylabel("Mean absolute log error")
            ax.set_xlabel("Saudi history (years)")
            ax.ticklabel_format(axis="y", style="plain", useOffset=False)
            ax.legend(frameon=False, loc="best", fontsize=11)
            panel(ax, "CDEF"[(row - 1) * 2 + col], f"{outcome.capitalize()}: {sex.lower()}")
    save_figure(fig, 2)

    fig, axes = plt.subplots(3, 2, figsize=(12, 11.5))
    fig.subplots_adjust(left=.16, right=.96, top=.91, bottom=.07, hspace=.65, wspace=.4)
    order = [(sex, age_scope) for sex in ["Male", "Female"] for age_scope in ["45+", "80+"]]
    for col, outcome in enumerate(["prevalence", "incidence"]):
        part = rates[(rates.target == "Saudi Arabia") & (rates.outcome == outcome) & (rates.horizon == 5) & (rates.scale == "rate")]
        ax = axes[0, col]
        ax.axvline(80, linestyle=":", color="#555555", linewidth=1.7)
        for k, (sex, scope) in enumerate(order):
            a = part[(part.sex == sex) & (part.age_scope == scope) & (part.family == "tcn_adapted__cdf")].iloc[0]
            b = part[(part.sex == sex) & (part.age_scope == scope) & (part.family == "mixture__equal_weight")].iloc[0]
            ax.plot([100*a.coverage, 100*b.coverage], [k-.07, k+.07], color="#bbbbbb", zorder=1)
            ax.scatter(100*a.coverage, k-.07, c=COL["TCN"], marker="o", s=80, zorder=2)
            ax.scatter(100*b.coverage, k+.07, c=COL["Mixture"], marker="D", s=80, zorder=3)
            value = 100*(b.wis/a.wis-1)
            axes[1, col].barh(k, value, color="#0072B2" if value < 0 else "#D55E00", height=.55)
            axes[1, col].text(value + (-.9 if value < 0 else .9), k, f"{value:+.1f}", va="center", ha="right" if value < 0 else "left", fontsize=12)
        for axis in [ax, axes[1, col]]:
            axis.set_yticks(range(4), [f"{s} {a}" for s, a in order])
            axis.invert_yaxis()
        ax.set_xlim(0, 100)
        ax.set_xlabel("Rate interval coverage (%)")
        panel(ax, "AB"[col], outcome.capitalize())
        axes[1, col].axvline(0, color="black", linewidth=1)
        axes[1, col].set_xlim(-27, 42)
        axes[1, col].set_xlabel("WIS change (%) | negative = better")
        panel(axes[1, col], "CD"[col], "Mixture vs matched TCN")
        ax = axes[2, col]
        part = burdens[(burdens.target == "Saudi Arabia") & (burdens.outcome == outcome)
                       & (burdens.horizon == 5) & (burdens.population_method == "log_trend_last8")]
        for k, sex in enumerate(["Male", "Female"]):
            for offset, family, color, marker in [(-.12, "tcn_adapted__cdf", COL["TCN"], "o"),
                                                  (.12, "mixture__equal_weight", COL["Mixture"], "D")]:
                record = part[(part.node == sex + "__80+_within_45+") & (part.family == family)].iloc[0]
                ax.scatter(record.coverage*100, k+offset, s=85, color=color, marker=marker)
                ax.text(record.coverage*100+4, k+offset, f"{round(record.coverage*5)}/5", va="center", fontsize=12)
        ax.set_yticks([0, 1], ["Male", "Female"])
        ax.set_ylim(1.45, -.5)
        ax.set_xlim(0, 100)
        ax.axvline(80, linestyle=":", color="#555555", linewidth=1.7)
        ax.set_xlabel("80+ share interval coverage (%)")
        panel(ax, "EF"[col], "Oldest-age composition")
    fig.legend([Line2D([], [], marker="o", color=COL["TCN"], linestyle="none", markersize=9),
                Line2D([], [], marker="D", color=COL["Mixture"], linestyle="none", markersize=9)],
               ["Matched TCN control", "Equal-weight mixture"], loc="upper center", ncol=2, frameon=False)
    save_figure(fig, 3)

    fig, axes = plt.subplots(2, 2, figsize=(12, 8.8))
    fig.subplots_adjust(left=.12, right=.96, top=.93, bottom=.17, hspace=.62, wspace=.38)
    for col, sex in enumerate(["Male", "Female"]):
        ax = axes[0, col]
        record = population[(population.outcome == "prevalence") & (population.sex == sex) & (population.age_group == "80+")].iloc[0]
        values = [record.gbd_population/1000, record.un_population/1000, record.gastat_population/1000]
        ax.bar(range(3), values, color=["#666666", "#56B4E9", "#009E73"], width=.62)
        for k, value in enumerate(values):
            ax.text(k, value+1.5, f"{value:.1f}", ha="center", fontsize=13)
        ax.set_xticks(range(3), ["GBD-implied", "UN", "GASTAT"])
        ax.set_ylim(0, 100)
        ax.set_ylabel("Population aged 80+ (thousands)")
        panel(ax, "AB"[col], f"Saudi population, 2023: {sex.lower()}")
    styles = [("log_trend_last8", "Operational forecast", "#D55E00", "o"),
              ("gbd_realized_oracle", "Realised-population oracle", "#009E73", "s")]
    for col, outcome in enumerate(["prevalence", "incidence"]):
        ax = axes[1, col]
        for row, sex in enumerate(["Male", "Female"]):
            part = shares[(shares.outcome == outcome) & (shares.sex == sex)]
            values = [part.loc[part.scenario.eq(scenario), "value"].item() for scenario, _, _, _ in styles]
            truth = part.observed.iloc[0]
            for offset, ((scenario, label, color, marker), value) in zip([-.15, 0], zip(styles, values)):
                ax.scatter(value, row+offset, color=color, marker=marker, s=90, label=label if row == 0 else None)
                ax.text(value+.5, row+offset, f"{value:.2f}", va="center", fontsize=11)
            ax.scatter(truth, row+.15, color="#222222", marker="D", s=90, label="GBD verification" if row == 0 else None)
            ax.text(truth+.5, row+.15, f"{truth:.2f}", va="center", fontsize=11)
        ax.set_yticks([0, 1], ["Male", "Female"])
        ax.set_ylim(1.45, -.5)
        ax.set_xlim(8, 25)
        ax.set_xlabel("80+ share of sex-specific 45+ burden (%)")
        panel(ax, "CD"[col], f"2018→2023 {outcome}")
    handles, labels = axes[1, 0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="lower center", ncol=3, frameon=False, bbox_to_anchor=(.51, .03), fontsize=11)
    save_figure(fig, 4)


CAPTIONS = [
    ("Primary age-specific performance and Gulf benchmarking",
     "Panels A–B show Saudi prevalence absolute log errors at the 2018-origin, five-year endpoint for the adapted temporal convolutional network (TCN) and historically selected local and non-neural comparators. Shading identifies ages ≥80 years; the open ≥95-year band is retained. Panels C–D show relative TCN error improvements against each comparator, by sex and country, for prevalence and incidence; positive values indicate lower TCN error. These are descriptive contrasts without inferential confidence intervals. GCC, Gulf Cooperation Council; UAE, United Arab Emirates. All countries use Global Burden of Disease 2023 outcomes. Supplementary Tables S1–S2 contain complete model results.",
     "Two age-error line plots show mixed comparator rankings in Saudi males and females. Two heatmaps show heterogeneous prevalence and incidence improvements across six Gulf countries, with a negative Saudi female prevalence contrast against the non-neural comparator."),
    ("Donor restriction and Saudi target-history sensitivity",
     "Panels A–B show percentage changes in Saudi prevalence five-year mean absolute log error relative to the matched all-six-donor GPU reference; positive changes indicate harm. Random donor sets are identified by their fixed seeds. Panels C–F compare adapted and unadapted temporal convolutional networks for 15, 20 and 29 years of Saudi history, with donor checkpoints, settings, device and origin held fixed. These budgets provide 3, 8 and 17 complete target adaptation windows per age–sex stratum. Fixed-default learning-curve procedures are distinct from the tuned primary procedure. GCC, Gulf Cooperation Council; GPU, graphics processing unit. Supplementary Table S3 contains complete donor, age and model-family contrasts.",
     "Donor-restriction bars show lower male error with GCC-only donors but higher female error under every restriction. Four learning curves show that additional Saudi history does not consistently lower error; adaptation helps some sex–outcome combinations and harms others."),
    ("Interval scores, coverage and oldest-age composition in the exploratory mixture",
     "The two columns represent Saudi prevalence and incidence. Panels A–B compare nominal 80% rate interval coverage for matched temporal convolutional network (TCN) and equal-weight mixture procedures at ages ≥45 and ≥80 years. Panels C–D show mixture changes in weighted interval score (WIS; 50%/80% intervals), where negative changes are favourable. Panels E–F show coverage of the sex-specific ≥80-year share within ≥45-year burden; labels give the number of covered origins out of five. The population procedure is the preceding eight-year log trend. Both procedures use inverse empirical cumulative-distribution quantiles. Original interpolated-quantile results are separate controls, not a source of claimed mixture benefit. There are 55 age–origin cells for ages ≥45, 20 for ≥80 and five dependent observations for each share; dotted lines indicate nominal coverage. These summaries average horizon-five results across overlapping origins 2014–2018 and are exploratory. Supplementary Tables S5–S6 retain all alternatives, widths, misses and decision criteria.",
     "Coverage plots and score-change bars show improved prevalence interval scores but worse female coverage. Male incidence coverage improves while its score deteriorates. Oldest-age male share coverage stays at one of five origins for both outcomes, and female prevalence-share coverage falls from three of five to two of five."),
    ("Population-source differences and oldest-age burden composition",
     "Panels A–B compare Saudi populations aged ≥80 years in 2023 from Global Burden of Disease (GBD)-implied denominators, United Nations (UN) World Population Prospects 2024 and General Authority for Statistics (GASTAT) estimates. The GBD denominator is inferred from native Number/Rate, not an independent population export. Source residency, reference-date and vintage comparability remain unresolved. Panels C–D show the ≥80-year share of sex-specific burden within ages ≥45 years at the 2018→2023 endpoint. Operational forecasts use preceding eight-year population log trends; oracle values substitute realised GBD-implied populations while holding predicted disease rates fixed. Oracle substitutions are diagnostics, not operational improvements. All displayed values are point quantities, without error bars or probability claims. Data sources are cited in Table 1; complete scenarios appear in Supplementary Table S7.",
     "Population bars show substantially smaller GBD-implied oldest-age populations than UN and GASTAT estimates. Burden-share plots show that realised-population substitutions bring male forecasts closer to GBD verification values, while female changes are smaller and do not uniformly improve agreement.")
]


def main():
    FIG.mkdir(parents=True, exist_ok=True)
    # Preserve the scientific evidence; this script performs no model fitting.
    prior = json.loads((ROOT / "reports/study_synthesis_v1_2/validation.json").read_text())
    for name, expected in prior["source_sha256"].items():
        assert sha(ROOT / name) == expected, name
    for name, expected in prior["artifact_sha256"].items():
        assert sha(ROOT / "reports/study_synthesis_v1_2" / name) == expected, name
    age = read("results/primary_v1/primary_age_contrasts.csv")
    primary = read("results/primary_v1/primary_contrasts.csv")
    endpoint = read("reports/secondary_v1/endpoint_contrasts.csv")
    donors = read("reports/donor_comparisons_gpu_v1/endpoint_percent_changes_vs_gpu_reference.csv")
    learning = read("results/learning_curves_v1/summary.csv")
    rates = read("reports/distribution_mixture_v1/rate_comparison.csv")
    burdens = read("reports/distribution_mixture_v1/burden_comparison.csv")
    population = read("reports/population_sensitivity_v1/saudi_population_source_comparison.csv")
    shares = read("reports/population_sensitivity_v1/saudi_80plus_share_endpoint.csv")
    projection = read("reports/study_synthesis_v1_2/saudi_2028_scenario_totals.csv")
    figures(age, endpoint, donors, learning, rates, burdens, population, shares)
    tables(endpoint, rates, burdens, projection)
    manuscript()
    contradictions()
    missing_references()
    supplementary()
    review_package()
    print(json.dumps({"output": str(OUT), "status": "draft_built", "figures": 4, "tables": 4}, indent=2))


def tables(endpoint, rates, burdens, projection):
    doc = document(landscape=True)
    data_rows = [
        ["Regional disease panel", "GBD 2023; six GCC countries and Jordan; males/females, eleven ages ≥45", "1990–2023; deaths/YLLs also 1980–1989", "Age-specific rates per 100 000 and numbers, with source bounds", "Primary prevalence; secondary incidence; supporting mortality/disability", "[@gbd_results;@gbd_nonfatal;@gbd_fatal]"],
        ["Full-age standardised rates", "203 World Bank-classified locations; 7 independently checked", "1990–2023; regional fatal verification also 1980–1989", "Full-age age-standardised rate per 100 000", "Separate donor benchmark; 4284 overlapping points verified", "[@gbd_results;@hierarchy]"],
        ["UN population estimates/projections", "Age and sex; six GCC countries plus Jordan in regional preparation", "Historical estimates through 2023; medium projections used for 2024–2028", "Persons, converted from source thousands", "Demographic preparation and fixed-population scenarios", "[@wpp;@wpp_methods]"],
        ["UN probabilistic population", "Six GCC countries, both sexes, ages 45–49 through 100+", "2024–2028; 3600 marginal quantile values", "Marginal 80%/95% limits and medians", "Uncertainty-availability audit; not inserted into forecast intervals", "[@wpp;@wpp_methods]"],
        ["Saudi national population", "Saudi Arabia; age, sex and nationality; oldest workbook group 80+", "2023/2024 estimates, 2024 edition", "Population counts", "Denominator/source comparison; 80+ group preserved", "[@gastat]"],
        ["GCC-Stat population", "Six GCC countries; age, sex and nationality", "Available 2010–2024 records", "Population counts", "Supporting demographic consistency checks; overlaps GASTAT", "[@gccstat]"],
    ]
    add_table(doc, "Table 1. Data sources, estimands and analytical roles",
              ["Source", "Population", "Period", "Quantity", "Role", "Reference"], data_rows,
              "GBD, Global Burden of Disease; GCC, Gulf Cooperation Council; UN, United Nations; YLLs, years of life lost. Source uncertainty bounds are not forecast intervals. Full-age rates and all-age counts are separate from the primary ages-45+ estimand. The 203 locations comprise 201 main GBD national locations plus Hong Kong and Macao; Cook Islands, Niue and Tokelau are absent. See Supplementary Material S1 for the citation and provenance crosswalk.",
              [3, 5, 4, 4, 6, 2])
    doc.add_page_break()
    names = {"damped_ets": "Damped exponential smoothing", "pooled_boosting": "Pooled boosting",
             "arima": "Autoregressive integrated moving average", "donor_ridge_adapted": "Adapted donor ridge",
             "donor_boosting_adapted": "Adapted donor boosting", "pooled_ridge": "Pooled ridge"}
    subset = endpoint[endpoint.target.eq("Saudi Arabia")]
    rows = [[r.outcome.capitalize(), r.sex, names[r.comparator_source_family], n(r.tcn_error, 6),
             n(r.comparator_error, 6), f"{r.relative_improvement_percent:+.2f}"] for r in subset.itertuples()]
    add_table(doc, "Table 2. Saudi five-year primary prevalence and secondary incidence comparisons",
              ["Outcome", "Sex", "Historical comparator", "TCN error", "Comparator error", "TCN improvement (%)"], rows,
              "TCN, temporal convolutional network. Errors are equal-age mean absolute log errors across eleven age groups, 45–49 through ≥95 years, for the 2018→2023 endpoint. Positive improvements mean lower TCN error. Comparator families were selected from eligible historical data. The joint primary prevalence rule requires all four prevalence contrasts to favour TCN and was not met. These point contrasts have no independent-sample inferential confidence intervals.", [3, 2, 7, 4, 4, 4])
    doc.add_page_break()
    rows = []
    for outcome in ["prevalence", "incidence"]:
        for sex in ["Male", "Female"]:
            for family, label in [("tcn_adapted__original", "TCN, original"), ("tcn_adapted__cdf", "TCN, matched"), ("mixture__equal_weight", "Mixture")]:
                sub = rates[(rates.target == "Saudi Arabia") & (rates.outcome == outcome) & (rates.sex == sex)
                            & (rates.horizon == 5) & (rates.scale == "rate") & (rates.family == family)]
                a = sub[sub.age_scope == "45+"].iloc[0]
                b = sub[sub.age_scope == "80+"].iloc[0]
                share = burdens[(burdens.target == "Saudi Arabia") & (burdens.outcome == outcome)
                    & (burdens.horizon == 5) & (burdens.population_method == "log_trend_last8")
                    & (burdens.node == sex+"__80+_within_45+") & (burdens.family == family)].iloc[0]
                rows.append([outcome.capitalize(), sex, label, n(a.coverage*100, 1), n(a.mean_width, 2), n(a.wis, 2),
                             n(b.coverage*100, 1), f"{round(share.coverage*5)}/5"])
    add_table(doc, "Table 3. Saudi rate and age-composition interval reliability",
              ["Outcome", "Sex", "Procedure", "45+ coverage (%)", "45+ width", "45+ WIS", "80+ rate coverage (%)", "80+ share covered"], rows,
              "TCN, temporal convolutional network; WIS, weighted interval score. Coverage/width refer to nominal 80% intervals; WIS uses 50%/80% intervals. Width and WIS are rate units per 100 000. The five overlapping horizon-five evaluations include 55 age–origin cells at 45+ and 20 at 80+; share coverage counts five dependent origins. Original TCN uses linearly interpolated quantiles. Matched TCN and mixture use inverse empirical cumulative-distribution quantiles; only their paired comparison isolates pooling. Share results use the preceding eight-year population log trend. Lower WIS is favourable; higher coverage is not sufficient evidence of improvement. Supplementary Tables S5–S6 provide other levels, widths, directional misses, comparators and decision rules.", [2.7, 1.8, 3.2, 3.3, 2.8, 2.6, 3.7, 3.4])
    doc.add_page_break()
    rows = []
    for r in projection.itertuples():
        rows.append([r.outcome.capitalize(), r.population_scenario.replace("UN", "United Nations"), n(r.native_GBD_2023),
                     n(r.scenario_baseline_2023), n(r.value), f"{n(r.lower)}–{n(r.upper)}", n(r.change_percent_from_scenario_baseline, 1)])
    add_table(doc, "Table 4. Conditional Saudi burden scenarios for 2028",
              ["Outcome", "Population scenario", "Native 2023", "Scenario 2023", "2028 point", "Conditional 80% interval", "Change from scenario baseline (%)"], rows,
              "Both sexes, ages ≥45 years; values are modelled numbers, not all-age totals. GBD, Global Burden of Disease. The GBD-baseline scenario scales its 2023 denominator by United Nations medium growth. The unaligned scenario applies United Nations levels at baseline and projection. Forecasts were generated in 2026 from a 2023 disease-data cutoff, not issued in 2023. Sixteen original residual blocks describe rate-conditional intervals under fixed population paths; these omit full demographic and GBD estimation uncertainty and are not validated total-uncertainty intervals. Baseline source differences must not be interpreted as disease growth. The mixture did not replace these projections. Data sources: [@gbd_results;@wpp;@wpp_methods].", [2.6, 6, 2.8, 3, 2.5, 4, 3])
    doc.save(OUT / "tables.docx")


def manuscript():
    doc = document()
    doc.add_heading(TITLE, 0)
    doc.add_paragraph("Running head: Sex-specific Parkinson’s forecasting")
    doc.add_paragraph("Authors, affiliations and corresponding-author details: to be supplied by the research team.")
    count = word_count(BODY, include_headings=True)
    abstract_count = word_count(ABSTRACT.removeprefix("# Abstract\n"), include_headings=True)
    doc.add_paragraph(f"Main-text word count: {count}; abstract: {abstract_count} (including section headings). Main figures: 4. Main tables: 4. References: {len(KEYS)}.")
    doc.add_paragraph("Complete scientific-text draft with structured abstract, keywords, key messages, Introduction, Methods, Results, Discussion and Conclusion. Author details, final declarations and supplement deposition require completion before submission.")
    doc.add_page_break()
    add_markdown(doc, ABSTRACT)
    doc.add_paragraph("Keywords: " + "; ".join(FRONT_MATTER["keywords"]))
    doc.add_heading("Key messages", 1)
    for message in FRONT_MATTER["key_messages"]:
        doc.add_paragraph(message, style="List Bullet")
    doc.add_page_break()
    add_markdown(doc, BODY)
    doc.add_heading("Declarations requiring author completion", 1)
    declarations = [
        ("Ethics approval", "The study analysed aggregate published/modelled estimates and involved no individual participant recruitment or identifiable patient records. The authors must confirm the applicable institutional ethics determination; no approval number or exemption has been invented."),
        ("Acknowledgements", "To be completed by the authors, including consent from any named contributors."),
        ("Author contributions", "Author roles and the guarantor must be supplied and approved by the research team."),
        ("Supplementary data", "The draft refers to Supplementary Material S1–S3, Supplementary Tables S1–S8 and Supplementary Software S1. These refer to the accompanying supplementary index and planned archive. Confirm deposition before replacing this text with the journal’s availability statement."),
        ("Conflict of interest", "To be declared by all authors; no absence of conflicts is assumed."),
        ("Funding", "Funding sources and grant numbers, or confirmation of no funding, must be supplied by the authors."),
        ("Data availability", "Analysis source code is available from the versioned public repository cited above. GBD and demographic source data are available through their providers subject to access/reuse terms. The supplementary index identifies the analytical tables, source citations, figures, scripts and provenance records intended for deposition. Repository code alone omits the original data and locked execution artifacts; a supplementary archive DOI/URL remains to be supplied."),
        ("Use of artificial intelligence tools", "OpenAI Codex assisted code development, literature retrieval, computational checks, figure coding and manuscript drafting. Figures were plotted from retained numerical outputs rather than generated as synthetic images. Authors must review and take responsibility for the final data, analyses, references and text; AI tools are not authors."),
    ]
    for heading, content in declarations:
        doc.add_heading(heading, 2)
        doc.add_paragraph(content)
    doc.add_heading("References", 1)
    for key in KEYS:
        doc.add_paragraph(f"{REFNUM[key]}. {REFERENCES[key]['text']}")
    doc.add_page_break()
    doc.add_heading("Figure captions", 1)
    for i, (title, caption, alt) in enumerate(CAPTIONS, 1):
        doc.add_heading(f"Figure {i}. {title}", 2)
        doc.add_paragraph(caption)
        doc.add_paragraph("Alt text: " + alt)
    doc.save(OUT / "manuscript.docx")
    numbered = CITE_RE.sub(lambda m: "[" + ",".join(str(REFNUM[x]) for x in m.group(1).replace("@", "").split(";")) + "]", BODY)
    front = ABSTRACT.strip() + "\n\nKeywords: " + "; ".join(FRONT_MATTER["keywords"])
    front += "\n\n# Key messages\n\n" + "\n\n".join("- " + message for message in FRONT_MATTER["key_messages"])
    text = "# " + TITLE + "\n\n" + front + "\n\n" + numbered + "\n\n# References\n\n"
    text += "\n\n".join(f"{REFNUM[k]}. {REFERENCES[k]['text']}" for k in KEYS)
    text += "\n\n# Figure captions\n\n" + "\n\n".join(f"## Figure {i}. {t}\n\n{c}\n\nAlt text: {a}" for i, (t,c,a) in enumerate(CAPTIONS,1))
    (OUT / "manuscript.md").write_text(text + "\n")
    (OUT / "abstract.txt").write_text(re.sub(r"^#+\s*", "", ABSTRACT.removeprefix("# Abstract\n").strip(), flags=re.M) + "\n")
    pd.DataFrame([dict(key=k, number=REFNUM[k], **REFERENCES[k]) for k in KEYS]).to_csv(OUT / "reference_register.csv", index=False)


def contradictions():
    doc = document()
    doc.add_heading("Methodological and demographic comparisons with earlier Parkinson’s studies", 0)
    doc.add_paragraph("This companion distinguishes differences in study questions, estimands and evaluation procedures. The comparisons below do not establish that an earlier study is incorrect.")
    items = [
        ("Long-range burden growth versus short-horizon accuracy", "Su and colleagues’ projections extend to 2050 and employ Bayesian model averaging with historical validation. Our five-year transfer comparison and exploratory equal-weight residual mixture have different information sets, losses, intervals and targets. Failure of our mixture rule neither refutes their projections nor validates their intervals for our setting.[@su]"),
        ("Descriptive sex patterns versus sex-specific transfer benefit", "Male/female differences in disease burden do not imply that a transfer model must improve both sexes. Female negative transfer concerns forecast transportability, not evidence against published sex-specific burden patterns.[@safiri;@norway]"),
        ("Global ageing contributions versus Saudi denominator error", "Published demographic decomposition concerns changes in projected population burden. Our 88.1%/97.0% figures partition particular signed male oldest-age forecast errors. They are neither global ageing contributions, variance-explained estimates nor causal effects.[@su]"),
        ("Source uncertainty versus forecast reliability", "GBD 95% uncertainty bounds describe estimation uncertainty. Our 50%/80% empirical forecast intervals evaluate future modelled point estimates. Comparing their nominal levels as though they were the same uncertainty quantity is invalid.[@gbd_nonfatal;@wis]"),
        ("National coverage versus World Bank classification", "The supplied 203-location panel matches the World Bank classification sets: 201 main national locations are present, plus Hong Kong and Macao, and three main national locations are absent. The hierarchy resolves classification, not population-boundary overlap.[@hierarchy]"),
        ("Population estimates from different providers", "GBD-implied 80+ Saudi populations differ from GASTAT by −37.39% for males and −49.65% for females. This comparison involves different providers and potentially different population definitions and reference dates. It does not demonstrate that either population estimate is erroneous.[@gastat;@wpp]"),
    ]
    for heading, text in items:
        doc.add_heading(heading, 1)
        rich(doc.add_paragraph(), text)
    doc.add_heading("References", 1)
    for key in KEYS:
        if key in {"safiri", "gbd_nonfatal", "su", "norway", "hierarchy", "gastat", "wpp", "wis"}:
            doc.add_paragraph(f"{REFNUM[key]}. {REFERENCES[key]['text']}")
    doc.save(OUT / "contradictions.docx")


def missing_references():
    rows = [
        ["R1", "Native GBD export citation year", "Both native exports supply a citation year of 2024; verified GBD 2023 release records/methods are dated 2025.", "Use the verified tool citation and release methods; confirm the exact recommended export citation with IHME before submission."],
        ["R2", "Original workbook/extract history", "The original extraction date, query record and workbook construction history are incomplete; a creation timestamp is not extraction provenance.", "Recover original query/export metadata if available. Regional numerical agreement does not resolve the remaining 196 locations."],
        ["R3", "GASTAT publication date", "The 2024 population edition and official workbook/report are identified, but precise publication date is not established.", "Retain publication-date-not-stated wording; verify against the official release notice rather than infer a date from a download URL."],
        ["R4", "GCC-Stat publication/update date", "The population dataflow/version and access date are recorded; a stable original publication date is not identified.", "Retain version/access-date citation and obtain an official release/version history if required."],
        ["R5", "Supplementary archive identifier", "No deposited supplementary archive DOI or journal-hosted URL exists in this draft.", "Supply the final archive identifier after deposition; do not claim that the supplement is already hosted by IJE."],
        ["R6", "Software archival citation", "Public source commit and account creator are verified; no versioned archival DOI is available.", "The commit URL is citable now. Confirm preferred creator names and obtain an archival DOI if desired; do not invent either."],
    ]
    doc = document(landscape=True)
    doc.add_heading("Missing references and unresolved citation metadata", 0)
    doc.add_paragraph("All named data sources used in the reported analyses have identifiable source citations. The items below are unresolved bibliographic/provenance details, not invented publications or silently omitted references. Main reference numbers are recorded in the reference register.")
    add_table(doc, "Items requiring author or provider resolution", ["ID", "Item", "Unresolved detail", "Action"], rows, sizes=[1.2, 4.3, 9, 9.5])
    doc.add_heading("Evidence gaps distinct from missing references", 1)
    doc.add_paragraph("Coherent joint disease/population trajectories and comparable population definitions remain data needs distinct from missing references. The hierarchy information sheet retains a 2021 sentence despite its 2023 title; this metadata detail does not change the verified location roster.")
    doc.add_heading("Acquired material not used as forecasting inputs", 1)
    doc.add_paragraph("Saudi Ministry of Health workforce/rehabilitation yearbooks and the MENASA registry questionnaire were acquired for background/readiness work, not model training or clinical validation. They are identified in the supplementary source crosswalk and should not be described as analytical patient data. The socio-demographic index was not an acquired model covariate; its absent values are not a missing citation for the fitted models.")
    for text in [
        "Ministry of Health, Kingdom of Saudi Arabia. Statistical Yearbook 2023. Riyadh: Ministry of Health, publication date not stated. https://www.moh.gov.sa/Ministry/Statistics/book/Documents/Statistical-Yearbook-2023.xlsx (26 September 2026, date accessed).",
        "Ministry of Health, Kingdom of Saudi Arabia. Statistical Yearbook 2024. Riyadh: Ministry of Health, publication date not stated. https://www.moh.gov.sa/Ministry/Statistics/book/Documents/Statistical-Yearbook-2024.xlsx (26 September 2026, date accessed).",
        "Khalil H, Shraim M, Jaradat B et al. Parkinson’s disease database in the Middle East, North Africa, and South Asia countries. Int J Public Health 2025;70:1608016. https://doi.org/10.3389/ijph.2025.1608016. Questionnaire only; no patient-level dataset analysed.",
    ]:
        doc.add_paragraph(text)
    doc.save(OUT / "missing_references.docx")
    pd.DataFrame(rows, columns=["id", "item", "unresolved_detail", "action"]).to_csv(OUT / "missing_references.csv", index=False)


SUPPLEMENT = [
    ("Supplementary Material S1", "Source provenance, locked protocol and amendments", ["study_design/locked_v1/protocol.md", "study_design/locked_v1/design.json", "study_design/data_source_references.md", "study_design/data_file_citations.csv", "results/source_asr_verification_v1/workbook_comparison_4284.csv", "results/source_asr_verification_v1/coverage.csv", "reports/source_hierarchy_verification_v1/report.md", "results/source_compatibility_v1/native_coverage.csv", "results/source_compatibility_v1/saudi_2023_population_comparison.csv", "results/source_compatibility_v1/forecast_draw_schema.json"]),
    ("Supplementary Material S2", "Exploratory mixture specification and complete decision rule", ["study_design/distribution_mixture_v1.md", "study_design/distribution_mixture_v1.json", "reports/distribution_mixture_v1/decision_gates.csv", "reports/distribution_mixture_v1/decisions.csv"]),
    ("Supplementary Material S3", "Execution provenance and independent verification", ["reports/global_asr_cpu_recovery_v1/report.md", "work/completion-validation/global_gpu_failure_review.json", "reports/study_synthesis_v1_2/validation.json", "reports/distribution_mixture_v1/validation.json"]),
    ("Supplementary Table S1", "All primary model-family and age-specific comparisons", ["results/primary_v1/primary_by_family.csv", "results/primary_v1/primary_age_contrasts.csv", "study_design/locked_v1/design.json"]),
    ("Supplementary Table S2", "Complete GCC endpoint and rolling model-family results", ["reports/secondary_v1/endpoint_contrasts.csv", "reports/secondary_v1/reliability_all_families.csv"]),
    ("Supplementary Table S3", "Donor selection, adaptation harm and target-history budgets", ["reports/donor_comparisons_gpu_v1/endpoint_percent_changes_vs_gpu_reference.csv", "results/learning_curves_v1/summary.csv", "results/learning_curves_v1/adaptation_comparisons.csv"]),
    ("Supplementary Table S4", "Separate global standardised-rate benchmark", ["reports/study_synthesis_v1_2/saudi_global_asr_comparisons.csv", "results/global_asr_cpu_recovery_v1/predictions.csv"]),
    ("Supplementary Table S5", "Original and earlier exploratory reliability procedures", ["reports/secondary_v1/reliability_80_rate_coverage.csv", "reports/reliability_v1_3/saudi_all_procedures.csv", "reports/reliability_v1_3/burden_endpoint_summary.csv"]),
    ("Supplementary Table S6", "Complete mixture interval, derived-burden and quantile-control results", ["reports/distribution_mixture_v1/rate_comparison.csv", "reports/distribution_mixture_v1/burden_comparison.csv", "reports/distribution_mixture_v1/quantile_convention_effect.csv"]),
    ("Supplementary Table S7", "Demographic scenarios, accounting and native-count coherence", ["reports/population_sensitivity_v1/saudi_population_source_comparison.csv", "reports/population_sensitivity_v1/saudi_80plus_share_endpoint.csv", "reports/population_sensitivity_v1/final_2018_error_accounting.csv", "results/count_coherence_v1/summary.csv", "reports/study_synthesis_v1_2/saudi_2028_scenario_totals.csv"]),
    ("Supplementary Table S8", "Mortality, disability, component and history sensitivities", ["reports/supporting_v1/endpoint_tcn_comparisons.csv", "reports/supporting_v1/rate_interval_five_origin.csv", "reports/supporting_v1/component_daly_comparison.csv", "reports/supporting_v1/prevalence_ratio_yld_comparison.csv", "reports/supporting_v1/mortality_history_comparison.csv"]),
    ("Supplementary Software S1", "Analysis, plotting, acquisition, tests and audit source", ["src/gbd_park/", "scripts/", "tests/", "study_design/locked_v1/build_design.py", "supporting_data/2026-09-26/", "supporting_data/2026-09-30_source_audit/", "work/"]),
]


def include_in_supplement(path):
    """Keep retired comparisons and their editorial discussion out of this deposit."""
    name = str(path.relative_to(ROOT)).lower()
    if any(token in name for token in ["release_sensitivity", "published_release_asr_comparison",
                                      "discrepancy_status_update", "discrepancy_ledger", "source_status.csv"]):
        return False
    if path.suffix.lower() == ".md":
        content = path.read_text(errors="replace")
        if re.search(r"107[.,]6\d*", content) and re.search(r"241[.,]68\d*", content):
            return False
    return True


def supplementary():
    rows = []
    for identifier, title, members in SUPPLEMENT:
        for member in members:
            assert (ROOT / member).exists(), member
            rows.append({"supplementary_identifier": identifier, "title": title, "archive_member": member,
                         "member_type": "directory" if (ROOT/member).is_dir() else "file"})
    figures = sorted(p for p in (ROOT / "reports").rglob("*.svg") if p.is_file() and include_in_supplement(p))
    for i, path in enumerate(figures, 1):
        rows.append({"supplementary_identifier": f"Supplementary Figure S{i}", "title": path.stem.replace("_", " "),
                     "archive_member": str(path.relative_to(ROOT)), "member_type": "file"})
    pd.DataFrame(rows).to_csv(OUT / "supplementary_index.csv", index=False)
    doc = document(landscape=True)
    doc.add_heading("Supplementary material index", 0)
    doc.add_paragraph("Archive members below are relative names within the supplementary deposit, not locations on an author’s computer. Original folder structure and filenames are preserved. A named supplementary table may comprise a clearly identified bundle of related machine-readable tables; the member list defines the bundle. Supporting figure files retain their paired raster versions where available. Materials are intended for deposition and are not asserted to be hosted by the journal already.")
    add_table(doc, "Cited supplementary bundles", ["Identifier", "Title", "Members in supplementary archive"],
              [[identifier, title, "; ".join(members)] for identifier, title, members in SUPPLEMENT], sizes=[4, 6, 14])
    doc.add_heading("Additional datasets, figures and scripts", 1)
    doc.add_paragraph("The machine-readable supplementary index assigns S-prefix identifiers to the existing report figures. The accompanying archive inventory lists additional source/derived data, tables, figures, scripts and provenance files in their original folders. Model checkpoints and transient caches are not counted as supplementary tables. Publisher upload-format restrictions may require a repository-hosted archive for scripts and nested directories; retain these member names and supply the final archive identifier.")
    doc.add_heading("Data-source attribution", 1)
    doc.add_paragraph("Regional disease estimates and their bounds cite the GBD Results Tool and relevant GBD 2023 methods; hierarchy files cite the hierarchy dataset DOI. United Nations estimates, corrections and marginal quantiles cite World Population Prospects 2024 and its methods. GASTAT and GCC-Stat population tables retain their respective institutional citations. The original citation crosswalk distinguishes analytical data from acquired background-only yearbooks and registry documentation; the missing-reference report records unresolved metadata.")
    for key in ["gbd_results", "gbd_nonfatal", "gbd_fatal", "hierarchy", "wpp", "wpp_methods", "gastat", "gccstat", "safiri", "software"]:
        doc.add_paragraph(f"{REFNUM[key]}. {REFERENCES[key]['text']}")
    doc.save(OUT / "supplementary_index.docx")
    roots = ["More data", "age_standard", "data", "reports", "results", "src", "scripts", "tests", "study_design", "supporting_data", "work"]
    suffixes = {".csv", ".gz", ".json", ".md", ".svg", ".png", ".pdf", ".xlsx", ".xls", ".docx", ".py", ".bib", ".html", ".txt"}
    inventory = []
    for folder in roots:
        for path in sorted((ROOT/folder).rglob("*")):
            if (path.is_file() and path.suffix.lower() in suffixes
                    and "__pycache__" not in path.parts and include_in_supplement(path)):
                inventory.append({"archive_member": str(path.relative_to(ROOT)), "bytes": path.stat().st_size,
                    "category": "software" if path.suffix == ".py" else "supporting_data_or_output"})
    pd.DataFrame(inventory).to_csv(OUT / "supplementary_archive_inventory.csv", index=False)


def review_package():
    thumbs = []
    for i in range(1, 5):
        im = Image.open(FIG / f"figure_{i}.png").convert("RGB")
        im.thumbnail((900, 900))
        canvas = Image.new("RGB", (930, 960), "white")
        canvas.paste(im, ((930-im.width)//2, 40))
        ImageDraw.Draw(canvas).text((25, 10), f"Figure {i}", fill="black")
        thumbs.append(canvas)
    sheet = Image.new("RGB", (1860, 1920), "#dddddd")
    for i, im in enumerate(thumbs):
        sheet.paste(im, ((i % 2)*930, (i // 2)*960))
    sheet.save(OUT / "figure_contact_sheet.png")
    preview = document()
    preview.add_heading("Figure review copy", 0)
    preview.add_paragraph("For author review only. Submit the separate high-resolution figures, not this embedded-image copy.")
    for i, (title, caption, alt) in enumerate(CAPTIONS, 1):
        preview.add_heading(f"Figure {i}. {title}", 1)
        preview.add_picture(str(FIG / f"figure_{i}.png"), width=Cm(16))
        preview.add_paragraph(caption)
        preview.add_paragraph("Alt text: " + alt)
        if i < 4:
            preview.add_page_break()
    preview.save(OUT / "figure_review.docx")
    count = word_count(BODY, include_headings=True)
    abstract_count = word_count(ABSTRACT.removeprefix("# Abstract\n"), include_headings=True)
    assert count <= 3000 and abstract_count <= 250 and len(KEYS) <= 50
    assert re.findall(r"^## (.+)$", ABSTRACT, flags=re.M) == ["Background", "Methods", "Results", "Conclusions"]
    assert not CITE_RE.search(ABSTRACT)
    assert len(FRONT_MATTER["key_messages"]) == 3
    assert 3 <= len(FRONT_MATTER["keywords"]) <= 10
    for i in range(1, 5):
        assert f"Figure {i}" in BODY and f"Table {i}" in BODY
    figure_checks = []
    for path in sorted(FIG.glob("*.tiff")):
        with Image.open(path) as im:
            assert im.width >= 3600 and min(im.info["dpi"]) >= 300
            figure_checks.append({"file": path.name, "width_pixels": im.width, "height_pixels": im.height, "dpi": [float(value) for value in im.info["dpi"]]})
    for path in OUT.glob("*.docx"):
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode()
            assert "/home/saif" not in xml and "/mnt/c/" not in xml and "file://" not in xml
            assert "[ @" not in xml and "[@" not in xml
    assert len(Document(OUT / "tables.docx").tables) == 4
    assert len(Document(OUT / "manuscript.docx").inline_shapes) == 0
    assert len(Document(OUT / "manuscript.docx").tables) == 0
    for name, expected in SOURCES.items():
        assert sha(ROOT / name) == expected
    instructions = f"""# Draft package and submission notes

The complete Introduction, Methods, Results, Discussion (including a separate Limitations subsection) and Conclusion contain {count} words including headings. The structured abstract contains {abstract_count} words including its four section headings. Seven keywords and three key messages are included. The package uses four main figures, four editable tables and {len(KEYS)} references. The main document is double-spaced, uses 2.5 cm margins and page fields, and includes figure captions/alt text after the references. Tables are separate and have no vertical rules. TIFF figures are 7200 pixels wide at 600 dpi; EPS masters and PNG review copies are also supplied.

IJE original-article guidance currently specifies a 3000-word main-text limit, a 250-word structured abstract, eight tables/figures combined and up to 50 references. The main text is {3000-count} words below the limit. Author details, final declarations and a supplement deposit identifier remain outstanding; the package is not claimed to be submission-complete. Source code is already public; the larger archive has not been uploaded by this drafting step. Guidance checked 30 September 2026: https://academic.oup.com/ije/pages/General_Instructions

## Files

- `manuscript.docx`: editable main text and figure captions; author-confirmation declarations are explicit.
- `tables.docx`: four editable main tables.
- `figures/figure_1.tiff` through `figure_4.tiff`: submission-resolution figures; matching EPS and PNG files.
- `figure_review.docx` and `figure_contact_sheet.png`: review aids, not additional main figures.
- `contradictions.docx`: companion covering methodological, demographic and forecast-evaluation differences with earlier studies.
- `missing_references.docx` and `.csv`: six unresolved bibliographic/provenance/deposition items.
- `supplementary_index.docx`, `.csv` and `supplementary_archive_inventory.csv`: S-prefix mappings and relative archive members.
- `reference_register.csv`: ordered references and verification status.
- `body.md`, `abstract.md`, `abstract.txt`, `front_matter.json`, `references.json`, `manuscript.md` and the build script: editable/reproducible drafting sources and abstract text for journal submission.

Main text uses supplementary identifiers, never author-machine storage paths. Archive member names appear only in the supplementary index. GBD 2023 is the study’s disease-estimation source. The manuscript retains the observed forecasting, calibration and demographic findings. Main figures and forecast values are unchanged.
"""
    (OUT / "submission_notes.md").write_text(instructions)
    validation = {"created_utc": datetime.now(timezone.utc).isoformat(), "passed": True,
        "main_text_words": count, "main_text_words_excluding_headings": word_count(BODY),
        "abstract_words": abstract_count, "word_counts_include_section_headings": True,
        "key_messages": len(FRONT_MATTER["key_messages"]), "keywords": len(FRONT_MATTER["keywords"]),
        "figures": 4, "tables": 4, "combined_display_items": 8,
        "references": len(KEYS), "all_figure_and_table_callouts_present": True,
        "no_absolute_storage_paths_in_docx": True, "figure_resolution": figure_checks,
        "source_sha256": SOURCES, "builder_sha256": sha(Path(__file__)),
        "models_fitted": 0, "results_changed": False, "release_comparison_tables_in_companion": 0,
        "supplementary_tables": 8,
        "author_declarations_finalised": False, "supplement_deposited": False,
        "artifact_sha256": {str(p.relative_to(OUT)): sha(p) for p in OUT.rglob("*") if p.is_file() and p.name not in {"validation.json", "quality_review.json"}}}
    (OUT / "validation.json").write_text(json.dumps(validation, indent=2)+"\n")


if __name__ == "__main__":
    main()
