"""Assemble JDR manuscripts, six editable tables and reporting documentation."""
from copy import deepcopy
import csv
import hashlib
import importlib.util
import json
from pathlib import Path
import re
import shutil
import zipfile

import numpy as np
import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'manuscript/JDR_GBD_PARK'
AUTH=OUT/'authoring'
TITLE=json.loads((AUTH/'editorial_specification.json').read_text())['title']
BODY=(AUTH/'body.txt').read_text()
FULL_LIMITATIONS=(AUTH/'limitations_full.txt').read_text().strip()
ABSTRACT=(AUTH/'abstract.txt').read_text().strip()
REFERENCE_REGISTER=json.loads((OUT/'references/reference_register.json').read_text())
MAIN_REFERENCE_KEYS=set(re.findall(r'@([a-z_]+)',BODY))
assert MAIN_REFERENCE_KEYS<=set(REFERENCE_REGISTER)
REFS={key:record for key,record in REFERENCE_REGISTER.items() if key in MAIN_REFERENCE_KEYS}
CAPTIONS=json.loads((AUTH/'figure_captions.json').read_text())
PLACEMENTS=json.loads((AUTH/'inline_placements.json').read_text())
KEYWORDS=['Parkinson’s disease','disability','transfer learning','temporal convolutional network',
          'forecast reliability','population ageing','Saudi Arabia','Gulf Cooperation Council']
CITE=re.compile(r'\[@([a-z_]+(?:;@[a-z_]+)*)\]')
SOURCES={}
DATA_ACKNOWLEDGEMENT=(
 'Source: Institute for Health Metrics and Evaluation. Used with permission. All rights reserved. '
 'Population inputs are credited to the United Nations Department of Economic and Social Affairs, Population Division; '
 'the General Authority for Statistics, Saudi Arabia; and the Ministry of Health Portal (https://www.moh.gov.sa/). '
 'Source: GCC-Statistical Centre, Sultanate of Oman (https://gccstat.org/); the authors selected, aggregated and '
 'harmonised these population data. Insurance tabulations are credited to the Council of Health Insurance. '
 'UN population data are © 2024 United Nations, CC BY 3.0 IGO (https://creativecommons.org/licenses/by/3.0/igo/). '
 'All forecasts, demographic transformations, decompositions and graphics are the authors’ analyses. '
 'The data providers do not endorse these analyses or conclusions.')


def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def read(path):
 SOURCES[path]=sha(ROOT/path)
 return pd.read_csv(ROOT/path,float_precision='round_trip')


def resolve(text,blind=False):
 def replacement(m):
  keys=m.group(1).replace('@','').split(';')
  keys.sort(key=lambda k:(REFERENCE_REGISTER[k]['sort_author'].lower(),REFERENCE_REGISTER[k]['display_year']))
  return '; '.join(('Study software, 2026' if blind and k=='software' else REFERENCE_REGISTER[k]['cite']) for k in keys)
 return CITE.sub(replacement,text)


def document(blind=False):
 doc=Document()
 example=Document(OUT/'Example_Manuscript_JDR.docx')
 # Reuse the example's semantic styles without copying its text, images or relationships.
 styles=doc.styles.element
 for item in list(styles):styles.remove(item)
 for item in example.styles.element:styles.append(deepcopy(item))
 sec=doc.sections[0]
 sec.page_width,sec.page_height=Inches(8.27),Inches(11.69)
 sec.top_margin=sec.bottom_margin=Inches(.75)
 sec.left_margin=sec.right_margin=Inches(.80)
 doc.styles['Normal'].font.name='Times New Roman';doc.styles['Normal'].font.size=Pt(11.5)
 doc.styles['Normal'].paragraph_format.line_spacing=1.18
 doc.styles['Normal'].paragraph_format.space_after=Pt(6)
 for name in ['Heading 1','Heading 2','Heading 3','Title']:
  doc.styles[name].font.name='Times New Roman';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
  doc.styles[name].paragraph_format.keep_with_next=True
 doc.styles['Heading 1'].font.size=Pt(16)
 doc.styles['Heading 2'].font.size=Pt(13)
 doc.styles['Heading 3'].font.size=Pt(11.5)
 for p in [sec.header.paragraphs[0],sec.footer.paragraphs[0]]:
  p.paragraph_format.line_spacing=1
 sec.header.paragraphs[0].text='Parkinson’s forecasting and disability care'
 sec.header.paragraphs[0].runs[0].font.size=Pt(9)
 footer=sec.footer.paragraphs[0];footer.alignment=2
 run=footer.add_run('Page ');run.font.size=Pt(9)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 doc.core_properties.title=TITLE
 doc.core_properties.author='' if blind else 'Saif Khan; Mahvish Khan; Mohtashim Lohani'
 doc.core_properties.last_modified_by=''
 doc.core_properties.comments=''
 # Preserve the full-resolution publication masters when the document is edited.
 compression=OxmlElement('w:doNotAutoCompressPictures')
 doc.settings.element.append(compression)
 return doc


def add_table(doc,table):
 p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
 p.add_run(f"Table {table['number']}. {table['title']}").bold=True
 t=doc.add_table(rows=1,cols=len(table['columns']));t.autofit=False
 total=sum(table['widths']);widths=[6.67*v/total for v in table['widths']]
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','bottom','left','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');e.set(qn('w:sz'),'8');borders.append(e)
 t._tbl.tblPr.append(borders)
 for cell,text in zip(t.rows[0].cells,table['columns']):
  cell.text=text
  for r in cell.paragraphs[0].runs:r.bold=True
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(shade)
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for values in table['rows']:
  cells=t.add_row().cells
  for cell,value in zip(cells,values):cell.text=str(value)
 for row in t.rows:
  no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
  for cell,w in zip(row.cells,widths):
   cell.width=Inches(w)
   for p in cell.paragraphs:
    p.paragraph_format.line_spacing=1.08;p.paragraph_format.space_after=Pt(4)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.keep_together=True
    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9.5)
 p=doc.add_paragraph(resolve(table['note']))
 p.paragraph_format.line_spacing=1.08
 p.paragraph_format.keep_together=True
 p.paragraph_format.keep_with_next=False
 for r in p.runs:r.font.size=Pt(9)


def add_figure(doc,caption,blind=False):
 """Embed the unchanged 600-dpi TIFF master with its title and complete caption."""
 p=doc.add_paragraph()
 p.paragraph_format.keep_with_next=True
 p.paragraph_format.keep_together=True
 p.add_run(f"Figure {caption['number']}. {caption['title']}").bold=True
 p=doc.add_paragraph()
 p.paragraph_format.line_spacing=1
 p.paragraph_format.keep_with_next=True
 p.paragraph_format.keep_together=True
 shape=p.add_run().add_picture(str(OUT/'figures'/f"figure_{caption['number']}.tiff"),width=Inches(6.55))
 shape._inline.docPr.set('descr',caption['alt'])
 shape._inline.docPr.set('title',caption['title'])
 p=doc.add_paragraph(resolve(caption['caption'],blind))
 p.paragraph_format.line_spacing=1.08
 p.paragraph_format.keep_together=True
 p.paragraph_format.keep_with_next=False
 for r in p.runs:r.font.size=Pt(10)


def tables():
 result=[]
 def add(n,title,columns,rows,note,widths):
  t=dict(number=n,title=title,columns=columns,rows=rows,note=note,widths=widths);result.append(t)
 add(1,'Data sources, estimands and analytical roles',
     ['Source','Data and scope','Role'],[
     ['GBD 2023 results','Seven countries; male/female; eleven ages ≥45; annual 1990–2023 rates and numbers','Primary prevalence, secondary incidence, supporting YLDs, YLLs, DALYs and deaths'],
     ['Earlier fatal outcomes','1980–1989 death and YLL estimates','History-length sensitivity'],
     ['Native standardised export and hierarchy','4,704 regional export records; official cause/location classifications','Verification of 4,284 overlapping values; roster checking'],
     ['Global standardised-rate panel','203 World Bank-classified GBD locations','Separate global-versus-regional donor benchmark'],
     ['UN World Population Prospects 2024','Age–sex populations; medium estimates and projections through 2028','Counts and growth paths'],
     ['Saudi GASTAT and MOH yearbook','National 2022–2024 age–sex populations; citizenship categories','Denominator assessment and national-anchor scenarios'],
     ['GCC Statistical Centre','Regional population tables','Source cross-checks'],
     ['Council of Health Insurance','Seven quarterly releases, original subscriber/insured categories','Insurance-population composition'],
     ],'GBD, Global Burden of Disease; UN, United Nations; GASTAT, General Authority for Statistics; MOH, Ministry of Health; GCC, Gulf Cooperation Council; YLDs, years lived with disability; YLLs, years of life lost; DALYs, disability-adjusted life-years. Disease outcomes are modelled estimates; rates and counts remain separate analytical quantities. Full source citations and raw-to-derived mappings: Methods and Supplementary Materials S1 and S4.',[1.45,2.75,2.47])
 endpoint=read('reports/secondary_v1/endpoint_contrasts.csv')
 names={'damped_ets':'Damped exponential smoothing','pooled_boosting':'Pooled boosting','arima':'ARIMA',
        'donor_ridge_adapted':'Adapted donor ridge','donor_boosting_adapted':'Adapted donor boosting','pooled_ridge':'Pooled ridge'}
 rows=[]
 for o in ['prevalence','incidence']:
  for s in ['Male','Female']:
   for c in ['local_champion','nonneural_champion']:
    r=endpoint.loc[endpoint.target.eq('Saudi Arabia')&endpoint.outcome.eq(o)&endpoint.sex.eq(s)&endpoint.comparator.eq(c)].iloc[0]
    rows.append([o.capitalize(),s,names[r.comparator_source_family],f'{r.comparator_error:.6f}',f'{r.tcn_error:.6f}',f'{r.relative_improvement_percent:+.2f}'])
 add(2,'Saudi five-year endpoint comparisons at the 2018 origin',
     ['Outcome','Sex','Comparator','Comparator error','TCN error','Improvement (%)'],rows,
     'Error is mean absolute log error across eleven equally weighted age groups, verified against 2023 rates. Positive improvement means lower TCN error: 100 × (comparator error − TCN error)/comparator error. TCN, temporal convolutional network; ARIMA, autoregressive integrated moving average. The joint primary prevalence criterion requires all four prevalence contrasts to be positive.',[.85,.55,1.8,1.1,1.1,1.15])
 rates=read('reports/distribution_mixture_v1/rate_comparison.csv')
 burdens=read('reports/distribution_mixture_v1/burden_comparison.csv')
 rows=[]
 families=[('tcn_adapted__original','Original TCN'),('tcn_adapted__cdf','Matched TCN'),('mixture__equal_weight','Mixture')]
 for o in ['prevalence','incidence']:
  for s in ['Male','Female']:
   for f,label in families:
    r=rates.loc[rates.target.eq('Saudi Arabia')&rates.outcome.eq(o)&rates.sex.eq(s)&rates.family.eq(f)&rates.horizon.eq(5)&rates.scale.eq('rate')&rates.age_scope.eq('45+')].iloc[0]
    b=burdens.loc[burdens.target.eq('Saudi Arabia')&burdens.outcome.eq(o)&burdens.sex.eq(s)&burdens.family.eq(f)&burdens.horizon.eq(5)&burdens.population_method.eq('log_trend_last8')&burdens.node.eq(s+'__80+_within_45+')].iloc[0]
    rows.append([o.capitalize(),s,label,f'{100*r.coverage:.1f}',f'{r.mean_width:.2f}',f'{r.wis:.2f}',f'{round(b.coverage*5)}/5'])
 add(3,'Five-origin rate and oldest-age-share interval evaluation',
     ['Outcome','Sex','Procedure','Rate coverage (%)','Rate width','Rate WIS','80+ share covered'],rows,
     'Coverage and width refer to nominal 80% intervals; rate WIS combines 50% and 80% intervals and is expressed in rate units per 100,000. Each rate row contains 55 dependent age–origin cells; the share uses five overlapping-origin observations. Original TCN uses interpolated empirical quantiles; matched TCN and the equal-weight mixture use inverse empirical-distribution quantiles. Shares use the preceding eight-year population log trend. WIS, weighted interval score.',[.85,.55,1.10,1.1,.8,.8,1.1])
 d=read('manuscript/JDR_GBD_PARK/analysis/disability_reporting_summary.csv')
 rows=[]
 labels={'ylds':'YLDs','ylls':'YLLs','dalys':'DALYs','deaths':'Deaths'}
 for _,r in d.iterrows():rows.append([labels[r.outcome],r.sex,f'{r.tcn_error:.5f}',f'{r.nonneural_error:.5f}',f'{r.rate_80_coverage_percent:.1f}',f'{r.forecast_80plus_share:.2f} / {r.gbd_80plus_share:.2f}',f'{round(r.share_80_coverage_percent/20)}/5'])
 add(4,'Disability and mortality forecasting in Saudi Arabia',
     ['Outcome','Sex','TCN rate error','Non-neural error','Rate coverage (%)','80+ share: forecast / GBD (%)','Share covered'],rows,
     'Rate errors refer to the 2018→2023 endpoint. Coverage is for original nominal 80% intervals over origins 2014–2018. Rate coverage uses 55 dependent age–origin cells per sex. Age shares use five aggregate observations, direct outcome forecasts and operational population trends. YLDs, years lived with disability; YLLs, years of life lost; DALYs, disability-adjusted life-years; GBD, Global Burden of Disease.',[.65,.55,.9,.95,1.05,1.7,.87])
 forecasts=read('forecast_percentage_change.csv')
 p=forecasts.loc[forecasts.population_scenario.eq('gbd_2023_aligned_un_growth')]
 rows=[]
 for o in ['prevalence','incidence']:
  for s in ['Male','Female','Both']:
   part=p.loc[p.outcome.eq(o)&p.sex.eq(s)].sort_values('forecast_year')
   rows.append([o.capitalize(),s,f'{part.baseline_count_2023.iloc[0]:,.0f}']+[f'{v:,.0f}' for v in part.forecast_count]+[f'{part.percentage_change_from_2023.iloc[-1]:+.2f}'])
 add(5,'Saudi case-count forecasts and changes from the 2023 baseline',
     ['Outcome','Sex','2023','2024','2025','2026','2027','2028','Change (%)'],rows,
     'Saudi residents aged ≥45; adapted TCN, original GBD baseline and UN population growth. Prevalence denotes prevalent cases; incidence, annual incident cases. Both sums male and female counts. Displayed counts are rounded; 2023–2028 changes use unrounded values. Disease-data cutoff: 2023; conditional projections computed in 2026. Exact values and alternatives: Supplementary Table S12.',[.95,.6,.72,.72,.72,.72,.72,.72,.80])
 dec=read('manuscript/JDR_GBD_PARK/analysis/forecast_growth_decomposition.csv')
 dec=dec.loc[dec.year.eq(2028)&dec.scenario.eq('gbd_2023_aligned_un_growth')]
 rows=[]
 for o in ['prevalence','incidence']:
  for s in ['Male','Female','Both']:
   r=dec.loc[dec.outcome.eq(o)&dec.sex.eq(s)].iloc[0]
   rows.append([o.capitalize(),s]+[f'{r[k]:+.2f}' for k in ['population_size_percentage_point_contribution','composition_percentage_point_contribution','rates_percentage_point_contribution','total_change_percent']])
 add(6,'Exploratory decomposition of the 2023–2028 percentage change',
     ['Outcome','Sex','Population size (points)','Composition (points)','Disease rates (points)','Total change (%)'],rows,
     'Components average incremental contributions over six factor-replacement orders, divided by matched 2023 burden. Population size covers residents aged ≥45; composition represents age within sex and age–sex jointly for both sexes. Original GBD-baseline/UN-growth scenario. Unrounded contributions sum exactly to total change; rounding affects displayed sums. Full results: Supplementary Table S12.',[1,.65,1.3,1.25,1.25,1.22])
 (OUT/'tables').mkdir(exist_ok=True)
 for t in result:
  with (OUT/'tables'/f"table_{t['number']}.csv").open('w',newline='',encoding='utf-8-sig') as h:
   w=csv.writer(h);w.writerow(t['columns']);w.writerows(t['rows'])
 for t in result[1:]:
  t['note']+=' Source: authors’ calculations using GBD 2023 estimates ([@gbd_results])'+(' and UN World Population Prospects 2024 ([@wpp])' if t['number'] in [5,6] else '')+'.'
 (AUTH/'tables.json').write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n')
 return result


def add_title_page(doc):
 doc.add_heading(TITLE,0)
 doc.add_paragraph('RESEARCH ARTICLE','Article Type')
 source=Document(OUT/'Example_Manuscript_JDR.docx')
 p=doc.add_paragraph(style='Author Line')
 for i,(name,numbers) in enumerate([('Saif Khan','1,4'),('Mahvish Khan','2,4*'),('Mohtashim Lohani','3,4')]):
  if i:p.add_run(', ')
  p.add_run(name);p.add_run(numbers).font.superscript=True
 for i in range(3,7):
  text=source.paragraphs[i].text.replace('College of science','College of Science')
  doc.add_paragraph(text,'Author Line' if i==2 else 'Affiliation')
 doc.add_paragraph('Corresponding author: Mahvish Khan\nDepartment of Biology, College of Science, University of Ha’il, Ha’il 55473, Saudi Arabia\nEmail: mk.khan@uoh.edu.sa')
 doc.add_paragraph('Running title: Parkinson’s forecasting and disability care')
 doc.add_paragraph('Funding: The authors acknowledge funding from the King Salman Center for Disability Research through Research Group Number KSRG-2026-489.')
 doc.add_paragraph('Competing interests: The authors declare no competing interests.')
 doc.add_paragraph('Ethics and consent: This secondary aggregate-data analysis used public population estimates and demographic tabulations, without recruitment, intervention, biological specimens or identifiable individual records. Individual informed consent was not applicable.')
 doc.add_paragraph('Data and code availability: The sources are identified in the references and supplementary citation crosswalk. Derived results, analysis specifications and reporting scripts are indexed in the supplementary material. Core forecasting code is available at https://github.com/saifkhanbio/Parkinson-TCN-forecast; Supplementary Software S1 includes the later reporting extensions.')
 doc.add_paragraph('Acknowledgements: We acknowledge the organisations providing the disease estimates, demographic data and administrative tabulations used in this study.')


def add_references(doc,blind):
 doc.add_heading('References',1)
 for key,r in sorted(REFS.items(),key=lambda kv:(kv[1]['sort_author'].lower(),kv[1]['display_year'],kv[1]['title'])):
  p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.3);p.paragraph_format.first_line_indent=Inches(-.3)
  if blind and key=='software':p.add_run('Study software. (2026). Core forecasting source snapshot [Computer software]. Author-identifying repository details supplied separately for blinded review.');continue
  p.add_run(r['reference_prefix']+r['reference_title']);p.add_run(r['reference_italic']).italic=True;p.add_run(r['reference_suffix'])


def manuscript(tables,blind=False):
 doc=document(blind)
 if not blind:add_title_page(doc);doc.add_page_break()
 doc.add_heading(TITLE,0)
 doc.add_heading('Abstract',1);doc.add_paragraph(ABSTRACT)
 p=doc.add_paragraph();p.add_run('Keywords: ').bold=True;p.add_run('; '.join(KEYWORDS))
 inserted=[]
 in_limitations=False
 for block in BODY.strip().split('\n\n'):
  if block.startswith('## '):
   in_limitations=block=='## Limitations'
   doc.add_heading(block[3:],2)
  elif block.startswith('# '):
   in_limitations=False
   doc.add_heading(block[2:],1)
  else:
   p=doc.add_paragraph(resolve(block,blind))
   if in_limitations:p.paragraph_format.keep_together=True
  for placement in PLACEMENTS:
   if block.startswith(placement['after']):
    for kind,number in placement['items']:
     assert (kind,number) not in inserted
     if kind=='table':add_table(doc,next(t for t in tables if t['number']==number))
     else:add_figure(doc,next(c for c in CAPTIONS if c['number']==number),blind)
     inserted.append((kind,number))
 assert sorted(inserted)==sorted([('table',i) for i in range(1,7)]+[('figure',i) for i in range(1,9)])
 doc.add_heading('Supplementary material and availability',1)
 doc.add_paragraph('Supplementary Materials S1–S5, Supplementary Tables S1–S13, Supplementary Figures S1–S35 and Supplementary Software S1 provide the protocol, source documentation, complete results, sensitivity analyses and computational resources. The supplementary index identifies each item by its archive member and retains the original folder organisation. Figures and tables in the main manuscript were regenerated from these numerical records.')
 doc.add_paragraph('The complete provider bibliography and file-to-citation crosswalk in Supplementary Material S1 also identify acquisition-only and eligibility-review resources, including the WHO Mortality Database, CHI/NPHIES summaries and NHRSP survey documentation, separately from the datasets used for modelling and reported numerical comparisons.')
 doc.add_heading('Data-source acknowledgements',2)
 doc.add_paragraph(DATA_ACKNOWLEDGEMENT)
 if blind:
  doc.add_paragraph('Author-identifying funding, acknowledgements, institutional and code-repository details are supplied on the separate title page. The scientific methods describe the aggregate public-data design and software versions.')
 add_references(doc,blind)
 name='JDR_Manuscript_Blinded.docx' if blind else 'JDR_Manuscript.docx'
 doc.save(OUT/name)
 return OUT/name


def supplementary_index():
 old=read('manuscript/ije_draft_v2_raw_integration/supplementary_index.csv')
 # The current provider register supersedes the older citation wording, without editing it.
 old=old.loc[~old.archive_member.str.startswith('study_design/dataset_citations_2026-10-01/') &
             ~old.archive_member.isin(['study_design/data_source_references.md','study_design/data_file_citations.csv'])].copy()
 new=[]
 def add(identifier,title,path,kind='file'):
  assert (ROOT/path).exists(),path
  new.append(dict(supplementary_identifier=identifier,title=title,archive_member=path,member_type=kind))
 for name in ['Data_Citation_Audit.docx','data_citation_audit.csv','provider_citation_audit/retrieval_manifest.json',
              'Supplementary_Data_References.docx','data_source_register.csv','data_file_citation_crosswalk.csv',
              'citation_audit_validation.json','provider_citation_audit/unresolved_metadata.csv']:
  add('Supplementary Material S1','Provider citation and attribution audit','manuscript/JDR_GBD_PARK/references/'+name)
 for p,title in [('forecast_percentage_change.csv','Exact 2024–2028 changes from matched 2023 baselines'),
     ('forecast.csv','Exact forecasts and published projection context'),
     ('manuscript/JDR_GBD_PARK/analysis/forecast_growth_decomposition.csv','All demographic decompositions'),
     ('manuscript/JDR_GBD_PARK/analysis/saudi_2023_burden_composition.csv','Baseline burden by sex and age'),
     ('reports/forecast_percentage_change_v1/scenario_baseline_2023_age_cells.csv','Matched scenario baseline cells')]:add('Supplementary Table S12',title,p)
 add('Supplementary Table S13','Resources and reporting-rigor crosswalk','manuscript/JDR_GBD_PARK/reporting_checklist.csv')
 add('Supplementary Material S5','JDR methods reproducibility and reporting audit','manuscript/JDR_GBD_PARK/Reporting_Rigor_Checklist.docx')
 add('Supplementary Material S5','Full study limitations','manuscript/JDR_GBD_PARK/Limitations.docx')
 add('Supplementary Material S5','Exploratory decomposition validation','manuscript/JDR_GBD_PARK/analysis/validation.json')
 add('Supplementary Material S5','Recent literature-window verification','study_design/literature_review/saudi_forecast_window_2026-10-03/verification.json')
 for script in ['jdr_additional_analysis.py','plot_jdr_figures.py','plot_forecast_percentage_change.py','build_jdr_manuscript.py','build_jdr_references.py','jdr_reference_metadata.py']:
  add('Supplementary Software S1','JDR reporting and exploratory analysis software','scripts/'+script)
 index=pd.concat([old,pd.DataFrame(new)],ignore_index=True)
 assert not index.archive_member.str.contains('/home/|/mnt/').any()
 for p in index.archive_member:assert (ROOT/p).exists(),p
 index.to_csv(OUT/'supplementary_index.csv',index=False)
 doc=document();doc.add_heading('Supplementary material index',0)
 doc.add_paragraph('Archive members below are paths within the supplementary archive, not local storage locations. Original folders are retained. This index extends the previous supplementary package; it does not duplicate the full raw-data archive.')
 for identifier,part in index.groupby('supplementary_identifier',sort=False):
  doc.add_heading(identifier,1)
  for _,r in part.iterrows():doc.add_paragraph(r.title+' — '+r.archive_member)
 doc.save(OUT/'Supplementary_Index.docx')


def reporting_checklist():
 rows=[
 ('Study design','Applicable','Population-level retrospective forecasting and conditional projection','Methods: Design, setting and units of analysis'),
 ('Inclusion and exclusion','Applicable','Seven countries, two source sex categories, eleven ages ≥45; complete common-period grids; younger ages outside scope','Methods: Data sources, eligibility and integrity'),
 ('Demographics / sex','Applicable','Sex-separated and age-specific reporting; citizens and non-citizens in resident totals','Methods: Design; Table 1'),
 ('Sample size / power','Applicable','All available eligible annual cells; 34 years per common stratum; no hypothesis-test power calculation','Methods: Sample-size determination'),
 ('Randomisation','No participant allocation','Calendar-time partitions; fixed computational and random-donor seeds','Methods: Sample-size determination; Donor choice'),
 ('Blinding','Analysts used labelled data','Bias-control procedures stated; prior outcome inspection disclosed in Limitations','Methods: Sample-size determination; Discussion: Limitations'),
 ('Replication','Computational and temporal','Five seeds, five overlapping origins and six target-country applications separately defined','Methods: Sample-size determination'),
 ('Attrition / failed fits','Applicable','Status-ledger persistence fallback retains rows; CPU-recovered global benchmark identified','Methods: Computational attrition; Supplementary Tables S1–S4'),
 ('Ethics / consent','Aggregate secondary data','No recruitment, intervention, specimens or identifiable records; no fabricated approval identifier','Methods: Ethics; separate title page'),
 ('Antibodies / cell lines / organisms','Not applicable','No wet-laboratory or experimental-animal resources','Methods: Ethics'),
 ('Protocol','Applicable','Local lock 29 September 2026; amendments and later exploratory work identified','Methods: Bias control; Supplementary Material S1'),
 ('Training/test separation','Applicable','Complete five-year labels before origin; fit-only scaling; target excluded from donors','Methods: Chronological development; TCN'),
 ('Algorithms / hyperparameters','Applicable','Architecture, loss, optimiser, grid, seeds and two-parameter adaptation reported','Methods: Comparators; TCN; Supplementary Table S1'),
 ('Statistics','Applicable','Loss, relative improvement, interval score and WIS defined; dependence-aware descriptive comparisons','Methods: Accuracy and predictive intervals'),
 ('Data identifiers','Applicable','Provider URLs, hierarchy DOI and dataset citation crosswalk','References; Supplementary Material S1'),
 ('Code availability','Applicable','Public core snapshot plus supplementary later reporting scripts; blinded manuscript withholds identifying link','Methods: Resources; title page; Supplementary Software S1'),
 ('Resources / RRIDs','Applicable','Python SCR_008394; PyTorch SCR_018536; NumPy SCR_008633; SciPy SCR_008058; scikit-learn SCR_002577; Matplotlib SCR_008624','Methods: Resources'),
 ('GATHER','Applicable items mapped','Sources, definitions, processing, models, evaluation and uncertainty documented','Methods; Limitations; source crosswalk'),
 ('TRIPOD+AI','Applicable transparency items','Used as a reporting aid for population forecasting; not labelled a clinical individual-risk validation','Methods; Limitations'),
 ('New decomposition','Exploratory','Six-order average; arithmetic identity; independent closed form and synthetic boundary checks','Methods; Figure 8; Supplementary Table S12'),
 ('Figure accessibility','Applicable','Eight figures with embedded alternative text; distinguishable markers and vector masters','Main figures; Figure_Captions.docx'),
 ('Actual SciScore assessment','Not run','This is an author reporting crosswalk, not a new SciScore score or endorsement','Submission notes'),
 ]
 d=pd.DataFrame(rows,columns=['criterion','applicability','evidence','manuscript_location']);d.to_csv(OUT/'reporting_checklist.csv',index=False)
 doc=document();doc.add_heading('Reporting rigor and resource checklist',0)
 doc.add_paragraph('Prepared against the supplied example SciScore report and the public Core/MDAR descriptions. This document maps actual study procedures to reporting items. It is not an automated SciScore assessment and assigns no score.')
 for criterion,applicability,evidence,where in rows:
  doc.add_heading(criterion,2);doc.add_paragraph(applicability+'. '+evidence+'. '+where+'.')
 doc.add_paragraph('Source guidance: https://sciscore.com/reports/Core-Report.php ; supplied Example_SciscoreReport.pdf. Reporting-guideline and core-software references are retained below for this supplementary crosswalk.')
 for key in ['gather','tripod','software']:doc.add_paragraph(REFERENCE_REGISTER[key]['apa_text'])
 doc.save(OUT/'Reporting_Rigor_Checklist.docx')


def missing_reference_metadata():
 source=read('manuscript/JDR_GBD_PARK/references/provider_citation_audit/unresolved_metadata.csv')
 source.to_csv(OUT/'references/missing_reference_metadata.csv',index=False)
 doc=document();doc.add_heading('Reference metadata requiring source confirmation',0)
 doc.add_paragraph('Every dataset used in the manuscript has an identifiable provider citation. The GBD 2023 citation year is now resolved to 2025 by current IHME guidance. This list records the remaining publication-date and export-provenance details; these are distinguished from a missing institutional attribution. The public software URL supplies an identifier even without an archival DOI.')
 for _,row in source.iterrows():doc.add_paragraph(' | '.join(str(v) for v in row.values if pd.notna(v)))
 doc.save(OUT/'references/Missing_Reference_Metadata.docx')


def full_limitations_document():
 doc=document(blind=True)
 doc.add_heading('Limitations',0)
 for paragraph in FULL_LIMITATIONS.split('\n\n'):
  p=doc.add_paragraph(resolve(paragraph));p.paragraph_format.keep_together=True
 keys=set(re.findall(r'@([a-z_]+)',FULL_LIMITATIONS))
 if keys:
  doc.add_heading('References',1)
  for key in sorted(keys,key=lambda k:(REFERENCE_REGISTER[k]['sort_author'].lower(),REFERENCE_REGISTER[k]['display_year'])):
   record=REFERENCE_REGISTER[key]
   p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.3);p.paragraph_format.first_line_indent=Inches(-.3)
   p.add_run(record['reference_prefix']+record['reference_title']);p.add_run(record['reference_italic']).italic=True;p.add_run(record['reference_suffix'])
 doc.save(OUT/'Limitations.docx')


def main():
 short_limitations=re.search(r'(?<=## Limitations\n\n)(.*?)(?=\n\n# Conclusions)',BODY,re.S).group(1)
 limitations_words=len(resolve(short_limitations).split())
 assert limitations_words+1<=125,limitations_words
 full_limitations_document()
 caption_counts={str(c['number']):len(f"Figure {c['number']}. {c['title']} {resolve(c['caption'])}".split()) for c in CAPTIONS}
 assert all(count<=100 for count in caption_counts.values()),caption_counts
 (OUT/'references/manuscript_reference_register.json').write_text(json.dumps(REFS,ensure_ascii=False,indent=2)+'\n')
 protected=[OUT/'Example_Manuscript_JDR.docx',OUT/'Example_SciscoreReport.pdf']
 protected += [p for folder in ['ije_draft_v1','ije_draft_v2_raw_integration'] for p in (ROOT/'manuscript'/folder).rglob('*') if p.is_file()]
 before={str(p.relative_to(ROOT)):sha(p) for p in protected}
 ts=tables()
 table_caption_counts={str(t['number']):len(f"Table {t['number']}. {t['title']} {resolve(t['note'])}".split()) for t in ts}
 assert all(count<=100 for count in table_caption_counts.values()),table_caption_counts
 for t in ts:
  d=document();add_table(d,t);d.save(OUT/'tables'/f"Table_{t['number']}.docx")
 d=document();d.add_heading('Main manuscript tables',0)
 for i,t in enumerate(ts):
  if i:d.add_page_break()
  add_table(d,t)
 d.save(OUT/'JDR_Tables.docx')
 d=document();add_title_page(d);d.save(OUT/'JDR_Title_Page.docx')
 for blind in [False,True]:manuscript(ts,blind)
 d=document();d.add_heading('Figure captions and alternative text',0)
 for c in CAPTIONS:
  d.add_heading(f"Figure {c['number']}. {c['title']}",1);d.add_paragraph(resolve(c['caption']));d.add_paragraph('Alternative text: '+c['alt'])
 d.save(OUT/'Figure_Captions.docx')
 reporting_checklist();supplementary_index();missing_reference_metadata()
 notes=[
 'JDR author review package — 3 October 2026',
 'Primary title: '+TITLE,
 'Alternative title: Parkinson’s Disease, Population Ageing and Disability Care in Saudi Arabia: Evaluating Sex-Specific Forecasts Across the Gulf.',
 'Journal website instructions checked: https://jdr.kscdr.org/about/?tab=manuscript . These specify Word, a maximum 250-word abstract, 5–8 keywords, APA 7 references and double-blind review. The ScienceOpen page differs (Harvard in-text/PLOS entries and PDF). This package follows the journal website and supplies a blinded copy and separate title page. The example guides layout and depth, not its biological methods or data.',
 'Eight figures are embedded from the 600 dpi TIFF masters at their relevant Methods/Results locations, with complete captions and alternative text. Figure titles, panel labels, axes, tick values, legends and annotations use bold lettering to match the supplied example manuscript. Word automatic picture compression is disabled. Six editable Word tables appear beside their first substantive discussion, with greyscale headers, repeated header rows and no vertical rules. Separate 600 dpi TIFF, 300 dpi PNG and vector PDF/SVG figure files remain available.',
 'Figures 6–8 use at least 16-point lettering in the 12.6-inch masters, equivalent to 8.32 points at the 6.55-inch manuscript width. Legends, endpoint labels and notes are reflowed for readability; all annual forecast values remain plotted.',
 'Each main figure caption contains no more than 100 words, including its figure number and title. Panel definitions, interpretation, source links through Table 1 and supplementary-result links are retained. The main reference list contains only cited entries; additional reporting-guideline and software references remain in the supplementary crosswalk.',
 'Each main table caption contains no more than 100 words, counting the table number, title, explanatory note and expanded citations together. Definitions, numerical interpretation and source links are retained. Full source citations for Table 1 appear in the Methods, References and supplementary source crosswalk.',
 'The main body has '+str(len(BODY.split()))+' words before resolved citations; the abstract has '+str(len(ABSTRACT.split()))+' words. All '+str(len(REFS))+' references are author–date citations with APA-style entries and primary-source metadata.',
 'The new analysis is a post-primary arithmetic decomposition of unchanged forecasts, with no model refitting. The Discussion Limitations subsection is limited to 125 words. The complete original six-paragraph account and its citation are preserved in Limitations.docx, indexed in Supplementary Material S5. Negative results remain visible in the abstract and Results as scientific findings.',
 'The user confirmed use of the example authors/affiliations and then the funding and competing-interest declarations. No institutional approval/exemption identifier was supplied, and none is invented. Authors should verify the final ethics wording and specify their study-specific contributions before submission.',
 'The package is prepared for author review, not uploaded to the journal. The reporting checklist is not an actual SciScore run or score. The existing IJE package and supplied examples are preserved.',
 'The provider citation audit adds the IHME-required source identifier, provider-specific attribution and transformation wording, and corrections to the hierarchy and UN report citations. The current IHME FAQ explicitly confirms 2025 as the GBD 2023 citation year, resolving the prior citation-year query. Some Saudi export-level metadata remain incomplete. See references/Data_Citation_Audit.docx.',
 'IHME’s current non-commercial user agreement permits publication of analytical results but restricts providing third-party downloads of its source datasets from user-hosted facilities without written permission. The supplementary index is an inventory, not redistribution clearance. Before deposition, supply IHME source links in place of restricted raw downloads unless written permission covers redistribution. The required phrase Used with permission follows the public agreement and does not document a separate correspondence-based permission.',
 'Supplementary_Index.docx and supplementary_index.csv extend the existing archive mapping. They retain original relative folders. The full raw-data archive is not duplicated into this package; data access/reuse conditions continue to apply.',
 'Blinded copy: author names, affiliations, funding and identifying public-code URL are withheld. Before uploading the supplementary archive for double-blind review, inspect file metadata and repository links in those separate materials as well.',
 ]
 (OUT/'Submission_Notes.txt').write_text('\n\n'.join(notes)+'\n')
 d=document();d.add_heading('Submission and author review notes',0)
 for p in notes[1:]:d.add_paragraph(p)
 d.save(OUT/'Submission_Notes.docx')
 for p in protected:assert sha(p)==before[str(p.relative_to(ROOT))]
 for path,hashval in SOURCES.items():assert sha(ROOT/path)==hashval
 # Read-back checks include object counts and author-identifying strings.
 full=Document(OUT/'JDR_Manuscript.docx');blind=Document(OUT/'JDR_Manuscript_Blinded.docx')
 assert len(full.tables)==len(blind.tables)==6
 assert len(full.inline_shapes)==len(blind.inline_shapes)==8
 assert all(s._inline.docPr.get('descr') for s in full.inline_shapes)
 blindtext='\n'.join(p.text for p in blind.paragraphs)
 assert not any(s in blindtext for s in ['Saif Khan','Mahvish Khan','Mohtashim Lohani','saifkhanbio','mk.khan@','KSRG-2026'])
 assert '[@' not in blindtext and len(ABSTRACT.split())<=250 and 5<=len(KEYWORDS)<=8
 assert len(re.findall(r'^## Limitations$',BODY,re.M))==1
 assert all(f'Figure {i}' in BODY for i in range(1,9))
 assert all(f'Table {i}' in BODY for i in range(1,7))
 assert not any(s in blindtext for s in ['/home/','/mnt/','C:\\'])
 result={'status':'passed','abstract_words':len(ABSTRACT.split()),'body_words_before_citation_expansion':len(BODY.split()),
  'references':len(REFS),'main_figures':8,'main_tables':6,'figure_alt_text':8,'new_model_fits':0,
  'figure_caption_words_including_title':caption_counts,'figure_caption_max_words':100,
  'table_caption_words_including_title_and_notes':table_caption_counts,'table_caption_max_words':100,
  'limitations_words':limitations_words,'limitations_words_including_heading':limitations_words+1,'limitations_max_words':125,
  'full_limitations_words':len(resolve(FULL_LIMITATIONS).split()),'full_limitations_paragraphs':len(FULL_LIMITATIONS.split('\n\n')),
  'display_placement':'inline at relevant Methods/Results paragraphs','embedded_figure_format':'original 600 dpi TIFF',
  'source_sha256':SOURCES,'original_manuscripts_and_examples_unchanged':True,
  'protected_sha256':before,'blinded_identity_check':'passed','actual_sciscore_run':False,
  'pdf_rendering':'pending separate rendering and inspection'}
 (OUT/'validation.json').write_text(json.dumps(result,indent=2)+'\n')
 print(json.dumps({k:v for k,v in result.items() if 'sha256' not in k},indent=2))


if __name__=='__main__':main()
