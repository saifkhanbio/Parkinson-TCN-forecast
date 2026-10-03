"""Build review documents, provenance crosswalks and independent score checks."""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import zipfile

import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
EVIDENCE=ROOT/'supporting_data/2026-10-01_saudi_evidence'
REPORT=ROOT/'reports/bayesian_age_time_v1'
POP=ROOT/'reports/population_denominator_audit_v2'


def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def inline(paragraph,text):
    for part in re.split(r'(\[[^\]]+\]\(https?://[^)]+\))',text):
        match=re.fullmatch(r'\[([^\]]+)\]\((https?://[^)]+)\)',part)
        if match:
            hyperlink=OxmlElement('w:hyperlink');hyperlink.set(qn('r:id'),paragraph.part.relate_to(match[2],RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');properties=OxmlElement('w:rPr');style=OxmlElement('w:rStyle');style.set(qn('w:val'),'Hyperlink');properties.append(style);run.append(properties)
            node=OxmlElement('w:t');node.text=match[1];run.append(node);hyperlink.append(run);paragraph._p.append(hyperlink)
        else:
            for i,piece in enumerate(part.split('**')):
                run=paragraph.add_run(piece.replace('`',''));run.bold=bool(i%2)


def append_markdown(doc,text):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x) for x in row):rows.append(row)
                i+=1
            table=doc.add_table(rows=1,cols=len(rows[0]));table.style='Light Shading Accent 1'
            for j,value in enumerate(rows[0]):inline(table.rows[0].cells[j].paragraphs[0],value)
            header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
            for row in rows[1:]:
                cells=table.add_row().cells
                for j,value in enumerate(row):inline(cells[j].paragraphs[0],value)
            continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));inline(doc.add_heading(level=min(level,3)),line[level:].strip())
        else:
            style='List Bullet' if line.startswith('- ') else None
            inline(doc.add_paragraph(style=style),line[2:] if style else line)
        i+=1


def document():
    d=Document();d.sections[0].top_margin=Inches(.7);d.sections[0].bottom_margin=Inches(.7)
    d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10)
    return d


def main():
    final=EVIDENCE/'package_validation.json'
    if final.exists():raise FileExistsError('Existing completed package is immutable')
    source_rows=[
        ('WHO_MDB','World Health Organization. WHO Mortality Database. Public files updated February 2026; accessed 1 October 2026.','https://www.who.int/data/data-collection-tools/who-mortality-database','public_downloaded','Registered Saudi G20 deaths confirmed; ascertainment and GBD input overlap unresolved'),
        ('MOH_ACCESS','Saudi Ministry of Health. Health Data Platform: Health Data Sharing Procedures. Accessed 1 October 2026.','https://hdp.moh.gov.sa/en/data-sharing','request_route_verified','PD-specific extract not confirmed'),
        ('CHI_ACCESS','Council of Health Insurance. Data Sharing Request. Accessed 1 October 2026.','https://www.chi.gov.sa/en/open-data/Pages/access-data.aspx','request_route_verified','Claims data and approval require custodian confirmation'),
        ('NHRSP_ACCESS','Saudi Health Council. National Health Research and Studies Portal: access workflow and user guide. Accessed 1 October 2026.','https://shc.gov.sa/ar/EServices/Pages/nhrsp.aspx','login_route_verified','PD project not verified behind login'),
        ('GASTAT_ACCESS','General Authority for Statistics. Microdata request and user guide. Accessed 1 October 2026.','https://www.stats.gov.sa/en/request-for-scientific-use-files','request_route_verified','PD survey item unverified; population-estimates product says microdata unavailable'),
        ('GASTAT_METHODS','General Authority for Statistics. Methodology and Quality Report for Population Projections and Estimates Statistics. Accessed 1 October 2026.','https://www.stats.gov.sa/en/w/methodology-and-quality-report-for-population-projections-and-estimates-statistics','public_downloaded','Population definitions; not disease outcomes'),
        ('ALMUBARAK_2015','Al-Mubarak BR, Bohlega SA, Alkhairallah TS, et al. Parkinson’s Disease in Saudi Patients: A Genetic Study. PLoS ONE. 2015;10:e0135950. doi:10.1371/journal.pone.0135950.','https://doi.org/10.1371/journal.pone.0135950','public_supplements_downloaded','Selected clinical/genetic sample; no population denominator'),
        ('MENASA_2025','Khalil H, Shraim M, Jaradat B, et al. Parkinson’s Disease Database in the Middle East, North Africa, and South Asia Countries. Int J Public Health. 2025;70:1608016. doi:10.3389/ijph.2025.1608016.','https://doi.org/10.3389/ijph.2025.1608016','protocol_public_patient_data_unverified','Two named Saudi sites; consortium request needed'),
        ('IQVIA_ABSTRACT','Orayj K. Effect of COVID-19 Pandemic on Utilization of Parkinson’s Medications in Saudi Arabia: A Repeated Cross-Sectional Study [abstract]. Mov Disord. 2025;40(suppl 1).','https://www.mdsabstracts.org/abstract/effect-of-covid-19-pandemic-on-utilization-of-parkinsons-medications-in-saudi-arabia-a-repeated-cross-sectional-study/','abstract_public_underlying_data_not_public','Saudi IQVIA sales 2019–2023; commercial access terms unverified'),
        ('KFMC_PROJECT','Saudi Ministry of Health. Research report 2016: Parkinson Disease and Movement Disorder Registry in KFMC project listing.','https://www.moh.gov.sa/Ministry/MediaCenter/Publications/Pages/MOH-2016-VIEW.pdf','historical_project_document_downloaded','Current registry status unverified'),
        ('CHI_GUIDANCE','Council of Health Insurance. Parkinson Disease: indication update. Accessed 1 October 2026.','https://www.chi.gov.sa/Style%20Library/IDF_Branding/Indication/155%20-%20Parkinson%20Disease-Indication%20Update.pdf','public_document_downloaded','Clinical guidance; not patient data'),
    ]
    pd.DataFrame(source_rows,columns=['reference_id','citation','canonical_url','availability','qualification']).to_csv(EVIDENCE/'source_register.csv',index=False)
    manifests=[]
    for p in sorted(EVIDENCE.glob('manifest_*.json')):manifests.extend(json.loads(p.read_text()))
    crosswalk=[]
    for r in manifests:
        if r['status']!='downloaded':continue
        ident=r['id']
        if ident.startswith('who_'):ref='WHO_MDB'
        elif ident.startswith('saudi_genetics'):ref='ALMUBARAK_2015'
        elif ident.startswith('gastat_population'):ref='GASTAT_METHODS'
        elif ident.startswith('gastat'):ref='GASTAT_ACCESS'
        elif ident.startswith('nhrsp'):ref='NHRSP_ACCESS'
        elif ident.startswith('menasa'):ref='MENASA_2025'
        elif ident.startswith('iqvia'):ref='IQVIA_ABSTRACT'
        elif ident.startswith('kfmc'):ref='KFMC_PROJECT'
        elif ident=='chi_pd_guidance':ref='CHI_GUIDANCE'
        elif ident.startswith('chi'):ref='CHI_ACCESS'
        else:ref='MOH_ACCESS'
        assert sha(EVIDENCE/r['path'])==r['sha256']
        crosswalk.append(dict(file=r['path'],reference_id=ref,url=r['url'],sha256=r['sha256'],bytes=r['bytes']))
    pd.DataFrame(crosswalk).to_csv(EVIDENCE/'file_to_citation.csv',index=False)
    queries=[
        'Saudi Arabia Parkinson disease registry data access MENASA prevalence community survey',
        'site.moh.gov.sa Parkinson open data research data request',
        'Saudi health research data access NPHIES SHC Parkinson registry',
        '"Saudi" "Parkinson" "data availability" dataset',
        '"Saudi" "Parkinson" "dataset" "figshare"',
        '"Saudi" "Parkinson" "data" "zenodo"',
        '"Saudi" "Parkinson" "Mendeley"',
        'site.chi.gov.sa "data" "request" research',
        'site.stat.gov.sa "microdata" "research" access health',
        '"Saudi Arabia" "WHO Mortality Database" data availability',
        'Parkinson Saudi Arabia IQVIA 2019 2023 medication utilization pandemic study 2025',
        'site.open.data.gov.sa "Parkinson"', 'site.open.data.gov.sa "باركنسون"',
        'site.hdp.moh.gov.sa Parkinson','site.sfda.gov.sa "Parkinson" "data"',
    ]
    pd.DataFrame([dict(date='2026-10-01',query=q,role='targeted_discovery_not_systematic_review') for q in queries]).to_csv(EVIDENCE/'search_log.csv',index=False)
    # Record population age formats without using WHO to compute disease rates.
    pop=pd.read_csv(EVIDENCE/'processed/who_saudi_population_raw_extract.csv',dtype={'Frmat':str})
    pop=pop[['Year','Sex','Frmat','Admin1','SubDiv']].copy()
    pop['top_age_group']=pop.Frmat.map({'00':'95+','01':'85+','02':'85+','04':'75+'})
    assert pop.top_age_group.notna().all()
    pop['used_to_compute_mortality_rates']=False
    pop.to_csv(EVIDENCE/'processed/who_population_age_format_audit.csv',index=False)
    # Replay key scientific scores without calling the reporting implementation.
    point=pd.read_csv(REPORT/'point_scores.csv')
    independent=np.abs(np.log(point.prediction.to_numpy()/point.observed_rate.to_numpy()))
    np.testing.assert_allclose(independent,point.absolute_log_error,rtol=1e-10,atol=1e-14)
    new=pd.read_csv(REPORT/'interval_scores.csv.gz')
    assert np.array_equal(new.covered.to_numpy(),((new.observed>=new.lower)&(new.observed<=new.upper)).to_numpy())
    keys=['target','outcome','origin','sex','age','horizon','family','scale','level']
    old=pd.read_csv(ROOT/'results/primary_v1/interval_scores.csv')
    old=old.loc[old.family.isin(['tcn_adapted','local_champion','nonneural_champion'])]
    check=new.loc[new.target.eq('Saudi Arabia')&new.outcome.eq('prevalence')&new.family.isin(old.family.unique())]
    a=old.set_index(keys).sort_index();b=check.set_index(keys).sort_index()
    assert a.index.equals(b.index)
    np.testing.assert_allclose(a[['lower','median','upper','width','covered']].to_numpy(dtype=float),b[['lower','median','upper','width','covered']].to_numpy(dtype=float),rtol=1e-12,atol=1e-12)
    ww=pd.read_csv(REPORT/'wis_scores.csv');ww=ww.loc[ww.target.eq('Saudi Arabia')&ww.outcome.eq('prevalence')&ww.family.isin(old.family.unique())]
    ow=pd.read_csv(ROOT/'results/primary_v1/wis_scores.csv');ow=ow.loc[ow.family.isin(old.family.unique())]
    wkeys=[x for x in keys if x!='level']
    np.testing.assert_allclose(ww.set_index(wkeys).sort_index().wis_50_80,ow.set_index(wkeys).sort_index().wis_50_80,rtol=1e-12,atol=1e-12)
    lock=json.loads((ROOT/'study_design/bounded_extension_2026-10-01_lock.json').read_text())
    for file,digest in lock['sha256'].items():assert sha(ROOT/file)==digest
    # Documents preserve the exploratory status and distinguish evidence from access leads.
    doc=document();append_markdown(doc,(EVIDENCE/'README.md').read_text());doc.save(EVIDENCE/'saudi_evidence_and_access_review.docx')
    doc=document();doc.add_heading('Saudi Parkinson’s data access request — unsent draft',0)
    for text in (EVIDENCE/'access_request_draft.txt').read_text().split('\n\n'):
        doc.add_paragraph(text)
    doc.save(EVIDENCE/'access_request_draft.docx')
    doc=document();append_markdown(doc,(REPORT/'interpretation.md').read_text());doc.add_page_break();append_markdown(doc,(POP/'interpretation.md').read_text())
    doc.save(REPORT/'bounded_improvements_addendum.docx')
    docs=[EVIDENCE/'saudi_evidence_and_access_review.docx',EVIDENCE/'access_request_draft.docx',REPORT/'bounded_improvements_addendum.docx']
    for path in docs:
        with zipfile.ZipFile(path) as archive:assert archive.testzip() is None
        reopened=Document(path);assert len(reopened.paragraphs)>=10
        assert '/home/' not in '\n'.join(p.text for p in reopened.paragraphs)
    request=(EVIDENCE/'access_request_draft.txt').read_text().split('Short CHI description (within 500 characters):\n')[1].split('\n\n')[0]
    assert len(request)<=500
    state=dict(status='complete',checked_utc=datetime.now(timezone.utc).isoformat(),downloaded_files=len(crosswalk),
        downloaded_bytes=sum(x['bytes'] for x in crosswalk),all_download_hashes_verified=True,
        independent_point_error_replay=True,original_saudi_interval_and_wis_replay=True,
        original_primary_data_and_protocol_unchanged=True,chi_request_characters=len(request),
        requests_submitted=0,external_people_contacted=0,who_data_fitted_or_scored=False,
        docx_xml_and_readback_passed=True,word_page_rendering_checked=False,
        output_sha256={str(p.relative_to(ROOT)):sha(p) for p in docs})
    final.write_text(json.dumps(state,indent=2)+'\n')
    (REPORT/'independent_validation.json').write_text(json.dumps(state,indent=2)+'\n')
    print(json.dumps({k:state[k] for k in ['status','downloaded_files','downloaded_bytes','independent_point_error_replay','original_saudi_interval_and_wis_replay','chi_request_characters']}))


if __name__=='__main__':main()
