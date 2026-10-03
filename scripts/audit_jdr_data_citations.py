"""Audit analytical and supplementary data citations against provider evidence."""
import csv
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'manuscript/JDR_GBD_PARK/references'
EVIDENCE = OUT / 'provider_citation_audit'
OLD = ROOT / 'study_design/dataset_citations_2026-10-01'
MAPPING = {
    'D1': 'gbd_results', 'gbd_nonfatal': 'gbd_nonfatal', 'gbd_fatal': 'gbd_fatal',
    'gbd_demography': 'gbd_demography', 'hierarchy': 'hierarchy', 'D2': 'wpp',
    'wpp_methods': 'wpp_methods', 'D3': 'gccstat', 'D4': 'gastat',
    'gastat_methods': 'gastat_methods', 'D5': 'moh_yearbook',
    'gastat_series': 'gastat_series', 'moh_population': 'moh_population',
    'chi_data': 'chi',
}


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as handle:
        return list(csv.DictReader(handle))


def write_csv(path, rows):
    with path.open('w', newline='', encoding='utf-8-sig') as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def document(title):
    doc = Document()
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(11)
    doc.sections[0].left_margin = doc.sections[0].right_margin = Inches(.8)
    doc.add_heading(title, 0)
    return doc


def main():
    refs = json.loads((OUT / 'reference_register.json').read_text())
    body = (OUT.parent / 'authoring/body.txt').read_text()
    source_register = read_csv(OLD / 'source_register.csv')
    crosswalk = read_csv(OLD / 'file_to_citation_crosswalk.csv')
    source_ids = {row['id'] for row in source_register}
    for row in crosswalk:
        assert set(row['reference_ids'].split(';')) <= source_ids
    assert all(key in refs and '@' + key in body for key in MAPPING.values())
    assert refs['gbd_results']['year'] == '2025'
    assert refs['hierarchy']['authors'][0]['literal'] == 'Global Burden of Disease Collaborative Network'
    assert 'NO. 10' in refs['wpp_methods']['version']
    assert 'GBD 2023' in (EVIDENCE / 'ihme_faq.txt').read_text()
    assert 'IHME), 2025.' in (EVIDENCE / 'ihme_faq.txt').read_text()
    policy = (EVIDENCE / 'ihme_agreement.txt').read_text()
    assert 'Source: Institute for Health Metrics and Evaluation. Used with permission. All rights reserved' in policy

    audit_rows = []
    for row in source_register:
        key = MAPPING.get(row['id'])
        row['citation'] = refs[key]['apa_text'] if key else row['citation']
        if key:
            row['year'] = refs[key]['year']
            row['verification'] = refs[key]['provenance']
        if row['id'] == 'D1':
            row['note'] = 'Current IHME FAQ supplies citation year 2025. Export template is preserved as provenance, not used to override current provider citation guidance. Original extraction metadata remain incomplete.'
        if row['id'] == 'hierarchy':
            row['note'] = 'Explicit provider release-sheet suggested citation names the Global Burden of Disease Collaborative Network; applied in preference to the catalog creator field.'
        audit_rows.append({
            'source_id': row['id'], 'source': row['title'], 'role': row['role'],
            'citation_location': 'Main Methods, reference list and supplementary bibliography' if key else 'Supplementary Material S1 provider bibliography',
            'main_reference_key': key or '', 'files_mapped': sum(row['id'] in x['reference_ids'].split(';') for x in crosswalk),
            'status': 'Provider credited; export metadata qualifications recorded separately' if key else 'Supplementary/context/catalog source credited; not asserted to be a model input',
            'provider_url': row['url'], 'citation': row['citation'],
        })
    write_csv(OUT / 'data_citation_audit.csv', audit_rows)
    write_csv(OUT / 'data_source_register.csv', source_register)
    for row in crosswalk:
        row['citation_review_update'] = '2026-10-03; use the current JDR data_source_register.csv for corrected provider citation wording'
    write_csv(OUT / 'data_file_citation_crosswalk.csv', crosswalk)

    pending = read_csv(ROOT / 'manuscript/ije_draft_v2_raw_integration/missing_references.csv')
    pending = [row for row in pending if row['id'] != 'R1']
    for row in pending:
        row['action'] = row['action'].replace('IJE', 'the journal')
    write_csv(EVIDENCE / 'unresolved_metadata.csv', pending)

    provider_rows = [
        ('IHME/GBD results', 'Required IHME source identifier, bibliographic citation and no implied endorsement.',
         'Added exact source identifier to both full and blinded manuscripts; added author-derived-result wording. Current provider FAQ resolves the citation year to 2025.',
         'https://www.healthdata.org/data-tools-practices/data-practices/ihme-free-charge-non-commercial-user-agreement'),
        ('IHME hierarchy', 'Use the explicit suggested citation supplied with the dataset.',
         'Corrected corporate author to Global Burden of Disease Collaborative Network; retained 2025 and DOI 10.6069/kmah-et96.',
         'https://doi.org/10.6069/kmah-et96'),
        ('United Nations WPP 2024', 'Credit UN DESA Population Division, edition and licensed source; identify derived work.',
         'Online-edition citation verified against workbook. Added methodology report number and advance-unedited-version statement; added CC BY 3.0 IGO credit.',
         'https://population.un.org/wpp/assets/Files/WPP2024_Methodology-Report_Final.pdf'),
        ('GASTAT', 'Acknowledge GASTAT as official source and identify modifications.',
         'Cited population tables, methodology and the 2010–2024 duplicate-check export. National projections and harmonisations explicitly attributed to the authors.',
         'https://stats.gov.sa/en/use-policy'),
        ('GCC-Stat', 'Identify GCC-Statistical Centre, Sultanate of Oman, link its website, and mark edited content.',
         'Added the stipulated source identity and geographic designation; identified selection, aggregation and harmonisation by the authors.',
         'https://dp.marsa.gccstat.org/terms-use'),
        ('Saudi MOH', 'Name the MOH Portal and provide a source link; preserve origin and history.',
         'Cited both the population export and Yearbook 2023, Table 1-3; credited GASTAT as the original source of the reproduced population table.',
         'https://www.moh.gov.sa/en/Ministry/OpenData/Pages/OpenDataUsagePolicy.aspx'),
        ('Council of Health Insurance', 'Acknowledge CHI as the source and preserve the context of the published categories.',
         'Provider credited in Methods, Table 1, references and data acknowledgement. Original age and subscriber/insured definitions retained. Exact quarterly resource URLs remain incomplete.',
         'https://www.chi.gov.sa/en/open-data/Pages/requirements.aspx'),
        ('Supplementary-only sources', 'Distinguish materials inspected or acquired from inputs actually used in fitted models and reported comparisons.',
         'WHO mortality, CHI class/NPHIES, NHRSP HISS/ABSHER, additional MOH yearbook and registry/clinical documentation remain separately credited in the complete supplementary provider bibliography.',
         'Supplementary Material S1; complete file-to-citation crosswalk'),
    ]
    report = document('Data and database citation audit')
    report.add_paragraph('Review completed 3 October 2026. This audit checks source credit and citation wording for the JDR manuscript, not forecast validity or blanket permission to redistribute third-party data.')
    report.add_heading('Conclusion', 1)
    report.add_paragraph('All data providers contributing to the reported analyses are explicitly credited in the revised main manuscript. Supplementary acquisition, eligibility and contextual resources have separate references in the complete provider bibliography. Required attribution omissions and verified citation errors have been corrected. Exact metadata for some user-supplied exports remain incomplete and are listed separately; citation coverage must not be confused with complete acquisition provenance.')
    report.add_heading('Provider checks and changes', 1)
    for name, requirement, action, url in provider_rows:
        report.add_heading(name, 2)
        report.add_paragraph('Provider requirement or citation guidance: ' + requirement)
        report.add_paragraph('Manuscript action: ' + action)
        report.add_paragraph('Evidence: ' + url)
    report.add_heading('Current IHME citation guidance', 1)
    report.add_paragraph('The directly retrieved current IHME GBD FAQ explicitly gives 2025 for GBD 2023. The retained download citation template is preserved unchanged as evidence; the main reference now follows current provider guidance. This resolves the previous R1 bibliographic query. Source: https://www.healthdata.org/gbd/faq .')
    report.add_paragraph('The hierarchy release sheet explicitly names the Global Burden of Disease Collaborative Network in its suggested citation. That instruction takes precedence over the differently populated DataCite creator field. The original release sheet is retained in the source documentation.')
    report.add_heading('Source coverage', 1)
    report.add_paragraph(f'The main manuscript cites {len(MAPPING)} data or source-method references. The supplementary source register contains {len(source_register)} records and the crosswalk maps {len(crosswalk)} retained files. Every crosswalk reference identifier resolves to a source record. World Bank classifications came through the IHME hierarchy; no independent World Bank outcome dataset or unacquired SDI dataset is represented as a model input.')
    report.add_heading('Remaining metadata qualifications', 1)
    for row in pending:
        report.add_paragraph(row['id'] + ': ' + row['item'] + '. ' + row['unresolved_detail'] + ' ' + row['action'])
    report.add_heading('Data redistribution is a separate check', 1)
    report.add_paragraph('The current IHME agreement permits publication of analyses but restricts offering third-party downloads of IHME datasets through user-hosted facilities without written permission. The source inventory is not a redistribution licence. Use provider download links for restricted source files unless written permission covers redistribution. The required source-identifier wording follows the public agreement; it does not establish a separate individually issued permission or IHME endorsement.')
    report.add_heading('Verification boundaries', 1)
    report.add_paragraph('Provider pages and retained original citation documents were inspected. No provider was contacted and no acceptance letter was obtained. The WPP TermsOfUse URL returned 404; the citation and licence evidence instead comes from the retained official UN workbook and methodology report. Unspecified publication dates remain n.d.; no dates, dataset DOIs or access approvals were invented.')
    report.save(OUT / 'Data_Citation_Audit.docx')

    bibliography = document('Supplementary data-source references and analytical roles')
    bibliography.add_paragraph('Supplementary Material S1. This 3 October 2026 revision supersedes citation wording in earlier internal source lists. Main Methods citations identify analytical inputs; the entries below also credit resources used only for eligibility review, acquisition or context. They do not imply that every listed resource trained or validated a model.')
    for row in source_register:
        bibliography.add_heading(row['id'] + ' — ' + row['title'], 2)
        bibliography.add_paragraph(row['citation'])
        bibliography.add_paragraph('Role: ' + row['role'] + '. ' + row['note'])
    bibliography.save(OUT / 'Supplementary_Data_References.docx')
    summary = {'date': '2026-10-03', 'main_source_references': len(MAPPING),
               'supplementary_source_records': len(source_register), 'crosswalk_files': len(crosswalk),
               'orphan_crosswalk_reference_ids': [], 'gbd_citation_year_resolved': True,
               'unresolved_metadata_items': [r['id'] for r in pending],
               'all_analytical_providers_credited': True,
               'all_export_metadata_complete': False, 'blanket_redistribution_clearance': False,
               'verification_basis': 'Provider policies, provider suggested citations, retained source files and manuscript citation checks'}
    (OUT / 'citation_audit_validation.json').write_text(json.dumps(summary, indent=2) + '\n')
    print(json.dumps(summary, indent=2))


if __name__ == '__main__':
    main()
