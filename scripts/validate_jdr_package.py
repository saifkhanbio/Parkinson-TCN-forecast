"""Final numerical, document, image and rendered-page checks for the JDR package."""
import csv
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import zipfile

import fitz
import numpy as np
import pandas as pd
from docx import Document
from docx.oxml.ns import qn
from PIL import Image,ImageDraw

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'manuscript/JDR_GBD_PARK'


def sha(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def normalise_rendered_text(text):
 # PDF extraction inserts a line break inside compounds wrapped after a hyphen.
 text=re.sub(r'(?<=\w)([-–])\s*\n\s*(?=\w)',r'\1',text)
 text=re.sub(r'(?<=\w)\s*\n\s*([-–])(?=\w)',r'\1',text)
 return ' '.join(text.split())


def section_paragraphs(doc,heading):
 paragraphs=doc.paragraphs
 start=next(i for i,p in enumerate(paragraphs) if p.text==heading)
 selected=[]
 for paragraph in paragraphs[start+1:]:
  if paragraph.style.name.lower().startswith('heading'):break
  if paragraph.text.strip():selected.append(paragraph.text)
 return selected


def rendered_body_text(pdf):
 # Omit running headers and footers when checking paragraphs across page breaks.
 lines=[line for page in pdf for line in page.get_text().splitlines()
        if line.strip()!='Parkinson’s forecasting and disability care' and not re.fullmatch(r'Page\s+\d+',line.strip())]
 return normalise_rendered_text('\n'.join(lines))


def validate_table_captions(doc,numbers=range(1,7)):
 elements=list(doc.element.body)
 texts=[''.join(x.text or '' for x in node.xpath('.//w:t')) for node in elements]
 details={}
 for number in numbers:
  matches=[i for i,text in enumerate(texts) if text.startswith(f'Table {number}. ')]
  assert len(matches)==1,(number,'Missing or duplicated table caption')
  pos=matches[0]
  assert elements[pos+1].tag==qn('w:tbl')
  title,note=texts[pos],texts[pos+2]
  words=len((title+' '+note).split())
  assert words<=100,(number,'Table caption exceeds 100 words',words)
  details[str(number)]={'title':title,'note':note,'words':words}
 return details


def validate_inline_displays(doc):
 placements=json.loads((OUT/'authoring/inline_placements.json').read_text())
 elements=list(doc.element.body)
 texts=[''.join(x.text or '' for x in node.xpath('.//w:t')) for node in elements]
 reference_index=texts.index('References')
 placed=[]
 for placement in placements:
  matches=[i for i,t in enumerate(texts) if t.startswith(placement['after'])]
  assert len(matches)==1,(placement['after'],matches)
  pos=matches[0]+1
  for kind,number in placement['items']:
   assert texts[pos].startswith(f'{kind.capitalize()} {number}.'),(kind,number,texts[pos])
   assert pos<reference_index
   if kind=='table':assert elements[pos+1].tag==qn('w:tbl')
   else:assert elements[pos+1].xpath('.//w:drawing')
   assert len((texts[pos]+' '+texts[pos+2]).split())<=100,(kind,number,'Caption exceeds 100 words')
   placed.append((kind,number,pos))
   pos+=3
 assert len(placed)==14
 assert doc.settings.element.find(qn('w:doNotAutoCompressPictures')) is not None
 embedded=[]
 for number,shape in enumerate(doc.inline_shapes,1):
  rid=shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
  blob=doc.part.related_parts[rid].blob
  assert hashlib.sha256(blob).hexdigest()==sha(OUT/'figures'/f'figure_{number}.tiff')
  with Image.open(BytesIO(blob)) as im:
   assert im.format=='TIFF' and im.info['dpi']==(600,600)
   embedded.append({'figure':number,'pixels':list(im.size),'source_dpi':600,
                    'effective_dpi':round(im.width/shape.width.inches),'original_master_bytes_preserved':True})
 return {'all_displays_at_specified_paragraphs':True,'displays_before_references':True,
         'editable_tables':len(doc.tables),'embedded_figures':embedded}


def main():
 validation=json.loads((OUT/'validation.json').read_text())
 for path,expected in validation['protected_sha256'].items():assert sha(ROOT/path)==expected,path
 for path,expected in validation['source_sha256'].items():assert sha(ROOT/path)==expected,path
 body=(OUT/'authoring/body.txt').read_text()
 refs=json.loads((OUT/'references/manuscript_reference_register.json').read_text())
 assert set(re.findall(r'@([a-z_]+)',body))==set(refs)
 short_limitations=re.search(r'(?<=## Limitations\n\n)(.*?)(?=\n\n# Conclusions)',body,re.S).group(1)
 limitations_words=len(short_limitations.split())
 assert limitations_words+1<=125
 full_limitations=(OUT/'authoring/limitations_full.txt').read_text().strip()
 register=json.loads((OUT/'references/reference_register.json').read_text())
 def expand_citation(match):
  keys=match.group(1).replace('@','').split(';')
  keys.sort(key=lambda k:(register[k]['sort_author'].lower(),register[k]['display_year']))
  return '; '.join(register[key]['cite'] for key in keys)
 full_limitations=re.sub(r'\[@([a-z_]+(?:;@[a-z_]+)*)\]',expand_citation,full_limitations)
 full_doc=Document(OUT/'Limitations.docx')
 assert section_paragraphs(full_doc,'Limitations')==full_limitations.split('\n\n')
 assert '[@' not in '\n'.join(p.text for p in full_doc.paragraphs)
 with fitz.open(OUT/'Limitations.pdf') as full_pdf:
  full_rendered=rendered_body_text(full_pdf)
  assert all(normalise_rendered_text(paragraph) in full_rendered for paragraph in full_limitations.split('\n\n'))
 supplementary=pd.read_csv(OUT/'supplementary_index.csv')
 assert ((supplementary.archive_member=='manuscript/JDR_GBD_PARK/Limitations.docx') &
         (supplementary.supplementary_identifier=='Supplementary Material S5')).any()
 captions={c['number']:c for c in json.loads((OUT/'authoring/figure_captions.json').read_text())}
 caption_counts={str(n):len(f"Figure {n}. {c['title']} {c['caption']}".split()) for n,c in captions.items()}
 assert all(count<=100 for count in caption_counts.values()),caption_counts
 table_captions=validate_table_captions(Document(OUT/'JDR_Tables.docx'))
 table_caption_counts={key:record['words'] for key,record in table_captions.items()}
 for number in range(1,7):
  single=validate_table_captions(Document(OUT/'tables'/f'Table_{number}.docx'),[number])
  assert single[str(number)]==table_captions[str(number)]
 with fitz.open(OUT/'JDR_Tables.pdf') as table_pdf:
  assert len(table_pdf)==6
  for i,page in enumerate(table_pdf,1):
   text=normalise_rendered_text(page.get_text())
   assert table_captions[str(i)]['title'] in text
   assert normalise_rendered_text(table_captions[str(i)]['note']) in text
 assert len(re.findall(r'^## Limitations$',body,re.M))==1
 for i in range(1,9):assert f'Figure {i}' in body
 for i in range(1,7):assert f'Table {i}' in body
 abstract=(OUT/'authoring/abstract.txt').read_text()
 assert len(abstract.split())<=250
 # Independent check of all supplementary changes and demographic contributions.
 changes=pd.read_csv(ROOT/'forecast_percentage_change.csv')
 np.testing.assert_allclose(changes.percentage_change_from_2023,100*(changes.forecast_count/changes.baseline_count_2023-1),rtol=1e-12)
 decomposition=pd.read_csv(OUT/'analysis/forecast_growth_decomposition.csv')
 components=['population_size','composition','rates']
 np.testing.assert_allclose(decomposition[[c+'_count_contribution' for c in components]].sum(axis=1),decomposition.total_change,rtol=1e-11,atol=1e-8)
 np.testing.assert_allclose(decomposition[[c+'_percentage_point_contribution' for c in components]].sum(axis=1),decomposition.total_change_percent,rtol=1e-11,atol=1e-8)
 assert len(decomposition)==len(changes)==180
 # Compare main Table 5 to stored exact forecasts, independent of document builder.
 t5=pd.read_csv(OUT/'tables/table_5.csv')
 for _,r in t5.iterrows():
  part=changes.loc[changes.outcome.eq(r.Outcome.lower())&changes.sex.eq(r.Sex)&changes.population_scenario.eq('gbd_2023_aligned_un_growth')].sort_values('forecast_year')
  assert str(r['2023'])==f'{part.baseline_count_2023.iloc[0]:,.0f}'
  for year in range(2024,2029):assert str(r[str(year)])==f"{part.loc[part.forecast_year.eq(year),'forecast_count'].iloc[0]:,.0f}"
  np.testing.assert_allclose(float(r['Change (%)']),round(part.percentage_change_from_2023.iloc[-1],2),atol=1e-12)
 image_metadata={}
 typography=json.loads((OUT/'figure_typography.json').read_text())
 assert [item['figure'] for item in typography]==list(range(1,9))
 assert all(item['font_weights']==['bold'] and not item['text_outside_canvas'] for item in typography)
 for i in range(1,9):
  for ext in ['tiff','png','pdf','svg']:assert (OUT/'figures'/f'figure_{i}.{ext}').exists()
  with fitz.open(OUT/'figures'/f'figure_{i}.pdf') as figure_pdf:
   spans=[span for block in figure_pdf[0].get_text('dict')['blocks'] if block['type']==0
          for line in block['lines'] for span in line['spans'] if span['text'].strip()]
   assert spans,(i,'No figure text extracted')
   fonts=sorted({span['font'] for span in spans})
   assert all('bold' in font.lower() for font in fonts),(i,fonts)
   minimum_font=min(span['size'] for span in spans)
   minimum_embedded_font=minimum_font*6.55*72/figure_pdf[0].rect.width
   if i in [6,7,8]:assert minimum_embedded_font>=8,(i,'Figure lettering below 8 pt in manuscript',minimum_embedded_font)
  with Image.open(OUT/'figures'/f'figure_{i}.tiff') as im:
   assert all(abs(float(v)-600)<.1 for v in im.info['dpi'])
   assert im.width>=7000
   image_metadata[f'figure_{i}']={'pixels':list(im.size),'tiff_dpi':600,'vector_formats':['PDF','SVG'],
                                 'all_figure_text_bold':True,'pdf_fonts':fonts,
                                 'minimum_source_font_pt':round(minimum_font,2),
                                 'minimum_font_at_manuscript_width_pt':round(minimum_embedded_font,2)}
 rendered={}
 for stem in ['JDR_Manuscript','JDR_Manuscript_Blinded']:
  doc=Document(OUT/(stem+'.docx'));pdf=fitz.open(OUT/(stem+'.pdf'))
  assert len(doc.tables)==6 and len(doc.inline_shapes)==8
  assert all(s._inline.docPr.get('descr') for s in doc.inline_shapes)
  inline_validation=validate_inline_displays(doc)
  assert section_paragraphs(doc,'Limitations')==short_limitations.split('\n\n')
  assert len(('Limitations '+' '.join(section_paragraphs(doc,'Limitations'))).split())<=125
  assert normalise_rendered_text(short_limitations) in rendered_body_text(pdf)
  assert validate_table_captions(doc)==table_captions
  texts=[];image_pages=[];blank=[];overflow=[]
  for i,page in enumerate(pdf):
   text=page.get_text();texts.append(text)
   # LibreOffice shares image resources across all pages; count displayed images.
   displayed_images=page.get_image_info()
   if displayed_images:image_pages.append(i+1)
   if len(text.strip())<100 and not displayed_images:blank.append(i+1)
   for block in page.get_text('dict')['blocks']:
    if block['type']!=0:continue
    for line in block['lines']:
     for span in line['spans']:
      x0,y0,x1,y1=span['bbox']
      if span['text'].strip() and (x0 < -1 or y0 < -1 or x1>page.rect.width+1 or y1>page.rect.height+1):overflow.append((i+1,span['text']))
  assert not blank,(stem,blank)
  assert not overflow,(stem,overflow)
  assert len(image_pages)==8
  figure_caption_pages={}
  for number in range(1,9):
   matches=[i for i,t in enumerate(texts) if re.search(rf'Figure {number}\.\s',t)]
   assert len(matches)==1,(stem,number,matches)
   page=pdf[matches[0]]
   assert page.get_image_info(),(stem,number,'caption separated from figure')
   assert normalise_rendered_text(captions[number]['caption']) in normalise_rendered_text(page.get_text()),(stem,number,'caption split across pages')
   info=page.get_image_info()[0]
   assert info['width']>=7000,(stem,number,'PDF resolution reduced',info['width'])
   figure_caption_pages[str(number)]=matches[0]+1
  table_pages={}
  table_titles={t['number']:f"Table {t['number']}. {t['title']}" for t in json.loads((OUT/'authoring/tables.json').read_text())}
  for number in range(1,7):
   matches=[i for i,t in enumerate(texts) if table_titles[number] in ' '.join(t.split())]
   assert len(matches)==1,(stem,number,matches)
   assert normalise_rendered_text(table_captions[str(number)]['note']) in normalise_rendered_text(texts[matches[0]]),(stem,number,'Table note split across pages')
   table_pages[str(number)]=matches[0]+1
  reference_page=next(i for i,t in enumerate(texts) if re.search(r'^References\s*$',t,re.M))
  assert max(image_pages)<reference_page+1
  text='\n'.join(texts)
  assert not any(s in text for s in ['Use of artificial intelligence:','OpenAI Codex',
      'The public core source snapshot is cited separately','The reporting audit applies GATHER items',
      'Automated checks covered future-row perturbation invariance'])
  assert not any(s in text for s in ['[@','/home/','/mnt/'])
  assert 'Source: Institute for Health Metrics and Evaluation. Used with permission. All rights reserved.' in ' '.join(text.split())
  assert 'Sultanate of Oman' in text and 'CC BY 3.0 IGO' in text
  assert 'Data-source acknowledgements' in text
  if 'Blinded' in stem:assert not any(s in text for s in ['Saif Khan','Mahvish Khan','Mohtashim Lohani','saifkhanbio','KSRG-2026','mk.khan@'])
  assert '34.56' in text and '26.37' in text and '47.52' in text
  rendered[stem]={'pages':len(pdf),'blank_pages':blank,'off_page_text':overflow,'figure_pages':image_pages,
                   'figure_caption_pages':figure_caption_pages,'table_pages':table_pages,
                   'inline_display_validation':inline_validation,
                   'docx_sha256':sha(OUT/(stem+'.docx')),'pdf_sha256':sha(OUT/(stem+'.pdf'))}
 # Contact sheet for the final, corrected Word-to-PDF rendering.
 pdf=fitz.open(OUT/'JDR_Manuscript.pdf')
 figure_indices=[i for i,p in enumerate(pdf) if p.get_image_info()]
 table_indices=[i for i,p in enumerate(pdf) if any(f'Table {n}.' in p.get_text() for n in range(1,7)) and not p.get_image_info()]
 selected=list(dict.fromkeys([0,1,5,12]+table_indices+figure_indices))
 canvas=Image.new('RGB',(1920,695*((len(selected)+3)//4)),'#DDDDDD')
 for k,i in enumerate(selected):
  pix=pdf[i].get_pixmap(matrix=fitz.Matrix(.75,.75));im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples);im.thumbnail((460,650))
  tile=Image.new('RGB',(480,695),'white');tile.paste(im,((480-im.width)//2,25));ImageDraw.Draw(tile).text((8,5),f'Page {i+1}',fill='black');canvas.paste(tile,((k%4)*480,(k//4)*695))
 canvas.save(OUT/'manuscript_contact_sheet.png')
 validation.update(pdf_rendering='LibreOffice Word-to-PDF; rendered pages visually inspected',
   rendered_documents=rendered,images=image_metadata,independent_percentage_and_decomposition_checks=True,
   figure_caption_words_including_title=caption_counts,figure_caption_max_words=100,
   table_caption_words_including_title_and_notes=table_caption_counts,table_caption_max_words=100,
   standalone_table_captions_match_manuscripts=True,
   limitations_words=limitations_words,limitations_words_including_heading=limitations_words+1,
   limitations_max_words=125,full_limitations_words=len(full_limitations.split()),
   full_limitations_paragraphs=len(full_limitations.split('\n\n')),full_limitations_preserved_in_separate_document=True,
   main_table_5_reproduces_exact_forecasts=True,actual_sciscore_run=False,final_status='validated_for_author_review')
 validation['data_citation_audit']=json.loads((OUT/'references/citation_audit_validation.json').read_text())
 (OUT/'validation.json').write_text(json.dumps(validation,indent=2)+'\n')
 excluded={'Example_Manuscript_JDR.docx','Example_SciscoreReport.pdf'}
 outputs=[p for p in sorted(OUT.rglob('*')) if p.is_file() and p.name not in excluded
          and not p.name.endswith(':Zone.Identifier') and p.suffix!='.zip' and p.name!='output_manifest.csv']
 with (OUT/'output_manifest.csv').open('w',newline='') as h:
  w=csv.writer(h);w.writerow(['package_member','bytes','sha256'])
  for p in outputs:w.writerow([str(p.relative_to(OUT)),p.stat().st_size,sha(p)])
 package=OUT/'JDR_Author_Review_Package.zip'
 with zipfile.ZipFile(package,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
  for p in outputs+[OUT/'output_manifest.csv']:z.write(p,str(p.relative_to(OUT)))
  for name in ['jdr_additional_analysis.py','plot_jdr_figures.py','build_jdr_references.py','build_jdr_manuscript.py','jdr_reference_metadata.py','validate_jdr_package.py','acquire_jdr_citation_evidence.py','audit_jdr_data_citations.py']:
   z.write(ROOT/'scripts'/name,'reporting_scripts/'+name)
 with zipfile.ZipFile(package) as z:assert z.testzip() is None
 print(json.dumps({'status':validation['final_status'],'documents':rendered,'figures':8,'tables':6,
                   'figure_caption_words_including_title':caption_counts,
                   'table_caption_words_including_title_and_notes':table_caption_counts,
                   'limitations_words':limitations_words,'full_limitations_words':len(full_limitations.split()),
                   'references':len(refs),'abstract_words':len(abstract.split()),'package_MB':round(package.stat().st_size/1e6,2)},indent=2))


if __name__=='__main__':main()
