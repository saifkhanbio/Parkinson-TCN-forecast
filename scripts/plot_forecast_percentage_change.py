"""Report sex-specific changes from matched 2023 baselines, without refitting."""

from __future__ import annotations

import csv
import hashlib
import json
import os
from decimal import Decimal, localcontext
from pathlib import Path

os.environ.setdefault("MPLCONFIGDIR", "/tmp/gbd_forecast_change_mpl")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Pt
from matplotlib.lines import Line2D
from matplotlib.ticker import MultipleLocator, PercentFormatter
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/forecast_percentage_change_v1"
SUMMARY = ROOT / "reports/saudi_raw_integration_v1/conditional_projection_summary_2024_2028.csv"
BASELINE = ROOT / "reports/projections_v1/baseline_2023_burden.csv"
UN = ROOT / "data/processed/design_v1/un_population_1990_2028.csv"
NATIONAL = ROOT / "data/processed/saudi_raw_integration_v1/national_population_2022_2024.csv"
CELLS = ROOT / "reports/saudi_raw_integration_v1/conditional_projection_cells_2024_2028.csv"
PRIOR = ROOT / "reports/projections_v1/projection_summary_all_years.csv"
MAIN_SCENARIO = "gbd_2023_aligned_un_growth"
SEXES = ["Male", "Female", "Both"]
OUTCOMES = ["prevalence", "incidence"]
INK, MUTED, GRID = "#213547", "#536371", "#E3E9EE"
OUTCOME_STYLES = {
    "prevalence": ("Prevalent cases", "#C55300", "o", "-"),
    "incidence": ("Annual incident cases", "#0072B2", "s", (0, (4, 1.6))),
}
SCENARIOS = [
    (MAIN_SCENARIO, "GBD 2023 baseline + UN growth", "#0072B2", "-", "o"),
    ("un_medium_unaligned", "UN medium population", "#00866B", (0, (1, 1.5)), "s"),
    ("national_2024_un_growth", "National 2024 baseline + UN growth", "#C55300", (0, (4, 2)), "^"),
]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def manuscript_hashes():
    return {str(p.relative_to(ROOT)): sha(p)
            for p in sorted((ROOT / "manuscript").rglob("*")) if p.is_file()}


def read_strings(path):
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def population_path(year, scenario, allocation, un, national, gbd_base):
    """Reproduce the preserved scenario's age-sex populations at any year."""
    part = un.loc[un.year.eq(year)].merge(gbd_base, on=["sex", "age"], validate="one_to_one")
    ub = un.loc[un.year.eq(2023), ["sex", "age", "population_persons"]].rename(
        columns={"population_persons": "un_2023"})
    part = part.merge(ub, on=["sex", "age"], validate="one_to_one")
    part["gbd_aligned"] = part.implied_population * part.population_persons / part.un_2023
    if scenario == MAIN_SCENARIO:
        part["population"] = part.gbd_aligned
    elif scenario == "un_medium_unaligned":
        part["population"] = part.population_persons
    else:
        anchor = int(scenario.split("_")[1])
        ug = un.groupby(["year", "sex", "broad_age"]).population_persons.sum()
        ng = national.groupby(["year", "sex", "age"]).population.sum()
        part["broad_population"] = [
            ng.loc[(anchor, sex, age)] * ug.loc[(year, sex, age)] / ug.loc[(anchor, sex, age)]
            for sex, age in zip(part.sex, part.broad_age)]
        part["raw_weight"] = (part.gbd_aligned if allocation.startswith("gbd")
                              else part.population_persons)
        denominator = part.groupby(["sex", "broad_age"]).raw_weight.transform("sum")
        part["population"] = part.broad_population * part.raw_weight / denominator
        np.testing.assert_allclose(
            part.groupby(["sex", "broad_age"]).population.sum(),
            part.groupby(["sex", "broad_age"]).broad_population.first(), rtol=1e-12)
    assert len(part) == 22 and (part.population > 0).all()
    return part[["sex", "age", "population"]]


def make_data():
    forecasts = [r for r in read_strings(SUMMARY) if r["family"] == "tcn_adapted"]
    assert len(forecasts) == 180
    original = [r for r in read_strings(BASELINE)
                if r["target"] == "Saudi Arabia" and r["measure"] == "count"
                and r["age_group"] == "45+"]
    baseline = {(r["outcome"], r["sex"], r["scenario"], "native_age_detail"): r["value"]
                for r in original}
    un = pd.read_csv(UN)
    un = un.loc[un.location_name.eq("Saudi Arabia") & un.age_start.ge(45)].copy()
    un["broad_age"] = np.where(un.age_start.ge(80), "80+", un.age)
    national = pd.read_csv(NATIONAL)
    rates = {outcome: pd.read_csv(ROOT / f"results/projections_v1/trials/SAU_{outcome}/baseline_2023_age_rates.csv")
             for outcome in OUTCOMES}
    gbd_base = rates["prevalence"][["sex", "age", "implied_population"]]
    combinations = sorted({(r["scenario"], r["within_80plus_allocation"]) for r in forecasts})
    existing_cells = pd.read_csv(CELLS)
    existing_cells = existing_cells.loc[existing_cells.family.eq("tcn_adapted")]
    baseline_cells = []
    for scenario, allocation in combinations:
        # Every 2024-2028 age-sex population must reproduce the saved analyses.
        for year in range(2024, 2029):
            reconstructed = population_path(year, scenario, allocation, un, national, gbd_base)
            saved = existing_cells.loc[existing_cells.scenario.eq(scenario)
                                       & existing_cells.within_80plus_allocation.eq(allocation)
                                       & existing_cells.year.eq(year)]
            matched = saved.merge(reconstructed, on=["sex", "age"], validate="many_to_one",
                                  suffixes=("_saved", "_reconstructed"))
            assert len(matched) == 44
            np.testing.assert_allclose(matched.population_saved, matched.population_reconstructed,
                                       rtol=1e-12, atol=1e-8)
        pops = population_path(2023, scenario, allocation, un, national, gbd_base)
        for outcome in OUTCOMES:
            cells = rates[outcome][["sex", "age", "rate"]].merge(pops, on=["sex", "age"], validate="one_to_one")
            cells["count"] = cells.rate * cells.population / 1e5
            cells["outcome"], cells["scenario"], cells["within_80plus_allocation"] = outcome, scenario, allocation
            baseline_cells.append(cells)
            for sex in SEXES:
                value = float(cells["count"].sum() if sex == "Both"
                              else cells.loc[cells.sex.eq(sex), "count"].sum())
                key = (outcome, sex, scenario, allocation)
                if key in baseline:
                    np.testing.assert_allclose(value, float(baseline[key]), rtol=1e-12)
                else:
                    baseline[key] = repr(value)
    pd.concat(baseline_cells, ignore_index=True).to_csv(OUT / "scenario_baseline_2023_age_cells.csv", index=False)
    rows = []
    for r in forecasts:
        key = (r["outcome"], r["sex"], r["scenario"], r["within_80plus_allocation"])
        b, v = baseline[key], r["count_45plus"]
        with localcontext() as ctx:
            ctx.prec = 34
            delta = Decimal(v) - Decimal(b)
            change = 100 * delta / Decimal(b)
        rows.append({
            "geography": "Saudi Arabia", "outcome": r["outcome"], "sex": r["sex"],
            "age_scope": "45+", "model": "tcn_adapted", "baseline_year": 2023,
            "forecast_year": r["year"], "population_scenario": r["scenario"],
            "within_80plus_allocation": r["within_80plus_allocation"],
            "baseline_count_2023": b, "forecast_count": v, "absolute_change": str(delta),
            "percentage_change_from_2023": str(change),
            "baseline_basis": ("GBD 2023 age-specific rates applied to the same population scenario at 2023"),
            "comments": ("Cumulative change from fixed 2023 baseline; case counts, not age-standardised rates. "
                         + ("The national 2024 anchor is backcast to 2023 using UN growth; "
                            "this baseline is scenario-derived. " if "national_2024" in r["scenario"] else "")
                         + "Point estimates; interval reliability is not inferred from these changes."),
        })
    dest = ROOT / "forecast_percentage_change.csv"
    with dest.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    data = pd.read_csv(dest)
    keys = ["outcome", "sex", "forecast_year", "population_scenario", "within_80plus_allocation"]
    assert len(data) == 180 and not data.duplicated(keys).any()
    assert np.isfinite(data.percentage_change_from_2023).all()
    for _, part in data.groupby(["outcome", "forecast_year", "population_scenario", "within_80plus_allocation"]):
        p = part.set_index("sex")
        for field in ["baseline_count_2023", "forecast_count"]:
            np.testing.assert_allclose(p.loc["Both", field], p.loc[["Male", "Female"], field].sum(), rtol=1e-12)
        weighted = np.average(p.loc[["Male", "Female"], "percentage_change_from_2023"],
                              weights=p.loc[["Male", "Female"], "baseline_count_2023"])
        np.testing.assert_allclose(p.loc["Both", "percentage_change_from_2023"], weighted, rtol=1e-12)
    prior = pd.read_csv(PRIOR)
    prior = prior.loc[prior.target.eq("Saudi Arabia") & prior.family.eq("tcn_adapted")
                      & prior.measure.eq("count") & prior.age_group.eq("45+")]
    match = data.merge(prior, left_on=["outcome", "sex", "forecast_year", "population_scenario"],
                       right_on=["outcome", "sex", "forecast_year", "scenario"], validate="many_to_one")
    assert len(match) == 60
    np.testing.assert_allclose(match.percentage_change_from_2023,
                               match.change_percent_from_scenario_baseline, rtol=1e-11, atol=1e-10)
    return data


def setup_style():
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 11, "text.color": INK,
        "axes.labelcolor": INK, "xtick.color": MUTED, "ytick.color": MUTED,
        "axes.spines.top": False, "axes.spines.right": False, "axes.spines.left": False,
        "axes.edgecolor": "#A4AFB7", "axes.linewidth": .65,
        "axes.titlesize": 13, "axes.titleweight": "bold", "axes.labelsize": 11,
        "xtick.labelsize": 10, "ytick.labelsize": 10, "pdf.fonttype": 42,
        "svg.fonttype": "none", "savefig.facecolor": "white", "figure.facecolor": "white",
    })


def format_axis(ax, title, ylabel=None, top=40):
    ax.set_title(title, loc="left", pad=16)
    ax.set_axisbelow(True)
    ax.grid(axis="y", color=GRID, lw=.8)
    ax.axhline(0, color="#A4AFB7", lw=.8)
    ax.set_xlim(2022.8, 2029.35)
    ax.set_ylim(-1, top)
    ax.set_xticks(range(2023, 2029))
    ax.yaxis.set_major_locator(MultipleLocator(10))
    ax.yaxis.set_major_formatter(PercentFormatter(xmax=100, decimals=0))
    ax.tick_params(length=0, pad=7)
    ax.spines["bottom"].set_visible(False)
    if ylabel:
        ax.set_ylabel(ylabel, labelpad=12)


def export(fig, stem):
    for extension in ["pdf", "svg", "png", "tiff"]:
        kwargs = {"dpi": 600 if extension == "tiff" else 300}
        if extension == "tiff":
            kwargs["pil_kwargs"] = {"compression": "tiff_lzw"}
        fig.savefig(OUT / f"{stem}.{extension}", **kwargs)
    fig.savefig(OUT / f"{stem}_preview.png", dpi=125)
    plt.close(fig)


def main_figure(data):
    primary = data.loc[data.population_scenario.eq(MAIN_SCENARIO)]
    fig = plt.figure(figsize=(12.6, 6.3))
    fig.text(.068, .953, "Parkinson’s disease burden: projected change from 2023",
             fontsize=18, weight="bold")
    fig.text(.068, .908, "Saudi Arabia · ages 45+ · adapted temporal convolutional network",
             color=MUTED, fontsize=11.5)
    handles = [Line2D([], [], color=c, marker=m, ls=ls, lw=2.2, ms=5, mfc="white", label=label)
               for label, c, m, ls in OUTCOME_STYLES.values()]
    fig.legend(handles=handles, loc="upper left", bbox_to_anchor=(.064, .866), ncol=2,
               frameon=False, handlelength=2.8, columnspacing=2.5)
    grid = fig.add_gridspec(1, 3, left=.068, right=.955, bottom=.265, top=.755, wspace=.20)
    for i, sex in enumerate(SEXES):
        ax = fig.add_subplot(grid[0, i])
        label = "Both sexes combined" if sex == "Both" else sex
        format_axis(ax, f"{'ABC'[i]}   {label}", "Change from 2023 (%)" if i == 0 else None)
        for outcome in OUTCOMES:
            part = primary.loc[primary.sex.eq(sex) & primary.outcome.eq(outcome)].sort_values("forecast_year")
            assert part.forecast_year.tolist() == list(range(2024, 2029))
            _, color, marker, ls = OUTCOME_STYLES[outcome]
            x = np.r_[2023, part.forecast_year.to_numpy()]
            y = np.r_[0., part.percentage_change_from_2023.to_numpy()]
            ax.plot(x, y, color=color, marker=marker, ls=ls, lw=2.2, ms=5.5,
                    mfc="white", mew=1.3, clip_on=False)
            ax.text(2028.22, y[-1], f"+{y[-1]:.2f}%", color=color, va="center", fontsize=11.4, weight="bold")
        ax.plot(2023, 0, "o", color=INK, mfc="white", ms=5, mew=1.1, clip_on=False)
        pos = ax.get_position()
        fig.text(pos.x0, .160, "2023 → 2028 · modelled case counts", fontsize=9.7, color=MUTED)
        for j, outcome in enumerate(OUTCOMES):
            r = primary.loc[primary.sex.eq(sex) & primary.outcome.eq(outcome) & primary.forecast_year.eq(2028)].iloc[0]
            label, color, _, _ = OUTCOME_STYLES[outcome]
            label = "Prevalence" if outcome == "prevalence" else "Incidence/year"
            fig.text(pos.x0, .120 - .035 * j,
                     f"{label}: {r.baseline_count_2023:,.0f} → {r.forecast_count:,.0f}",
                     color=color, fontsize=10.4)
    fig.text(.5, .211, "Year", ha="center", fontsize=11)
    fig.text(.068, .035, "GBD 2023 population baseline + UN growth. Point forecasts; count changes include population growth and ageing.",
             fontsize=9.2, color=MUTED)
    export(fig, "saudi_forecast_percentage_change")


def sensitivity_figure(data):
    fig = plt.figure(figsize=(12.6, 9.1))
    fig.text(.068, .963, "Sensitivity to population assumptions", fontsize=18, weight="bold")
    fig.text(.068, .926, "Saudi Arabia · ages 45+ · percentage change from each scenario’s matched 2023 baseline",
             color=MUTED, fontsize=11)
    handles = [Line2D([], [], color=c, marker=m, ls=ls, lw=1.8, ms=4.8, mfc="white", label=label)
               for _, label, c, ls, m in SCENARIOS]
    fig.legend(handles=handles, loc="upper left", bbox_to_anchor=(.064, .890), ncol=3,
               frameon=False, handlelength=2.5, columnspacing=1.6, fontsize=9.8)
    grid = fig.add_gridspec(2, 3, left=.068, right=.955, bottom=.14, top=.802, wspace=.20, hspace=.42)
    for row, outcome in enumerate(OUTCOMES):
        for col, sex in enumerate(SEXES):
            ax = fig.add_subplot(grid[row, col])
            label = "Both sexes combined" if sex == "Both" else sex
            format_axis(ax, f"{'ABCDEF'[row * 3 + col]}   {label}",
                        ("Prevalence" if row == 0 else "Annual incidence") + " change (%)" if col == 0 else None,
                        top=45)
            endings = []
            for scenario, _, color, ls, marker in SCENARIOS:
                allocation = "gbd_aligned_within_80plus" if scenario.startswith("national") else "native_age_detail"
                part = data.loc[data.population_scenario.eq(scenario) & data.within_80plus_allocation.eq(allocation)
                                & data.sex.eq(sex) & data.outcome.eq(outcome)].sort_values("forecast_year")
                x = np.r_[2023, part.forecast_year.to_numpy()]
                y = np.r_[0., part.percentage_change_from_2023.to_numpy()]
                ax.plot(x, y, color=color, marker=marker, ls=ls, lw=1.8, ms=4.8, mfc="white", mew=1, clip_on=False)
                endings.append((y[-1], color))
            # Keep close numerical endpoints legible without moving the curves.
            ordered = sorted(endings)
            label_y = [v for v, _ in ordered]
            for k in range(1, len(label_y)):
                label_y[k] = max(label_y[k], label_y[k - 1] + 4.1)
            if label_y[-1] > 43:
                label_y = [v - (label_y[-1] - 43) for v in label_y]
            for (value, color), ly in zip(ordered, label_y):
                ax.plot([2028.06, 2028.30, 2028.38], [value, ly, ly], color=color, lw=.65)
                ax.text(2028.45, ly, f"+{value:.2f}%", color=color, fontsize=9.4, va="center")
    fig.text(.5, .092, "Year", ha="center", fontsize=11)
    fig.text(.068, .045, "National 2024 baseline is backcast to 2023 along its UN growth path; ages 80+ use GBD-aligned allocation.",
             color=MUTED, fontsize=9.2)
    fig.text(.068, .024, "Population scenarios are conditional point projections, not uncertainty intervals.", color=MUTED, fontsize=9.2)
    export(fig, "saudi_forecast_percentage_change_sensitivity")


def captions():
    paragraphs = [
        "Standalone figure. Forecast percentage change in Parkinson’s disease case counts from 2023 to 2028 in Saudi Arabia.",
        "Panels A–C show male, female and combined-sex projections for residents aged 45 years and older, including citizens and non-citizens. Prevalence denotes prevalent cases in each year; incidence denotes new cases during each year. The adapted temporal convolutional network forecasts age-specific disease rates using a 2023 disease-data cutoff. Case counts apply the GBD 2023 population baseline advanced with age- and sex-specific United Nations World Population Prospects 2024 growth. Curves show point estimates; the 2023 reference is a modelled GBD estimate. Endpoint labels show changes at 2028. Counts beneath each panel are rounded to whole cases for display only.",
        "Percentage change = 100 × (projected count in year t / matched count in 2023 − 1). Every year is compared with 2023, not with the preceding year. Combined-sex changes are computed from summed counts, equivalent to weighting sex-specific changes by their 2023 case counts. They are not simple averages of male and female percentage changes. The 2023 plotting value is zero by definition. Increasing case counts reflect the joint influence of disease-rate trajectories and demographic change; these percentages are not changes in age-standardised rates or individual disease risk.",
        "Companion sensitivity figure. Panels A–C show prevalence and panels D–F annual incidence under the original GBD-aligned population, unaligned UN medium population and national 2024 population scenarios. Each scenario is compared with its own internally consistent 2023 baseline. The national 2024 population anchor is backcast to 2023 using the same UN growth path used in the saved projections; its 2023 count is therefore scenario-derived. Its open age group 80+ is allocated using the GBD-aligned age distribution. The older national 2022 anchor and the alternative UN within-80+ allocation are retained in the accompanying CSV. The primary figure uses the original GBD-aligned scenario to compare directly with the native GBD 2023 burden.",
        "No prediction intervals are assigned to these percentage changes. Valid uncertainty intervals would require joint uncertainty in the baseline, future disease rates and population path; scenario differences do not supply such intervals. The study’s previously documented interval-calibration limitations remain applicable. The projections were computed after 2023 and are conditional on the named demographic sources; a 2023 disease-data cutoff does not imply a forecast issued in 2023.",
        "Data provenance: preserved current-study adapted TCN projections; GBD 2023 Parkinson’s disease estimates; United Nations World Population Prospects 2024 medium population estimates and projections; and Saudi national population tables supplied in the Ministry of Health statistical yearbook downloads and attributed to the General Authority for Statistics in the study source audit. These figures introduce no external disease dataset. Full source citations remain in the study’s dataset bibliography and citation crosswalk. Forecast values retain their original stored decimal representation in forecast_percentage_change.csv; derived percentages retain calculation precision, which should not be interpreted as biological precision.",
        "Reproducibility: the export verifies every scenario’s 2024–2028 age-sex population against preserved outputs, reproduces the 60 previously saved original/UN percentage changes, checks combined-sex aggregation and preserves manuscript and source file hashes. Both figures are standalone artifacts and have not been added to the manuscript.",
    ]
    (OUT / "figure_caption.txt").write_text("\n\n".join(paragraphs) + "\n", encoding="utf-8")
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)
    for i, text in enumerate(paragraphs):
        p = doc.add_paragraph()
        p.add_run(text).bold = (i == 0)
    doc.save(OUT / "figure_caption.docx")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    sources = [SUMMARY, BASELINE, UN, NATIONAL, CELLS, PRIOR, ROOT / "forecast.csv"]
    sources += [ROOT / f"results/projections_v1/trials/SAU_{outcome}/baseline_2023_age_rates.csv"
                for outcome in OUTCOMES]
    before = {str(p.relative_to(ROOT)): sha(p) for p in sources}
    manuscript_before = manuscript_hashes()
    data = make_data()
    setup_style()
    main_figure(data)
    sensitivity_figure(data)
    captions()
    assert before == {str(p.relative_to(ROOT)): sha(p) for p in sources}
    assert manuscript_before == manuscript_hashes()
    image_metadata = {}
    for p in sorted(OUT.glob("*")):
        if p.suffix in [".png", ".tiff"]:
            with Image.open(p) as img:
                image_metadata[p.name] = {"size": list(img.size), "dpi": [float(v) for v in img.info.get("dpi", [])]}
    validation = {
        "status": "passed", "forecast_rows": len(data), "population_scenario_allocation_combinations": 6,
        "previous_percentage_changes_reproduced": 60, "forecast_population_cells_reproduced": 1320,
        "combined_sex_aggregation": "verified for counts and baseline-weighted percentage changes",
        "main_estimand": "percentage change in Saudi age-45-plus case counts from native GBD 2023 baseline",
        "source_files_unchanged": True, "source_sha256": before,
        "manuscript_files_unchanged": len(manuscript_before),
        "percentage_formula": "100 * (forecast_count / matched_baseline_count_2023 - 1)",
        "images": image_metadata,
        "output_sha256": {str(p.relative_to(ROOT)): sha(p) for p in sorted(OUT.glob("*"))
                          if p.is_file() and p.name != "validation.json"},
        "data_sha256": sha(ROOT / "forecast_percentage_change.csv"),
    }
    (OUT / "validation.json").write_text(json.dumps(validation, indent=2) + "\n", encoding="utf-8")
    primary = data.loc[data.population_scenario.eq(MAIN_SCENARIO)]
    print(primary.pivot(index="forecast_year", columns=["outcome", "sex"], values="percentage_change_from_2023").round(6).to_string())
    print("Validation passed; figures, caption and exact-value CSV saved.")


if __name__ == "__main__":
    main()
