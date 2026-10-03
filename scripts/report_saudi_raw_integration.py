"""Generate the supplementary report, publication figures and revised manuscript."""
from __future__ import annotations

import csv
import difflib
import importlib.util
import json
import os
import re
import shutil
from pathlib import Path

os.environ.setdefault('MPLCONFIGDIR','/tmp/gbd_raw_integration_mpl')
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image

from integrate_saudi_raw_data import ROOT, OUT, DATA, sha

REV = ROOT/'manuscript/ije_draft_v2_raw_integration'
FIG = OUT/'figures'
COLORS = ['#0072B2','#D55E00','#009E73','#7B3294']


def read(name):return pd.read_csv(OUT/name)


def document(title):
    d=Document();s=d.styles['Normal'];s.font.name='Times New Roman';s.font.size=Pt(11)
    d.add_heading(title,0)
    return d


def table(doc,title,headers,rows,note=''):
    doc.add_heading(title,2)
    t=doc.add_table(rows=1, cols=len(headers));t.style='Light Shading Accent 1'
    for c,text in zip(t.rows[0].cells,headers):c.text=str(text)
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=str(text)
    if note:doc.add_paragraph(note)


def save_figure(fig,stem):
    FIG.mkdir(exist_ok=True)
    for ext in ['pdf','svg','png','tiff']:
        args={'dpi':600 if ext=='tiff' else 180}
        if ext=='tiff':args['pil_kwargs']={'compression':'tiff_lzw'}
        fig.savefig(FIG/(stem+'.'+ext),facecolor='white',**args)
    plt.close(fig)


def figures(pop,counts,projection,chi,national):
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,
                         'svg.fonttype':'none','pdf.fonttype':42})
    fig,axes=plt.subplots(2,2,figsize=(12,9.2),layout='constrained')
    part=pop.loc[pop.age.eq('80+')].set_index(['year','sex']).loc[[(2022,'Male'),(2023,'Male'),(2022,'Female'),(2023,'Female')]]
    ax=axes[0,0]
    for off,key,label,color in [(-.2,'gbd_population','GBD-implied',COLORS[0]),(0,'un_population','UN WPP',COLORS[2]),(.2,'national_population','National source',COLORS[1])]:
        ax.scatter(np.arange(4)+off,part[key]/1000,s=65,label=label,color=color)
    ax.set_xticks(range(4),['Male\n2022','Male\n2023','Female\n2022','Female\n2023'])
    ax.set_ylabel('Population aged 80+ (thousands)');ax.set_ylim(0,100);ax.legend(frameon=False,fontsize=8)
    ax.set_title('A  Older population across sources',loc='left',fontweight='bold')
    ax=axes[0,1]
    part=counts.loc[counts.year.eq(2023)&counts.outcome.eq('prevalence')].set_index('sex').loc[['Male','Female']]
    for off,key,label,color in [(-.16,'native_80plus_share','GBD native',COLORS[0]),(.16,'national_80plus_share','National denominator',COLORS[1])]:
        ax.bar(np.arange(2)+off,part[key],width=.29,color=color,label=label)
        for i,val in enumerate(part[key]):ax.text(i+off,val+.7,f'{val:.1f}%',ha='center',fontsize=9)
    ax.set_xticks(range(2),['Male','Female']);ax.set_ylim(0,40);ax.set_ylabel('80+ share of prevalence at ages 45+ (%)')
    ax.legend(frameon=False,loc='upper left',fontsize=8);ax.set_title('B  2023 count composition sensitivity',loc='left',fontweight='bold')
    names=['gbd_2023_aligned_un_growth','un_medium_unaligned','national_2022_un_growth','national_2024_un_growth']
    part=projection.loc[projection.year.eq(2028)&projection.outcome.eq('prevalence')&projection.sex.eq('Both')&projection.family.eq('tcn_adapted')&projection.within_80plus_allocation.isin(['native_age_detail','gbd_aligned_within_80plus'])].set_index('scenario').loc[names]
    labels=['GBD 2023\n+ UN growth','UN medium','National 2022\n+ UN growth','National 2024\n+ UN growth']
    for ax,key,yl,title in [(axes[1,0],'count_45plus','Projected prevalent cases aged 45+ (thousands)','C  2028 conditional prevalence totals'),
                             (axes[1,1],'share_80plus','80+ share of projected prevalence at ages 45+ (%)','D  Similar totals can conceal different age profiles')]:
        vals=part[key]/(1000 if key=='count_45plus' else 1)
        ax.bar(range(4),vals,color=[COLORS[0],COLORS[2],COLORS[3],COLORS[1]],width=.62)
        for i,val in enumerate(vals):
            label=f'{val*1000:,.0f}' if key=='count_45plus' else f'{val:.2f}'
            ax.text(i,val+.7,label,ha='center',fontsize=10)
        ax.set_xticks(range(4),labels,fontsize=8);ax.set_ylabel(yl);ax.set_ylim(0,45)
        ax.set_title(title,loc='left',fontweight='bold',fontsize=10)
    save_figure(fig,'supplementary_figure_S34_demographic_integration')
    fig,axes=plt.subplots(1,2,figsize=(12,5.4),layout='constrained')
    groups=[('Female','Citizens'),('Female','Non-citizens'),('Male','Citizens'),('Male','Non-citizens')]
    q=read('chi_2024_population_composition_comparison.csv')
    q=q.loc[q.metric.eq('insured_persons')].set_index(['sex','nationality']).loc[groups]
    ax=axes[0];left=np.zeros(2)
    for (sex,nat),color in zip(groups,COLORS):
        r=q.loc[(sex,nat)];vals=np.array([r.national_composition_percent,r.chi_composition_percent])
        ax.barh([0,1],vals,left=left,color=color,label=f'{sex}, {nat.lower()}',height=.5)
        for i,v in enumerate(vals):ax.text(left[i]+v/2,i,f'{v:.1f}%',ha='center',va='center',color='white',fontsize=9)
        left+=vals
    ax.set_yticks([0,1],['National 2024','CHI Q2 2024']);ax.set_xlim(0,100);ax.set_xlabel('Share of source population (%)')
    ax.legend(frameon=False,loc='upper center',bbox_to_anchor=(.5,-.18),ncol=2,fontsize=8)
    ax.set_title('A  Sex and nationality composition',loc='left',fontweight='bold')
    ax=axes[1]
    part=chi.loc[chi.metric.eq('insured_persons')].copy()
    order=['12.2021','03.2022','12.2023','Q1-2024','Q2-2024','Q4-2025','Q2-2026']
    # Filename matching avoids interpreting mixed Arabic/English spreadsheet dates.
    keys=[]
    for token in order:
        matches=part.loc[part.file.str.contains(token,regex=False)]
        assert len(matches)==1;keys.append(matches.iloc[0])
    part=pd.DataFrame(keys)
    ax.scatter(range(7),part.male_percent,color=COLORS[0],marker='o',s=50,label='Male')
    ax.scatter(range(7),part.noncitizen_percent,color=COLORS[1],marker='s',s=50,label='Non-citizen')
    ax.set_xticks(range(7),['2021\nQ4','2022\nQ1','2023\nQ4','2024\nQ1','2024\nQ2','2025\nQ4','2026\nQ2'],fontsize=8)
    ax.set_ylim(0,100);ax.set_ylabel('Share of CHI insured persons (%)');ax.legend(frameon=False,loc='lower right')
    ax.set_title('B  Selected quarterly snapshots',loc='left',fontweight='bold')
    save_figure(fig,'supplementary_figure_S35_insurance_composition')


def report(pop,counts,projection,chi,national):
    d=document('Saudi raw-data integration: methods, findings and implications')
    d.add_paragraph('Exploratory supplementary analysis, 3 October 2026. All 19 files in the raw-data batch were assessed for their appropriate role. '
                    'Population and insurance observations were integrated into new analyses; documentation files informed source eligibility. '
                    'The original disease-rate models, primary comparisons and interval scores remain unchanged.')
    paragraphs=[
      ('Methods: population harmonisation',
       'The MOH 2022 table was reconstructed from its 68 age–sex–nationality cells and verified against Table 1-3 of the 2023 Statistical Yearbook, which credits GASTAT. '
       'It was combined with existing GASTAT 2023/2024 counts to form 204 cells. Totals included citizens and non-citizens. '
       'The 2022 export specifies a reference year but not an exact population date; the GASTAT methods define the later estimates at 1 July. '
       'The 840-cell GASTAT 2010–2024 export duplicated GCC-Stat numerically. Three date-like labels were mapped to 0–4, 5–9 and 10–14 solely in derived data, '
       'with exact cell matching retained. Four GCC-Stat cells labelled 65–69 equalled complete 65+ totals; the new 2022 linkage uses the verified MOH cells and flags this ambiguity. [D5; D4; D3; gastat_methods]'),
      ('Methods: fixed-rate count sensitivity',
       'For each sex, outcome and year (2022/2023), national and GBD/UN populations were harmonised to eight age bands, 45–49 through 75–79 and 80+. '
       'Within the open 80+ group, the reference rate was the GBD-population-weighted average of the four available age-specific rates. '
       'Alternative count = national population × reference rate / 100,000. These are fixed-rate counterfactual counts, not independently observed cases. '
       'Allocation limits assign the whole 80+ population to the minimum or maximum observed age-specific rate. They bound the consequence of an unspecified within-band mix while holding rates fixed; '
       'they are not confidence intervals, prediction intervals or comprehensive uncertainty bounds. No national 80–84/85–89/90–94/95+ counts are claimed. [D1; gbd_nonfatal; gbd_fatal]'),
      ('Methods: conditional projections',
       'Previously saved 2023-origin prevalence/incidence rate projections from the adapted TCN and the local/non-neural champions were reused without refitting. '
       'For national anchor year a (2022 or 2024), broad-age population in year t was N(a) × U(t)/U(a), using UN WPP 2024 growth in the same broad age and sex. '
       'The 80+ total was allocated alternatively by the GBD-2023-aligned UN-growth age profile or the UN age profile for year t. '
       'Both options exactly preserve the national broad-age total. The 2022 versus 2024 anchors also reveal sensitivity to the observed national age distribution. '
       'All four original TCN 2028 sex-combined count results were reproduced to numerical precision. The new scenarios use later-vintage information and are conditional projections, '
       'not prospective backtests or selected “best” forecasts. [D1; D2; wpp_methods; D4; D5]'),
      ('Methods: insurance composition',
       'All seven CHI quarterly workbooks were parsed from their preserved originals, including the ODS file. Subscriber and insured-person indicators were retained separately. '
       'Fourteen detailed-to-headline count reconciliations passed. All-age sex and nationality composition in Q2 2024 were compared with GASTAT 2024. '
       'The exported count-per-resident ratio is a descriptive snapshot ratio, not an insurance-coverage probability or person-year exposure. '
       'Original age categories were retained; 61 >=Age remained ambiguous and was excluded from any asserted oldest-age definition. '
       'The separate 2024 beneficiary-class CSV was reconciled independently and was not equated with a quarterly total. [chi_data; chi_class; chi_terms; D4]'),
      ('Findings: persistent oldest-age sensitivity',
       'In 2022, GBD-implied populations at ages 80+ were 38.23% below the MOH/GASTAT national count in males and 49.28% below in females. '
       'The corresponding 2023 differences were 37.39% and 49.65%. This shows that the broad discrepancy also occurs in the added year; it does not identify which demographic source is correct. '
       'At fixed 2023 GBD rates, national denominators raised conditional prevalent counts at ages 45+ by 12.41% in males and 18.18% in females. '
       'The oldest-age prevalence shares changed from 19.27% to 27.38% in males and from 17.24% to 28.97% in females. '
       'All six outcomes and both years are retained in Supplementary Table S9.'),
      ('Findings: similar total burden can conceal a different age profile',
       'For 2028, the adapted-TCN prevalence total was 31,844 under GBD-2023-aligned UN growth, 37,144 under unaligned UN medium populations, '
       '38,053 under the national-2022 anchor, and 37,165 under the national-2024 anchor. The national examples use the GBD-aligned within-80+ allocation. '
       'Their oldest-age shares were 24.73%, 29.55%, 35.22% and 34.25%, respectively. Thus the UN and national-2024 scenarios differed by only about 21 total cases '
       'while differing by 4.70 percentage points in the oldest-age share. This is a composition sensitivity, not a forecast-accuracy improvement. '
       'Using UN within-80+ weights instead changed the national-2024 prevalence total to 37,238 and its oldest-age share to 34.38%. '
       'Incidence and the two comparator families are included in Supplementary Table S11.'),
      ('Findings: insurance records represent a selected population',
       'In Q2 2024, the all-age CHI insured-person indicator summed to 12,381,657. Males comprised 73.10% and non-citizens 66.31%; all-age national population shares were 62.08% and 44.38%, respectively. '
       'Male non-citizens represented 55.59% of the insured-person snapshot versus 34.24% nationally. This identifies demographic differences in the source population, '
       'not observed Parkinson underdiagnosis or nationality-specific disease risk. A future claims validation must define a compatible insured population, ascertainment algorithm '
       'and observation period before comparing its outcomes with national forecasts. Both CHI indicators and all seven snapshots are retained in Supplementary Table S10.'),
      ('Data eligibility and limitations',
       'The NHRSP HISS and NCD files contain codebooks and a questionnaire, not observations. Neither inspected instrument identifies Parkinson’s disease specifically. '
       'The NPHIES visible top-10 summaries omit a Parkinson outcome series; omission from a top-10 list does not mean zero disease. The previously decoded embedded table contains '
       'aggregate service rows without the required age, sex, diagnosis and patient identifiers. The data-use agreement and access pages do not establish access approval. '
       'These files strengthen the eligibility assessment but cannot be added as disease-model observations. [nhrsp_hiss; nhrsp_ncd; chi_nphies; nhrsp_access]'),
      ('Interpretation and reporting decision',
       'The integrated data improve denominator handling, demonstrate that age composition matters beyond aggregate counts, and specify the demographic limits of claims-based validation. '
       'They provide no new observed Parkinson outcomes and therefore do not demonstrate improved forecast accuracy or interval calibration. National and GBD population definitions '
       'remain incompletely aligned; the 2022 reference date and within-80+ national detail are unresolved. Nationality is used for demographic description only. '
       'The primary age range remains 45+, the oldest ages remain visible, and the joint-primary and reliability conclusions are retained.'),
    ]
    for title,text in paragraphs:d.add_heading(title,1);d.add_paragraph(text)
    table(d,'Supplementary Table S9: national denominator sensitivity',
        ['Year','Sex','GBD population 80+','National population 80+','GBD difference (%)'],
        [[r.year,r.sex,f'{r.gbd_population:,.0f}',f'{r.national_population:,.0f}',f'{r.gbd_percent_difference_from_national:.2f}'] for r in pop.loc[pop.age.eq('80+')].itertuples()],
        'The complete machine-readable bundle includes all compatible ages and six outcomes. National 2022 reference date is not established; later GASTAT counts are mid-year.')
    q=chi.loc[chi.metric.eq('insured_persons')]
    table(d,'Supplementary Table S10: CHI insured-person composition',
        ['Period','Insured persons','Male (%)','Non-citizen (%)'],
        [[r.period.replace('\n',' '),f'{r.total:,}',f'{r.male_percent:.2f}',f'{r.noncitizen_percent:.2f}'] for r in q.itertuples()],
        'Selected snapshots; no continuous annual panel or Parkinson numerator. Source: Council of Health Insurance; authors’ calculations. Subscriber results are supplied separately.')
    q=projection.loc[projection.year.eq(2028)&projection.sex.eq('Both')&projection.family.eq('tcn_adapted')&projection.within_80plus_allocation.isin(['native_age_detail','gbd_aligned_within_80plus'])]
    labels={'gbd_2023_aligned_un_growth':'GBD 2023 + UN growth','un_medium_unaligned':'UN medium','national_2022_un_growth':'National 2022 + UN growth','national_2024_un_growth':'National 2024 + UN growth'}
    table(d,'Supplementary Table S11: conditional 2028 counts and composition',
        ['Outcome','Population scenario','Count aged 45+','80+ share (%)'],
        [[r.outcome,labels[r.scenario],f'{r.count_45plus:,.0f}',f'{r.share_80plus:.2f}'] for r in q.itertuples()],
        'The same saved TCN rate projections underlie every row. National rows assume a GBD-aligned within-80+ distribution; the full bundle includes UN-weight alternatives and both comparator families. No scenario is an independently verified forecast.')
    for stem,title,caption in [
      ('supplementary_figure_S34_demographic_integration','Supplementary Figure S34: demographic integration',
       'A, population at ages 80+ in 2022/2023. B, 2023 prevalence composition with GBD rates held fixed. C–D, 2028 prevalence totals and oldest-age shares under four population scenarios. '
       'National scenarios preserve observed broad-age totals and assume GBD-aligned within-80+ weights. The figure depicts conditional accounting, not new prevalence observations or prediction intervals. Sources: GBD 2023, UN WPP 2024, GASTAT and MOH; authors’ calculations.'),
      ('supplementary_figure_S35_insurance_composition','Supplementary Figure S35: insurance composition',
       'A, all-age sex–nationality composition of the Q2 2024 CHI insured-person snapshot and national 2024 population. B, male and non-citizen shares in seven selected CHI snapshots. '
       'No age-band interpolation or Parkinson outcome inference was performed. Sources: Council of Health Insurance and GASTAT; authors’ calculations.')]:
        d.add_heading(title,1);d.add_picture(str(FIG/(stem+'.png')),width=Inches(6.1));d.add_paragraph(caption)
    refs={r['id']:r for r in json.loads((ROOT/'study_design/dataset_citations_2026-10-01/references.json').read_text())}
    register=pd.read_csv(ROOT/'study_design/dataset_citations_2026-10-01/source_register.csv').set_index('id')
    keys=['D1','gbd_nonfatal','gbd_fatal','D2','wpp_methods','D3','D4','D5','gastat_methods','chi_data','chi_class','chi_terms','chi_nphies','nhrsp_hiss','nhrsp_ncd','nhrsp_access']
    d.add_heading('Source references',1)
    for key in keys:d.add_paragraph('['+key+'] '+register.loc[key,'citation'])
    d.save(OUT/'raw_data_integration_report.docx')
    (OUT/'interpretation.txt').write_text('\n\n'.join(title+'\n'+text for title,text in paragraphs),encoding='utf-8')


def manuscript_revision(projection):
    old=ROOT/'manuscript/ije_draft_v1'
    REV.mkdir(parents=True,exist_ok=True)
    body=(old/'body.md').read_text();before=body
    methods=('An exploratory source-integration analysis harmonised the MOH 2022 population table and GASTAT 2023–2024 estimates.[@moh_yearbook;@gastat] '
      'We held disease rates fixed while substituting national denominators and advanced national 2022/2024 anchors with United Nations growth. '
      'Ages ≥80 remained an open group with explicit alternative within-group weights. Seven CHI quarterly releases informed a separate insurance-composition analysis.[@chi] '
      'These analyses followed primary-result inspection (Supplementary Material S4; Supplementary Tables S9–S11).')
    needle='Analyses used Python, chronological input checks'
    assert needle in body;body=body.replace(needle,methods+'\n\n'+needle,1)
    original='GBD-implied Saudi populations aged ≥80 years were 37.39% below GASTAT for males and 49.65% below for females.'
    replacement=('GBD-implied Saudi populations aged ≥80 years were 37.39% below national counts for males and 49.65% below for females in 2023; '
                 'the added 2022 table showed corresponding differences of 38.23% and 49.28%.')
    assert original in body;body=body.replace(original,replacement)
    anchor='For Saudi mortality/disability, no supporting outcome met all four TCN endpoint comparisons.'
    added=('The national-2024 population anchor produced 37 165 conditional prevalent cases in 2028, close to the United Nations total, '
           'but with 34.25% aged ≥80 versus 29.55% (Supplementary Figure S34). CHI’s Q2 2024 all-age insured population was 73.10% male and 66.31% non-citizen, '
           'compared with 62.08% and 44.38% nationally (Supplementary Figure S35).')
    assert anchor in body;body=body.replace(anchor,added+'\n\n'+anchor,1)
    discussion='The denominator sensitivity demonstrates why changes in case counts cannot be interpreted directly as changes in disease risk.'
    replacement=('The denominator sensitivity demonstrates why case-count changes cannot be interpreted directly as disease-risk changes. '
      'Similar totals concealed different oldest-age composition. CHI’s demographic selection also means that future claims validation requires compatible populations and ascertainment.')
    assert discussion in body;body=body.replace(discussion,replacement)
    limitations='Finally, aggregate data cannot establish biological causation, individual prognosis, nationality-specific risk or validated staffing requirements.'
    replacement=('National rebasing supplies conditional counts, not independent disease validation or repaired intervals; within-80+ national detail remains unavailable. '
                 'Aggregate data cannot establish biological causation, individual prognosis, nationality-specific risk or validated staffing requirements.')
    assert limitations in body;body=body.replace(limitations,replacement)
    # Compress execution and exploratory-rule detail already fully documented in the supplement.
    compress={
      'Supplementary Software S1 and Supplementary Material S3 document execution, including the failed GPU standardised-rate attempt and successful matched CPU recovery.':'Supplementary Software S1 and Supplementary Material S3 document execution.',
      'The locally fixed exploratory rule required ≥5% lower rate WIS in each sex, comparator and point-error safeguards, no worse distance from nominal 80% coverage, and oldest-age/derived-burden safeguards, including a twofold width ceiling. These tolerances were research decisions, not clinical thresholds (Supplementary Material S2).':'The exploratory rule required ≥5% lower rate WIS in both sexes and safeguards for point error, coverage, oldest-age/derived burden and interval width (Supplementary Material S2).',
      'The 203-location roster matched the GBD World Bank region/income classifications, comprising 201 of 204 main GBD national locations plus Hong Kong and Macao. This clarified geographic coverage without changing the donor pool; independent disease-value verification remained limited to seven locations (Table 1; Supplementary Material S1).':'The 203-location roster matched GBD World Bank classifications: 201 main national locations plus Hong Kong and Macao. Disease-value verification remained limited to seven locations (Table 1; Supplementary Material S1).',
      'These projections were generated in 2026 from a 2023 disease-data cutoff.':'Projections used a 2023 disease-data cutoff and were generated in 2026.',
    }
    for source,target in compress.items():
        assert source in body,source;body=body.replace(source,target)
    body=body.replace('source code is publicly available.[@software]','core forecasting source code is publicly available.[@software]')
    # Reuse the existing styled writer without rerunning models, figures or retired comparison exports.
    spec=importlib.util.spec_from_file_location('ije_builder',ROOT/'scripts/build_ije_manuscript.py')
    builder=importlib.util.module_from_spec(spec);spec.loader.exec_module(builder)
    builder.OUT=REV;builder.BODY=body
    abstract=(old/'abstract.md').read_text()
    abstract=abstract.replace('subsequent distribution-mixture analyses were exploratory.',
        'subsequent distribution-mixture, population and insurance analyses were exploratory.')
    abstract=abstract.replace('\n## Conclusions',
        '\nNational and United Nations population scenarios yielded similar 2028 totals but ≥80-year shares of 34.25% and 29.55%.\n\n## Conclusions')
    builder.ABSTRACT=abstract
    assert builder.word_count(abstract.removeprefix('# Abstract\n'),include_headings=True)<=250
    builder.REFERENCES=json.loads((old/'references.json').read_text())
    source=pd.read_csv(ROOT/'study_design/dataset_citations_2026-10-01/source_register.csv').set_index('id')
    for key,sourcekey in [('moh_yearbook','D5'),('chi','chi_data')]:
        builder.REFERENCES[key]={'text':source.loc[sourcekey,'citation'],'url':source.loc[sourcekey,'url'],
          'status':'Verified provider source, used in the exploratory Saudi raw-data integration','type':'data'}
    builder.KEYS=[]
    for match in builder.CITE_RE.finditer(body):
        for key in match.group(1).replace('@','').split(';'):
            if key not in builder.KEYS:builder.KEYS.append(key)
    builder.REFNUM={k:i+1 for i,k in enumerate(builder.KEYS)}
    assert set(builder.KEYS)==set(builder.REFERENCES)
    count=builder.word_count(body,include_headings=True)
    assert count<=3000, f'New manuscript needs compression: {count} words'
    for fn in ['front_matter.json','tables.docx','missing_references.csv','missing_references.docx']:
        shutil.copyfile(old/fn,REV/fn)
    (REV/'abstract.md').write_text(abstract)
    (REV/'body.md').write_text(body)
    (REV/'references.json').write_text(json.dumps(builder.REFERENCES,indent=2,ensure_ascii=False))
    builder.manuscript()
    # Update the supplement declaration written by the common writer.
    d=Document(REV/'manuscript.docx')
    for p in d.paragraphs:
        if 'Supplementary Tables S1–S8' in p.text:
            p.text=p.text.replace('Supplementary Tables S1–S8','Supplementary Tables S1–S11').replace('Supplementary Material S1–S3','Supplementary Material S1–S4')
        if p.text.startswith('Analysis source code is available from the versioned public repository'):
            p.text=p.text.replace('Analysis source code','Core forecasting source code')+' New demographic and insurance integration scripts accompany Supplementary Software S1.'
    d.save(REV/'manuscript.docx')
    missing=pd.read_csv(REV/'missing_references.csv').to_dict('records')
    missing.append(dict(id='R7',item='CHI quarterly resource metadata',
        unresolved_detail='Provider collection and raw file hashes are documented; exact quarterly download URLs and publication/update timestamps remain incomplete.',
        action='Recover source export metadata before deposition; current attribution uses the verified CHI collection.'))
    missing.append(dict(id='R8',item='MOH 2022 population reference date',
        unresolved_detail='All 68 cells match Yearbook 2023 Table 1-3; exact population reference date and portal export timestamp are unverified.',
        action='Retain the 2022 reference-year qualification and avoid asserting full mid-year equivalence with other providers.'))
    pd.DataFrame(missing).to_csv(REV/'missing_references.csv',index=False)
    m=Document(REV/'missing_references.docx')
    m.add_heading('Additional metadata from Saudi raw-data integration',1)
    for r in missing[-2:]:m.add_paragraph(r['id']+' — '+r['item']+'. '+r['unresolved_detail']+' '+r['action'])
    m.save(REV/'missing_references.docx')
    (REV/'revision.diff').write_text(''.join(difflib.unified_diff(before.splitlines(True),body.splitlines(True),fromfile='ije_draft_v1/body.md',tofile='ije_draft_v2_raw_integration/body.md')))
    return {'main_text_words':count,'main_figures':4,'main_tables':4,'references':len(builder.KEYS),
            'abstract_words':builder.word_count(builder.ABSTRACT.removeprefix('# Abstract\n'),include_headings=True)}


def supplementary_and_provenance(info):
    old=ROOT/'manuscript/ije_draft_v1'
    existing=pd.read_csv(old/'supplementary_index.csv').to_dict('records')
    new=[]
    def add(identifier,title,path,kind='file'):
        new.append(dict(supplementary_identifier=identifier,title=title,archive_member=path,member_type=kind))
    root='reports/saudi_raw_integration_v1/'
    for name in ['population_comparison_2022_2023.csv','national_denominator_count_sensitivity_cells.csv','national_denominator_count_sensitivity_summary.csv','moh_2022_source_cells.csv','moh_gccstat_age_key_audit.csv','gastat_840_cell_duplicate_check.csv']:
        add('Supplementary Table S9','National population and fixed-rate burden sensitivity',root+name)
    for name in ['chi_quarter_composition.csv','chi_2024_population_composition_comparison.csv','chi_original_age_composition.csv','chi_excluded_total_rows.csv']:
        add('Supplementary Table S10','Insurance composition and source eligibility',root+name)
    for name in ['conditional_projection_cells_2024_2028.csv','conditional_projection_summary_2024_2028.csv','original_projection_reproduction.csv']:
        add('Supplementary Table S11','Nationally anchored conditional projections',root+name)
    for num,stem in [(34,'demographic_integration'),(35,'insurance_composition')]:
        add(f'Supplementary Figure S{num}',stem.replace('_',' '),root+f'figures/supplementary_figure_S{num}_{stem}.pdf')
    add('Supplementary Material S4','Saudi raw-data integration methods and interpretation',root+'raw_data_integration_report.docx')
    add('Supplementary Material S4','Source eligibility and hash validation',root+'validation.json')
    add('Supplementary Material S4','Scientific validation checks',root+'scientific_validation.json')
    add('Supplementary Material S4','Source eligibility and hash validation',root+'data_use_and_eligibility.csv')
    add('Supplementary Material S4','Exploratory analysis specification','study_design/raw_data_integration_2026-10-03.json')
    add('Supplementary Material S4','Prepared demographic and insurance tables','data/processed/saudi_raw_integration_v1/','directory')
    for fn in ['integrate_saudi_raw_data.py','report_saudi_raw_integration.py']:
        add('Supplementary Software S1','Saudi raw-data integration scripts','scripts/'+fn)
    add('Supplementary Software S1','Scientific validation checks','tests/test_saudi_raw_integration.py')
    pd.DataFrame(existing+new).to_csv(REV/'supplementary_index.csv',index=False)
    d=Document(old/'supplementary_index.docx');d.add_heading('Saudi raw-data integration — 3 October 2026',1)
    for identifier in dict.fromkeys(r['supplementary_identifier'] for r in new):
        d.add_heading(identifier,2)
        for r in new:
            if r['supplementary_identifier']==identifier:d.add_paragraph(r['title']+': '+r['archive_member'])
    d.save(REV/'supplementary_index.docx')
    for r in new:assert (ROOT/r['archive_member']).exists(),r
    # Link all new derived outputs back to their source citations.
    cross=[]
    for folder in [OUT,DATA]:
        for p in sorted(folder.glob('*.csv')):
            if p.name in ['derived_file_citations.csv','output_checksums.csv']:continue
            if p.name=='data_use_and_eligibility.csv':ids='D4;moh_population;chi_data;chi_class;chi_nphies;nhrsp_ncd;nhrsp_hiss;nhrsp_access;moh_access;gastat_access;chi_terms'
            elif 'chi' in p.name:ids='chi_data;chi_class;D4'
            elif 'projection' in p.name:ids='D1;D2;D4;D5;moh_population'
            elif 'moh' in p.name or '2022_corrected' in p.name:ids='D5;moh_population;D3'
            elif 'gastat' in p.name:ids='gastat_series;D3'
            else:ids='D1;D2;D3;D4;D5;moh_population'
            cross.append(dict(archive_member=p.relative_to(ROOT).as_posix(),reference_ids=ids,sha256=sha(p),role='Exploratory supporting analysis'))
    pd.DataFrame(cross).to_csv(OUT/'derived_file_citations.csv',index=False)
    manifest=json.loads((OUT/'validation.json').read_text())
    assert all(sha(ROOT/p)==h for p,h in manifest['source_sha256'].items())
    assert '[@' not in '\n'.join(p.text for p in Document(REV/'manuscript.docx').paragraphs)
    for stem in ['supplementary_figure_S34_demographic_integration','supplementary_figure_S35_insurance_composition']:
        with Image.open(FIG/(stem+'.tiff')) as im:
            assert im.width==7200 and round(im.info['dpi'][0])==600
    info.update(original_primary_results_unchanged=True,new_disease_models_fitted=0,
                source_hashes_rechecked=True,new_supplementary_tables=3,new_supplementary_figures=2,
                docx_readback_checked=True,docx_page_rendering_performed=False,
                ije_guidance_checked='2026-10-03',guidance_url='https://academic.oup.com/ije/pages/General_Instructions')
    (REV/'validation.json').write_text(json.dumps(info,indent=2))
    (REV/'revision_notes.txt').write_text(
       'SAUDI RAW-DATA INTEGRATION REVISION — 3 OCTOBER 2026\n\n'
       'This revision adds new demographic and insurance analyses to Methods, Results, Discussion and Limitations. '
       'The complete technical report is Supplementary Material S4; new tables are S9–S11 and figures S34–S35.\n'
       'Original primary disease models, results, four main figures and four main tables are unchanged. '
       'Use the original main figure files in manuscript/ije_draft_v1/figures/; no alternative main figure set was created. '
       'The structured abstract retains the primary findings and adds the conditional oldest-age-composition result.\n'
       f'Main text: {info["main_text_words"]} words; abstract: {info["abstract_words"]} words; {info["references"]} references. '
       'IJE original-article guidance checked 3 October 2026: 3000 main-text words, 250 abstract words, eight main displays, 50 references.\n'
       'The additions are exploratory and supply neither new disease observations nor independent validation. '
       'Source metadata still requiring confirmation remain in the dataset citation package; no release-contradiction content was added.\n'
       'Main DOCX readback was checked; Word pagination was not rendered. Original version remains available.\n',encoding='utf-8')
    print(json.dumps(info))


def main():
    pop=read('population_comparison_2022_2023.csv');counts=read('national_denominator_count_sensitivity_summary.csv')
    projection=read('conditional_projection_summary_2024_2028.csv');chi=read('chi_quarter_composition.csv')
    national=pd.read_csv(DATA/'national_population_2022_2024.csv')
    figures(pop,counts,projection,chi,national)
    report(pop,counts,projection,chi,national)
    info=manuscript_revision(projection)
    supplementary_and_provenance(info)


if __name__=='__main__':main()
