"""Reproduce supporting CSV extracts from preserved public source files.

Does not modify raw files or the existing Parkinson's disease outcomes.
Run with bundled Python and pandas/openpyxl. See README for analytic limitations.
"""
from pathlib import Path
import hashlib,json,re,zipfile,xml.etree.ElementTree as ET
import numpy as np
import pandas as pd
import openpyxl

ROOT=Path(__file__).resolve().parent
RAW=ROOT/'raw'
OUT=ROOT/'processed'
OUT.mkdir(exist_ok=True)
GCC={48:'Bahrain',414:'Kuwait',512:'Oman',634:'Qatar',682:'Saudi Arabia',784:'United Arab Emirates'}
REPORT={}

def save(df,name):
 path=OUT/name
 df.to_csv(path,index=False,encoding='utf-8-sig',float_format='%.12g')
 REPORT[name]={'rows':len(df),'columns':list(df.columns),'bytes':path.stat().st_size}
 return df

def prepare_un():
 filename='WPP2024_PopulationByAge5GroupSex_Medium.csv.gz'
 parts=[];total_rows=0
 for chunk in pd.read_csv(RAW/filename,chunksize=100000,low_memory=False):
  total_rows+=len(chunk)
  parts.append(chunk[(chunk.LocTypeName=='Country/Area') & chunk.Time.between(1990,2050)])
 data=pd.concat(parts,ignore_index=True)
 keys=['LocID','VarID','Time','AgeGrpStart']
 assert not data.duplicated(keys).any()
 with zipfile.ZipFile(RAW/'WPP2024_CSV_files_update.zip') as z:
  update=pd.read_csv(z.open('WPP2024_PopulationByAge5GroupSex_Medium_Update.csv'))
  update=update[update.Time.between(1990,2050)]
 replaced=data.LocID.isin(update.LocID).sum()
 data=pd.concat([data[~data.LocID.isin(update.LocID)],update],ignore_index=True)
 assert not data.duplicated(keys).any()
 assert data.Variant.eq('Medium').all()
 assert data[['PopMale','PopFemale','PopTotal']].notna().all().all()
 assert (data[['PopMale','PopFemale','PopTotal']]>=0).all().all()
 rounding_diff=(data.PopMale+data.PopFemale-data.PopTotal).abs().max()
 assert rounding_diff<0.0021 # Original population rounded to 0.001 thousand.
 gcc=data[data.LocID.isin(GCC)].copy()
 assert len(gcc)==6*61*21
 assert gcc.groupby(['LocID','Time']).AgeGrpStart.nunique().eq(21).all()
 rename={'LocID':'un_location_id','ISO3_code':'iso3','Location':'location','Time':'year','AgeGrp':'age_group','AgeGrpStart':'age_start','AgeGrpSpan':'age_span','Variant':'variant'}
 def longify(wide):
  ids=list(rename)
  long=wide[ids+['PopMale','PopFemale','PopTotal']].melt(id_vars=ids,var_name='sex',value_name='population_thousands_original').rename(columns=rename)
  long['sex']=long.sex.map({'PopMale':'Male','PopFemale':'Female','PopTotal':'Both'})
  long['population_persons']=long.population_thousands_original*1000
  long['period_type']=np.where(long.year<=2023,'Estimate','Projection')
  long['source']='UN WPP 2024; 1 July; medium variant'
  return long.sort_values(['location','year','sex','age_start']).reset_index(drop=True)
 glong=longify(gcc)
 save(glong,'un_wpp2024_gcc_age_sex_1990_2050.csv')
 save(glong[glong.location=='Saudi Arabia'],'un_wpp2024_saudi_age_sex_1990_2050.csv')
 # Optional GBD-compatible open age band: combine 95-99 with 100+.
 gbd=glong.copy();gbd.loc[gbd.age_start>=95,'age_group']='95+';gbd.loc[gbd.age_start>=95,'age_start']=95
 group=['un_location_id','iso3','location','year','sex','age_group','age_start','variant','period_type','source']
 gbd=gbd.groupby(group,as_index=False)[['population_thousands_original','population_persons']].sum()
 save(gbd[gbd.age_start>=45],'un_wpp2024_gcc_45plus_gbd_age_bands.csv')
 def features(long):
  idx=['un_location_id','iso3','location','year','sex','period_type']
  result=long.groupby(idx,as_index=False).population_persons.sum().rename(columns={'population_persons':'population_all_ages'})
  for age in [45,65,80]:
   sub=long[long.age_start>=age].groupby(idx,as_index=False).population_persons.sum().rename(columns={'population_persons':f'population_{age}plus'})
   result=result.merge(sub,on=idx,validate='one_to_one');result[f'share_{age}plus']=result[f'population_{age}plus']/result.population_all_ages
  result['source']='UN WPP 2024; 1 July; medium variant'
  return result
 save(features(glong),'un_wpp2024_gcc_demographic_features_1990_2050.csv')
 donors=features(longify(data[data.Time<=2023]))
 save(donors,'un_wpp2024_global_demographic_features_1990_2023.csv')
 REPORT['un_validation']={'raw_rows':total_rows,'gcc_locations':sorted(gcc.Location.unique()),'gcc_year_min':int(gcc.Time.min()),'gcc_year_max':int(gcc.Time.max()),'age_bands_per_location_year':21,'unit_original':'thousands of persons','unit_derived':'persons','max_sex_additivity_difference_thousands':float(rounding_diff),'correction_locations':sorted(update.Location.unique()),'rows_replaced_with_official_corrections':int(replaced),'gcc_affected_by_corrections':bool(update.LocID.isin(GCC).any()),'global_country_area_count':int(donors.location.nunique())}

def prepare_gccstat():
 data=pd.read_csv(RAW/'gccstat_population_age_sex_nationality.csv')
 data.insert(0,'source_row',np.arange(len(data))+2)
 names={'Emirates':'United Arab Emirates'}
 data['location']=data.COUNTRY.replace(names)
 data['nationality']=data.NATIONALITY.replace({'Citzens':'Citizens','Non Citzens':'Non-citizens'})
 data['sex']=data.GENDER.replace({'Total':'Both'})
 data['population_persons_original']=data.OBS_VALUE
 data['zero_value_requires_review']=data.OBS_VALUE.eq(0)
 data['missing_value']=data.OBS_VALUE.isna()
 data['source']='GCC-Stat population export retrieved 2026-09-26'
 keys=['COUNTRY','NATIONALITY','GENDER','AGE','TIME_PERIOD']
 assert not data.duplicated(keys).any()
 save(data[data.location.isin(GCC.values())],'gccstat_population_labeled.csv')
 save(data[data.location=='Saudi Arabia'],'gccstat_saudi_population_age_sex_nationality_2010_2024.csv')
 coverage=data.groupby(['location','nationality','sex']).agg(rows=('AGE','size'),first_year=('TIME_PERIOD','min'),last_year=('TIME_PERIOD','max'),zero_values=('zero_value_requires_review','sum'),missing_values=('missing_value','sum')).reset_index()
 save(coverage,'gccstat_coverage_and_quality.csv')
 REPORT['gccstat_validation']={'rows_raw':len(data),'zero_values_preserved_and_flagged':int(data.zero_value_requires_review.sum()),'missing_values':int(data.missing_value.sum()),'open_age_band':'80+','overlapping_aggregate_age_groups':['All Ages','0-14','15-64','65+'],'nationality_categories':['Citizens','Non-citizens','Total'],'warning':'Do not sum aggregate age bands with their constituent bands. Do not treat unverified zero populations as observed absence.'}

def prepare_moh():
 rows=[];index=[]
 for year,sheet in [(2023,'5.'),(2024,'2-19')]:
  file=f'saudi_moh_yearbook_{year}.xlsx'
  wb=openpyxl.load_workbook(RAW/file,read_only=True,data_only=True)
  ws=wb[sheet]
  targets=[r for r in ws.iter_rows() if any(c.value=='Neurology' for c in r)]
  assert len(targets)==1
  row=targets[0]
  for start,sector in [(2,'MOH'),(6,'Other governmental'),(10,'Private'),(14,'Total KSA')]:
   values=[]
   for offset,category in enumerate(['Resident','Registrar','Consultant','Total']):
    cell=row[start+offset-1];value=cell.value
    assert isinstance(value,(int,float)) and value>=0
    values.append(value)
    rows.append(dict(year=year,specialty='Neurology',sector=sector,professional_category=category,physician_count=value,source_file=file,source_sheet=sheet,source_cell=cell.coordinate,unit='persons',role='Health-system context; not PD-specific care capacity'))
   assert sum(values[:3])==values[3]
  assert row[4].value+row[8].value+row[12].value==row[16].value
  for ws in wb:
   title=[];matches=[]
   for n,r in enumerate(ws.iter_rows(values_only=True),1):
    for c,value in enumerate(r,1):
     if isinstance(value,str):
      if n<=3 and len(value)>35 and re.search('[A-Za-z]{4}',value):title.append(value.strip())
      if any(k in value.lower() for k in ['neurology','physiotherapy','occupational therapy']):matches.append(value.strip())
   if matches:index.append(dict(source_file=file,sheet=ws.title,title=' | '.join(dict.fromkeys(title)),relevant_terms=' | '.join(dict.fromkeys(matches))))
  wb.close()
 save(pd.DataFrame(rows),'saudi_moh_neurologists_2023_2024.csv')
 save(pd.DataFrame(index),'saudi_moh_relevant_sheet_index.csv')

def prepare_gastat():
 file='GASTAT_Population_Estimates_EN.xlsx'
 wb=openpyxl.load_workbook(RAW/file,read_only=True,data_only=True)
 ws=wb[wb.sheetnames[0]]
 rows=[]
 groups=[(2,2024,'Citizens'),(5,2024,'Non-citizens'),(8,2024,'Total'),(11,2023,'Citizens'),(14,2023,'Non-citizens'),(17,2023,'Total')]
 for source_row in range(4,22):
  age=ws.cell(source_row,1).value
  assert age=='Total' or re.fullmatch(r'\d+-\d+|\d+\+',age)
  for start,year,nationality in groups:
   cells=[ws.cell(source_row,start+i) for i in range(3)]
   assert cells[0].value+cells[1].value==cells[2].value
   for cell,sex in zip(cells,['Female','Male','Both']):
    rows.append(dict(location='Saudi Arabia',year=year,nationality=nationality,sex=sex,age_group=('All Ages' if age=='Total' else age),age_group_original=age,age_start=(None if age=='Total' else int(re.match(r'\d+',age).group())),population_persons=cell.value,source_file=file,source_sheet=ws.title,source_cell=cell.coordinate,reference_date='mid-year',source='GASTAT Population Estimates 2024 publication tables'))
 df=pd.DataFrame(rows)
 totals=df.pivot(index=['year','sex','age_group'],columns='nationality',values='population_persons')
 assert (totals.Citizens+totals['Non-citizens']).eq(totals.Total).all()
 for _,g in df.groupby(['year','nationality','sex']):
  assert g[g.age_group!='All Ages'].population_persons.sum()==g[g.age_group=='All Ages'].population_persons.iloc[0]
 assert len(df)==324
 save(df,'gastat_saudi_population_age_sex_nationality_2023_2024.csv')
 # Harmonize labels only for a source comparison; keep raw workbook intact.
 compare=df.copy();compare.loc[compare.age_start>=80,'age_group']='80+'
 compare=compare.groupby(['year','nationality','sex','age_group'],as_index=False).population_persons.sum()
 gcc=pd.read_csv(OUT/'gccstat_saudi_population_age_sex_nationality_2010_2024.csv')
 gcc=gcc.rename(columns={'TIME_PERIOD':'year','AGE':'age_group','population_persons_original':'gccstat_population_persons'})
 compare=compare.merge(gcc[['year','nationality','sex','age_group','gccstat_population_persons']],on=['year','nationality','sex','age_group'],how='left',validate='one_to_one')
 compare['gccstat_minus_gastat']=compare.gccstat_population_persons-compare.population_persons
 save(compare,'gastat_vs_gccstat_population_comparison_2023_2024.csv')
 REPORT['gastat_validation']={'rows':len(df),'years':[2023,2024],'specific_age_bands':17,'aggregate_age_band':'All Ages (original label Total)','open_age_band':'80+','sex_and_nationality_and_age_totals_verified':True,'comparison_missing_gccstat_rows':int(compare.gccstat_population_persons.isna().sum()),'comparison_nonzero_differences':int(compare.gccstat_minus_gastat.dropna().ne(0).sum()),'comparison_max_absolute_difference':float(compare.gccstat_minus_gastat.abs().max()),'note':'GCC-Stat and GASTAT need not be independent population sources; comparison documents concordance. The separate GASTAT PDF pyramid displays 80-84 and 85+, while its linked workbook uses 80+.'}
 wb.close()

def prepare_published_tables():
 ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
 rows=[]
 triple=re.compile(r'^\s*(-?[\d.]+)\s*\(\s*(-?[\d.]+)\s*,\s*(-?[\d.]+)\s*\)\s*$')
 for num,measure in [(5,'Prevalence'),(6,'Deaths'),(7,'DALYs')]:
  file=f'12889_2023_15018_MOESM{num}_ESM.docx'
  with zipfile.ZipFile(RAW/file) as z:root=ET.fromstring(z.read('word/document.xml'))
  tables=root.findall('.//w:tbl',ns);assert len(tables)==1
  for rownum,tr in enumerate(tables[0].findall('w:tr',ns),1):
   cells=[''.join(n.text or '' for n in c.findall('.//w:t',ns)).strip() for c in tr.findall('w:tc',ns)]
   if len(cells)!=6 or not triple.match(cells[1]):continue
   for i,(year,metric) in enumerate([(1990,'Number'),(1990,'Age-standardized rate'),(2019,'Number'),(2019,'Age-standardized rate'),('1990-2019','Percent change in ASR')],1):
    m=triple.match(cells[i]);assert m,(file,rownum,cells[i])
    value,lower,upper=map(float,m.groups());assert lower<=value<=upper
    rows.append(dict(location=cells[0],measure=measure,sex='Both',year_or_period=year,metric=metric,value=value,lower_95ui=lower,upper_95ui=upper,unit='per 100000' if metric=='Age-standardized rate' else ('percent' if metric.startswith('Percent') else ('DALYs' if measure=='DALYs' else 'persons')),original_cell_text=cells[i],gbd_release='GBD 2019',source_file=file,source_table_row=rownum,doi='10.1186/s12889-023-15018-x',role='Secondary release comparison; not independent validation'))
 df=pd.DataFrame(rows);assert len(df)==22*3*5
 save(df,'published_mena_gbd2019_tables_1990_2019.csv')
 save(df[df.location.isin(GCC.values())],'published_gcc_gbd2019_tables_1990_2019.csv')
 # Verify 2019 supplementary values against the separately downloaded main article XML.
 root=ET.parse(RAW/'mena_gbd2019_article.xml').getroot()
 checked=0
 for tr in root.findall('.//table-wrap/table/tbody/tr'):
  cells=[''.join(c.itertext()).strip() for c in tr]
  if len(cells)!=10:continue
  loc=cells[0]
  for off,measure in [(1,'Prevalence'),(4,'Deaths'),(7,'DALYs')]:
   for delta,metric in [(0,'Number'),(1,'Age-standardized rate')]:
    normalized=re.sub(r'(?<=\d),(?=\d{3}(?:\D|$))','',cells[off+delta])
    nums=re.findall(r'-?\d+(?:\.\d+)?',normalized)
    if len(nums)!=3:raise ValueError(('Unexpected published main table cell',cells[off+delta]))
    selected=df[(df.location==loc)&(df.measure==measure)&(df.year_or_period==2019)&(df.metric==metric)]
    assert len(selected)==1,(loc,measure,metric)
    assert np.allclose(selected[['value','lower_95ui','upper_95ui']].iloc[0].astype(float),list(map(float,nums)))
    checked+=1
 REPORT['published_table_validation']={'rows_extracted':len(df),'locations':int(df.location.nunique()),'main_article_2019_cells_crosschecked':checked,'sex':'Both; no sex-specific annual numeric series extracted','supplements_1_to_3':'Figures only, no values digitized','supplement_4':'Binary Word table preserved; not converted'}

if __name__=='__main__':
 manifest=json.loads((ROOT/'manifest.json').read_text(encoding='utf-8'))
 for item in manifest:
  file=RAW/item['filename']
  assert file.stat().st_size==item['bytes']
  assert hashlib.sha256(file.read_bytes()).hexdigest()==item['sha256']
 prepare_un();prepare_gccstat();prepare_moh();prepare_gastat();prepare_published_tables()
 REPORT['raw_downloads']={'files':len(manifest),'bytes':sum(i['bytes'] for i in manifest)}
 (ROOT/'validation_report.json').write_text(json.dumps(REPORT,indent=2,ensure_ascii=False),encoding='utf-8')
 print(json.dumps(REPORT,indent=2,ensure_ascii=False))
